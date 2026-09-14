"""Naqaa: Arabic lead ingestion, normalization, reporting and Streamlit UI.

Run with: streamlit run app.py
Pure processing functions are importable without starting the interface.
Customer files and exports remain in session memory; fonts are bundled assets.
"""
from __future__ import annotations

import codecs
import copy
import csv
import hashlib
import io
import json
import math
import numbers
import re
import struct
import threading
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date as _core_date, datetime as _core_datetime, time as _core_time
from decimal import Decimal, InvalidOperation
from pathlib import PurePath
from urllib.parse import parse_qs, unquote, urlsplit

import openpyxl
import pandas as pd
import phonenumbers
import xlrd
from charset_normalizer import from_bytes as _charset_from_bytes
from defusedxml.ElementTree import fromstring as _safe_xml_fromstring
from rapidfuzz import fuzz as _fuzz


PRODUCT_NAME = "Naqaa | نقاء"
PRODUCT_SLUG = "naqaa"

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_FILES = 5
MAX_ROWS = 25000
MAX_COLUMNS = 200
MAX_CELLS = 1500000
PDF_BATCH_SIZE = 500

FIELD_LABELS = {
    "name": "الاسم",
    "phone": "رقم الهاتف",
    "whatsapp": "رقم واتساب",
    "job": "الوظيفة",
    "budget": "الميزانية",
    "timeline": "موعد الشراء",
    "city": "المدينة",
}
COUNTRY_OPTIONS = {
    "EG": "مصر (+20)",
    "SA": "السعودية (+966)",
    "AE": "الإمارات (+971)",
    "QA": "قطر (+974)",
    "KW": "الكويت (+965)",
    "BH": "البحرين (+973)",
    "OM": "عُمان (+968)",
    "US": "الولايات المتحدة (+1)",
    "GB": "المملكة المتحدة (+44)",
    "JO": "الأردن (+962)",
    "LB": "لبنان (+961)",
    "IQ": "العراق (+964)",
}
CURRENCY_OPTIONS = ["EGP", "SAR", "AED", "QAR", "KWD", "BHD", "OMR", "USD", "EUR"]
CLEAN_COLUMNS = [
    "lead_id", "source", "source_row", "name", "phone_raw", "phone_e164",
    "phone_display", "phone_valid", "phone_region", "phone_error", "tel_url",
    "whatsapp_raw", "whatsapp_e164", "whatsapp_display", "whatsapp_valid",
    "whatsapp_error", "wa_url", "whatsapp_origin", "job", "budget_raw",
    "budget_value", "budget_currency", "vip", "timeline", "readiness", "city",
    "details", "original", "issues", "duplicate",
]


class UserDataError(ValueError):
    """A safe, Arabic explanation of invalid or unsafe input data."""


@dataclass
class ParsedTable:
    source_id: str
    filename: str
    sheet: str
    rows: list[list[str]]
    encoding: str
    delimiter: str
    warnings: list[str]
    header_row: int


_DIGIT_TRANSLATION = str.maketrans(
    "٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789"
)
_ARABIC_TRANSLATION = str.maketrans("أإآٱىة", "اااايه")
_ARABIC_DIACRITICS = re.compile(r"[\u0610-\u061a\u064b-\u065f\u0670\u06d6-\u06ed\u0640]")
_INVISIBLE_FORMATTING = re.compile(r"[\u061c\u200b-\u200f\u202a-\u202e\u2066-\u2069\ufeff]")
_GENERATED_HEADER = re.compile(r"^(?:عمود\s+\d+|column\s+\d+|unnamed(?::\s*\d+)?)(?:\s*\(\d+\))?$", re.I)
_CSV_READER_LOCK = threading.RLock()
_ZIP_MAX_MEMBERS = 2048
_ZIP_MAX_EXPANDED_BYTES = 128 * 1024 * 1024
_ZIP_MAX_PART_BYTES = 64 * 1024 * 1024
_ZIP_MAX_RATIO = 150
_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_GCC_REGIONS = ("SA", "AE", "QA", "KW", "BH", "OM")
_SUPPORTED_CALLING_CODES = tuple(sorted(
    {str(phonenumbers.country_code_for_region(region)) for region in COUNTRY_OPTIONS},
    key=lambda code: (-len(code), code),
))


def _raw_text(value) -> str:
    """Keep strings byte-for-character; stringify scalar spreadsheet values only."""
    if value is None or value is pd.NA or value is pd.NaT:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (_core_datetime, _core_date, _core_time)):
        return value.isoformat()
    try:
        missing = pd.isna(value)
        if isinstance(missing, (bool, numbers.Integral)) and missing:
            return ""
        # numpy.bool_ is intentionally not assumed to be a Python bool.
        if getattr(missing, "ndim", None) == 0 and bool(missing):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value)


def _digits(text: str) -> str:
    return unicodedata.normalize("NFKC", text).translate(_DIGIT_TRANSLATION)


def _fold_arabic(text: str) -> str:
    return _ARABIC_DIACRITICS.sub("", _digits(text)).translate(_ARABIC_TRANSLATION)


def _normal_words(value) -> str:
    text = _INVISIBLE_FORMATTING.sub("", _fold_arabic(_raw_text(value))).casefold()
    return " ".join(re.sub(r"[^\w]+", " ", text, flags=re.UNICODE).replace("_", " ").split())


def _is_formula(text: str) -> bool:
    value = _INVISIBLE_FORMATTING.sub("", _digits(text)).lstrip()
    return bool(value.startswith(("=", "@")) or re.match(r"^[+-]\s*[A-Za-z_][\w.]*\s*\(", value))


def _check_limits(row_count: int, column_count: int) -> None:
    if row_count > MAX_ROWS:
        raise UserDataError(f"يتجاوز المصدر الحد المسموح وهو {MAX_ROWS:,} صفًا؛ لم تُحذف أو تُختصر أي صفوف.")
    if column_count > MAX_COLUMNS:
        raise UserDataError(f"يتجاوز المصدر الحد المسموح وهو {MAX_COLUMNS} عمودًا؛ قسّم الملف قبل رفعه.")
    if row_count * column_count > MAX_CELLS:
        raise UserDataError(f"يتجاوز حجم الجدول الحد المسموح وهو {MAX_CELLS:,} خلية بعد استكمال الصفوف الناقصة.")


def _validate_input(data: bytes, filename: str) -> str:
    if not isinstance(data, (bytes, bytearray, memoryview)):
        raise UserDataError("محتوى الملف غير صالح؛ أعد رفع الملف الأصلي.")
    if not data:
        raise UserDataError("الملف فارغ ولا يحتوي على بيانات.")
    if len(data) > MAX_FILE_BYTES:
        raise UserDataError("حجم الملف يتجاوز ٢٠ ميجابايت؛ قسّمه إلى ملفات أصغر دون حذف بيانات.")
    suffix = PurePath(str(filename)).suffix.lower()
    if suffix not in {".csv", ".xlsx", ".xls"}:
        raise UserDataError("نوع الملف غير مدعوم؛ استخدم ملف CSV أو XLSX أو XLS.")
    prefix = bytes(data[:8])
    if suffix == ".xlsx" and not prefix.startswith(b"PK\x03\x04"):
        if prefix == _OLE_SIGNATURE:
            raise UserDataError("ملف Excel محمي بكلمة مرور أو لا يطابق صيغة XLSX؛ احفظ نسخة غير محمية بالصيغة الصحيحة.")
        raise UserDataError("توقيع الملف لا يطابق صيغة XLSX؛ قد يكون تالفًا أو أُعيدت تسمية امتداده.")
    if suffix == ".xls" and prefix != _OLE_SIGNATURE:
        raise UserDataError("توقيع الملف لا يطابق صيغة XLS؛ احفظ الملف من Excel بالصيغة الصحيحة.")
    if suffix == ".csv" and (prefix == _OLE_SIGNATURE or prefix.startswith((b"PK\x03\x04", b"%PDF-"))):
        raise UserDataError("المحتوى ليس ملف CSV نصيًا؛ لا يكفي تغيير امتداد الملف.")
    return suffix


def _preflight_xlsx(data: bytes) -> None:
    """Inspect archive sizes and safe XML before openpyxl decompresses anything."""
    try:
        with zipfile.ZipFile(io.BytesIO(data), "r") as archive:
            entries = archive.infolist()
            if len(entries) > _ZIP_MAX_MEMBERS:
                raise UserDataError("يحتوي ملف Excel على عدد غير آمن من الأجزاء المضغوطة.")
            names = [entry.filename for entry in entries]
            if len(names) != len(set(names)):
                raise UserDataError("يحتوي ملف Excel على أجزاء مضغوطة مكررة؛ أعد حفظه من المصدر.")
            expanded = compressed = 0
            for entry in entries:
                path = entry.filename.replace("\\", "/")
                if path.startswith("/") or ".." in path.split("/") or "\x00" in path:
                    raise UserDataError("يحتوي ملف Excel على مسارات داخلية غير آمنة.")
                if entry.flag_bits & 1:
                    raise UserDataError("الأجزاء المضغوطة محمية بكلمة مرور؛ ارفع نسخة غير محمية.")
                if entry.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                    raise UserDataError("يستخدم ملف Excel طريقة ضغط غير مدعومة بأمان.")
                if entry.file_size > _ZIP_MAX_PART_BYTES:
                    raise UserDataError("أحد أجزاء ملف Excel أكبر من الحد الآمن بعد فك الضغط.")
                expanded += entry.file_size
                compressed += entry.compress_size
                if expanded > _ZIP_MAX_EXPANDED_BYTES:
                    raise UserDataError("يتجاوز ملف Excel الحد الآمن للحجم بعد فك الضغط؛ قد يكون ضغطه ضارًا.")
                if entry.file_size >= 1024 * 1024 and entry.file_size / max(entry.compress_size, 1) > _ZIP_MAX_RATIO:
                    raise UserDataError("نسبة ضغط أحد أجزاء ملف Excel غير آمنة؛ ارفع نسخة أعيد حفظها من المصدر.")
            if expanded >= 1024 * 1024 and expanded / max(compressed, 1) > _ZIP_MAX_RATIO:
                raise UserDataError("نسبة ضغط ملف Excel غير آمنة؛ أُوقف الاستيراد لحماية الذاكرة.")
            required = {"[Content_Types].xml", "xl/workbook.xml", "_rels/.rels"}
            if not required.issubset(set(names)):
                raise UserDataError("ملف XLSX غير مكتمل أو ليس مصنف Excel صالحًا.")
            # Defused XML rejects DTD/entity expansion, including small XML bombs.
            for name in ("[Content_Types].xml", "xl/workbook.xml", "_rels/.rels"):
                _safe_xml_fromstring(archive.read(name), forbid_dtd=True, forbid_entities=True, forbid_external=True)
            # Verify every member's CRC with bounded buffers, not a full in-memory read.
            actual_total = 0
            for entry in entries:
                if entry.is_dir():
                    continue
                actual_part = 0
                previous = b""
                with archive.open(entry, "r") as part:
                    while True:
                        chunk = part.read(256 * 1024)
                        if not chunk:
                            break
                        actual_part += len(chunk)
                        actual_total += len(chunk)
                        if actual_part > _ZIP_MAX_PART_BYTES or actual_total > _ZIP_MAX_EXPANDED_BYTES:
                            raise UserDataError("تجاوز المحتوى الفعلي لملف Excel حدود فك الضغط الآمنة.")
                        # openpyxl uses defusedxml; also reject DTDs in all XML parts.
                        # A rolling prefix is kept so declarations split across chunks are seen.
                        if entry.filename.lower().endswith((".xml", ".rels")):
                            probe = previous + chunk
                            if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", probe.replace(b"\x00", b""), re.I):
                                raise UserDataError("يحتوي ملف Excel على تعريفات XML غير آمنة.")
                            previous = chunk[-64:]
                if actual_part != entry.file_size:
                    raise UserDataError("تختلف أحجام أجزاء ملف Excel عن بيانات الأرشيف؛ الملف تالف.")
    except UserDataError:
        raise
    except Exception as exc:
        raise UserDataError("تعذّر التحقق من سلامة ملف XLSX؛ قد يكون تالفًا أو محميًا بكلمة مرور.") from exc


def _open_xlsx(data: bytes):
    _preflight_xlsx(data)
    try:
        return openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False,
        )
    except Exception as exc:
        raise UserDataError("تعذّرت قراءة ملف XLSX؛ أعد حفظ نسخة سليمة وغير محمية من المصنف.") from exc


def _open_xls(data: bytes):
    try:
        return xlrd.open_workbook(file_contents=bytes(data), on_demand=True)
    except Exception as exc:
        raise UserDataError("تعذّرت قراءة ملف XLS؛ قد يكون تالفًا أو مشفّرًا بكلمة مرور.") from exc


def workbook_sheets(data: bytes, filename: str) -> list[str]:
    suffix = _validate_input(data, filename)
    if suffix == ".csv":
        return [""]
    if suffix == ".xlsx":
        book = _open_xlsx(data)
        try:
            names = [sheet.title for sheet in book.worksheets]
        finally:
            book.close()
    else:
        book = _open_xls(data)
        try:
            names = list(book.sheet_names())
        finally:
            book.release_resources()
    if not names:
        raise UserDataError("لا يحتوي المصنف على أوراق بيانات قابلة للقراءة.")
    return names


_FIELD_ALIASES_RAW = {
    "name": (
        "name", "full name", "contact name", "lead name", "customer name", "client name",
        "الاسم", "اسم", "الإسم", "اسم العميل", "اسم العميل بالكامل", "الاسم الكامل",
        "الاسم بالكامل", "الاسم الثلاثي", "اسم الزبون", "اسم الشخص",
    ),
    "phone": (
        "phone", "phone number", "telephone", "telephone number", "mobile", "mobile number",
        "mobile phone", "contact number", "primary phone", "cell", "cell phone", "tel", "p",
        "الهاتف", "هاتف", "رقم الهاتف", "رقم الموبايل", "الموبايل", "موبايل", "الجوال",
        "جوال", "رقم الجوال", "رقم التليفون", "التليفون", "تليفون", "رقم التواصل", "رقم الاتصال",
    ),
    "whatsapp": (
        "whatsapp", "whats app", "whatsapp number", "whatsapp phone", "whatsapp no", "wa",
        "wa number", "watsapp", "واتساب", "واتس", "واتس اب", "واتس آب", "الواتساب", "الواتس",
        "رقم واتساب", "رقم الواتساب", "رقم الواتس", "رقم الواتس اب", "رقم واتس اب",
    ),
    "job": (
        "job", "job title", "occupation", "profession", "position", "work", "career",
        "الوظيفة", "وظيفة", "المهنة", "مهنة", "المسمى الوظيفي", "العمل", "الوظيفة الحالية",
    ),
    "budget": (
        "budget", "expected budget", "budget amount", "purchase budget", "price budget", "max budget",
        "الميزانية", "ميزانية", "الميزانية المتوقعة", "ميزانية الشراء", "ميزانيه", "المبلغ المخصص",
    ),
    "timeline": (
        "timeline", "readiness", "ready", "ready to buy", "buying timeline", "buying timeframe",
        "purchase time", "purchase timeline", "timeframe", "when to buy", "purchase readiness",
        "موعد الشراء", "توقيت الشراء", "ميعاد الشراء", "وقت الشراء", "جاهزية الشراء", "الجاهزية",
        "جاهز للشراء", "هل انت جاهز للشراء", "هل أنت جاهز للشراء", "الوقت المتوقع للشراء",
    ),
    "city": (
        "city", "town", "governorate", "المدينة", "مدينة", "المحافظة", "محافظة", "مدينة السكن",
        "مكان السكن", "مكان الإقامة", "محل الإقامة",
    ),
}
_FIELD_ALIASES = {field: tuple(dict.fromkeys(_normal_words(alias) for alias in aliases))
                  for field, aliases in _FIELD_ALIASES_RAW.items()}


def _header_scores(value) -> dict[str, float]:
    raw = _raw_text(value).strip()
    if not raw or _GENERATED_HEADER.fullmatch(raw) or _is_formula(raw):
        return {}
    # Suffixes added to duplicate headings must not change the original meaning.
    text = _normal_words(re.sub(r"\s+\(\d+\)$", "", raw))
    if not text or len(text) > 100 or re.fullmatch(r"[\d\s]+", text):
        return {}
    scores = {}
    has_wa = bool(re.search(r"(?:whats?\s*app|watsapp|واتس)", text))
    for field, aliases in _FIELD_ALIASES.items():
        if has_wa and field != "whatsapp":
            continue
        if text in aliases:
            scores[field] = 100.0
            continue
        # Short aliases are exact-only; fuzzy matching 'wa', 'p', or 'job' is unsafe.
        candidates = [alias for alias in aliases if len(alias) >= 4 and len(text) >= 4]
        if not candidates:
            continue
        score = max(float(_fuzz.ratio(text, alias)) for alias in candidates)
        tokens = set(text.split())
        for alias in candidates:
            alias_tokens = set(alias.split())
            if len(alias_tokens) >= 2 and alias_tokens <= tokens and len(tokens) <= len(alias_tokens) + 2:
                score = max(score, 93.0)
        if score >= 86:
            scores[field] = score
    return scores


def _recognizable_header_cells(row: list[str], minimum: float = 93.0) -> list[tuple[int, str, float]]:
    hits = []
    for column, value in enumerate(row):
        scores = _header_scores(value)
        if scores:
            field, score = max(scores.items(), key=lambda item: item[1])
            if score >= minimum and len(_normal_words(value)) > 1:
                hits.append((column, field, score))
    return hits


def _looks_data_like(value: str) -> bool:
    text = _digits(_raw_text(value)).strip()
    return bool(text and (re.search(r"\d", text) or "@" in text or _is_formula(text)
                          or _normal_words(text) in {"yes", "no", "now", "later", "نعم", "لا", "فوري", "لاحقا"}))


def detect_header(rows: list[list[str]]) -> tuple[int, float]:
    """Use aliases plus neighboring structure; unknown data is never a header."""
    best_index, best_confidence = -1, 0.0
    for index, row in enumerate(rows[:100]):
        nonempty = [(i, _raw_text(value)) for i, value in enumerate(row) if _raw_text(value).strip()]
        if not nonempty:
            continue
        hits = _recognizable_header_cells(row)
        fields = {field for _, field, _ in hits}
        if not hits:
            continue
        density = len(hits) / len(nonempty)
        data_fraction = sum(_looks_data_like(value) for _, value in nonempty) / len(nonempty)
        # One recognized word among otherwise unknown text is usually metadata, not a heading.
        if len(fields) < 2 and (len(nonempty) > 3 or density < 0.5 or len(hits) != 1):
            continue
        if data_fraction > 0.35:
            continue
        following = [candidate for candidate in rows[index + 1:index + 9]
                     if any(_raw_text(value).strip() for value in candidate)][:4]
        structure = 0.0
        if following:
            header_width = max(i for i, _ in nonempty) + 1
            for candidate in following:
                occupied = [i for i, value in enumerate(candidate) if _raw_text(value).strip()]
                width = max(occupied, default=-1) + 1
                if width and 0.5 <= width / max(header_width, 1) <= 2:
                    structure += 0.08
                contrasts = sum(
                    column < len(candidate) and bool(_raw_text(candidate[column]).strip())
                    and not _header_scores(candidate[column])
                    for column, _, _ in hits
                )
                if contrasts:
                    structure += 0.07
            structure /= len(following)
        elif len(fields) < 2 and density < 1:
            continue
        if len(fields) == 1:
            confidence = 0.63 + 0.12 * density + structure - 0.2 * data_fraction
        else:
            confidence = 0.64 + min(len(fields), 5) * 0.045 + 0.08 * density + structure - 0.2 * data_fraction
        confidence = min(0.99, max(0.0, confidence))
        # Keep the earlier header when a later repeated heading has equivalent evidence.
        if confidence > best_confidence + 0.035:
            best_index, best_confidence = index, confidence
    return best_index, round(best_confidence, 3)


def _strict_decode(data: bytes, encoding: str) -> str:
    try:
        text = bytes(data).decode(encoding, errors="strict")
    except (UnicodeError, LookupError) as exc:
        raise UserDataError("لا يمكن فك ترميز الملف دون فقد بيانات؛ اختر الترميز الصحيح يدويًا أو احفظ نسخة UTF-8.") from exc
    if text.startswith("\ufeff"):
        text = text[1:]
    if "\x00" in text or re.search(r"[\x01-\x08\x0b\x0c\x0e-\x1f]", text):
        raise UserDataError("تحتوي القراءة النصية على محارف ثنائية غير صالحة؛ راجع الترميز أو نوع الملف.")
    return text


def _decode_csv(data: bytes, encoding: str) -> tuple[str, str, list[str]]:
    manual = str(encoding or "auto").strip().lower()
    if manual != "auto":
        try:
            canonical = codecs.lookup(manual).name
        except LookupError as exc:
            raise UserDataError("الترميز المحدد غير معروف؛ اختر ترميزًا من القائمة.") from exc
        allowed = {"utf-8", "utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "utf-32", "utf-32-le",
                   "utf-32-be", "cp1256", "cp1252", "iso8859-1", "iso8859-6", "cp864"}
        if canonical not in allowed:
            raise UserDataError("الترميز المحدد غير مدعوم بأمان؛ استخدم UTF-8 أو أحد ترميزات القائمة.")
        return _strict_decode(data, canonical), canonical, []
    for bom, codec in ((codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32"),
                       (codecs.BOM_UTF8, "utf-8-sig"), (codecs.BOM_UTF16_LE, "utf-16"),
                       (codecs.BOM_UTF16_BE, "utf-16")):
        if bytes(data).startswith(bom):
            return _strict_decode(data, codec), codec, []
    try:
        utf8 = bytes(data).decode("utf-8", errors="strict")
        if "\x00" not in utf8:
            return _strict_decode(data, "utf-8"), "utf-8", []
    except UnicodeDecodeError:
        pass
    predicted = None
    try:
        prediction = _charset_from_bytes(bytes(data[:512 * 1024])).best()
        if prediction is not None:
            predicted = codecs.lookup(prediction.encoding).name
    except (LookupError, UnicodeError, ValueError):
        pass
    if b"\x00" in bytes(data) and predicted in {"utf-16-le", "utf-16-be", "utf-32-le", "utf-32-be"}:
        text = _strict_decode(data, predicted)
        return text, predicted, ["اختير ترميز Unicode آليًا دون علامة تعريف؛ راجع النص وحدّد الترميز يدويًا إذا لزم."]
    candidates = {}
    for codec in ("cp1256", "cp1252"):
        try:
            candidates[codec] = _strict_decode(data, codec)
        except UserDataError:
            continue
    if not candidates:
        raise UserDataError("تعذّر تحديد ترميز نصي آمن؛ اختر الترميز الأصلي يدويًا دون استبدال المحارف.")
    if len(candidates) == 1:
        selected = next(iter(candidates))
    else:
        arabic = candidates["cp1256"]
        runs = re.findall(r"[\u0621-\u064a]{3,}", arabic[:30000])
        # Real Arabic words are a stronger signal than isolated accented Latin bytes.
        arabic_letters = sum(len(run) for run in runs)
        nonspace = sum(not char.isspace() for char in arabic[:30000])
        if len(runs) >= 2 and arabic_letters / max(nonspace, 1) >= 0.08:
            selected = "cp1256"
        elif predicted in candidates:
            selected = predicted
        else:
            selected = "cp1252"
    warning = (f"اختير الترميز {selected} آليًا، وقد يلتبس مع ترميز آخر؛ راجع الأسماء العربية والمحارف "
               "واختر الترميز يدويًا عند الحاجة. لم تُستبدل أو تُحذف أي بايتات أثناء فك الترميز.")
    return candidates[selected], selected, [warning]


def _csv_records(text: str, delimiter: str, preview: bool = False) -> tuple[list[list[str]], bool]:
    rows = []
    failed = False
    width = 0
    # csv.field_size_limit is process-global; protect and restore the temporary change.
    # This permits long raw cells without imposing a presentation-length limit.
    with _CSV_READER_LOCK:
        previous_limit = csv.field_size_limit()
        csv.field_size_limit(max(previous_limit, MAX_FILE_BYTES))
        try:
            reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            for row in reader:
                rows.append(row)
                width = max(width, len(row))
                if not preview:
                    _check_limits(len(rows), width)
                if preview and len(rows) >= 120:
                    break
        except csv.Error as exc:
            if not preview:
                raise UserDataError("بنية ملف CSV غير سليمة؛ راجع علامات الاقتباس والفاصل. لم يُتجاوز أي صف تالف.") from exc
            failed = True
        finally:
            csv.field_size_limit(previous_limit)
    return rows, failed


def _choose_delimiter(text: str, requested: str) -> tuple[str, list[str]]:
    selected = str(requested or "auto")
    aliases = {"tab": "\t", "\\t": "\t", "comma": ",", "semicolon": ";", "pipe": "|"}
    if selected != "auto":
        selected = aliases.get(selected.lower(), selected)
        if len(selected) != 1 or selected in "\r\n\x00\"" or selected.isalnum():
            raise UserDataError("الفاصل المحدد غير صالح؛ اختر الفاصلة أو الفاصلة المنقوطة أو الجدولة أو الخط العمودي.")
        return selected, []
    lines = []
    stream = io.StringIO(text)
    for _ in range(120):
        line = stream.readline()
        if not line:
            break
        lines.append(line)
    for line in lines[:5]:
        directive = re.fullmatch(r"\s*sep\s*=\s*([,;\t|])\s*", line, flags=re.I)
        if directive:
            return directive.group(1), ["استُخدم الفاصل المعلن في سطر تعريف الفاصل، مع الاحتفاظ بالسطر ضمن بيانات المصدر."]
        if line.strip():
            break
    sniffed = Counter()
    starts = {0, 1, 2, 3, 5, 10, 20, 50, 99}
    # Metadata can have a completely different delimiter from the actual table.
    for index, line in enumerate(lines[:100]):
        pieces = re.split(r"[,;\t|]", line.strip())
        if _recognizable_header_cells(pieces):
            starts.add(index)
    for start in sorted(starts):
        if start >= len(lines):
            continue
        sample = "".join(lines[start:start + 8])[:65536]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            sniffed[dialect.delimiter] += 1
        except csv.Error:
            continue
    scored = []
    for order, delimiter in enumerate((",", ";", "\t", "|")):
        rows, failed = _csv_records(text, delimiter, preview=True)
        populated = [row for row in rows if any(cell.strip() for cell in row)]
        widths = [len(row) for row in populated]
        if not widths:
            score = -100.0
        else:
            modal_width, count = Counter(widths).most_common(1)[0]
            multicol = sum(width > 1 for width in widths) / len(widths)
            alias_evidence = max((len({field for _, field, _ in _recognizable_header_cells(row)})
                                  for row in populated[:100]), default=0)
            score = alias_evidence * 35 + multicol * 25 + count / len(widths) * 15
            score += 8 if modal_width > 1 else -8
            score += min(sniffed[delimiter], 4) * 2
            if max(widths) > MAX_COLUMNS:
                score -= 12
            # Malformation must not make a wrong one-column delimiter look safer.
            if failed:
                score -= 1
        scored.append((score, -order, delimiter, max(widths, default=0)))
    best = max(scored)
    warnings = []
    if best[3] <= 1:
        warnings.append("لم يُرصد فاصل أعمدة واضح؛ قُرئ المصدر كعمود واحد. يمكنك تحديد الفاصل يدويًا.")
    return best[2], warnings


def _pad_rows(rows: list[list[str]]) -> tuple[list[list[str]], bool]:
    width = max((len(row) for row in rows), default=0)
    _check_limits(len(rows), width)
    lengths = {len(row) for row in rows if any(_raw_text(value).strip() for value in row)}
    return [[_raw_text(value) for value in row] + [""] * (width - len(row)) for row in rows], len(lengths) > 1


def _select_sheet(names: list[str], sheet: str) -> str:
    if not names:
        raise UserDataError("لا يحتوي المصنف على أوراق بيانات قابلة للقراءة.")
    if sheet == "" or sheet is None:
        return names[0]
    if sheet not in names:
        raise UserDataError("ورقة العمل المحددة غير موجودة في هذا الملف؛ أعد اختيارها من القائمة.")
    return sheet


def _read_xlsx_rows(data: bytes, sheet: str) -> tuple[list[list[str]], str, list[str]]:
    book = _open_xlsx(data)
    try:
        selected = _select_sheet([worksheet.title for worksheet in book.worksheets], sheet)
        worksheet = book[selected]
        # Check advertised dimensions, then ignore them so under-reported dimensions
        # cannot silently hide rows or columns from read-only iteration.
        _check_limits(worksheet.max_row or 0, worksheet.max_column or 0)
        worksheet.reset_dimensions()
        rows, width = [], 0
        for row in worksheet.iter_rows(values_only=True):
            width = max(width, len(row))
            _check_limits(len(rows) + 1, width)
            rows.append([_raw_text(value) for value in row])
        padded, inconsistent = _pad_rows(rows)
        warnings = []
        if inconsistent:
            warnings.append("استُكملت الخانات الناقصة بقيم فارغة دون إسقاط بيانات.")
        if any(_is_formula(value) for row in padded for value in row if value):
            warnings.append("احتُفظ بنصوص الصيغ كما هي ولم تُنفّذ؛ لا تُعد الصيغ أرقام تواصل أو ميزانيات صالحة.")
        return padded, selected, warnings
    except UserDataError:
        raise
    except Exception as exc:
        raise UserDataError("تعذّرت قراءة بيانات الورقة بالكامل؛ قد تكون تالفة. لم تُستورد نسخة مبتورة.") from exc
    finally:
        book.close()


def _preflight_xls_sheet(book, sheet_number: int) -> None:
    """Check BIFF dimensions and reject cached-only formulas before allocation.

    xlrd intentionally exposes formula results, not the formula text. Importing
    those cached results as contact data would bypass formula rejection. Legacy
    formula-bearing sheets therefore fail closed with a conversion instruction.
    """
    memory = book.mem
    positions = book._sh_abs_posn
    if memory is None or sheet_number >= len(positions):
        raise UserDataError("تعذّر فحص بنية ملف XLS بأمان؛ احفظ نسخة بصيغة XLSX.")
    position = positions[sheet_number]
    nesting = 0
    maximum_row = maximum_column = 0
    cell_opcodes = {0x0001, 0x0002, 0x0003, 0x0004, 0x0005, 0x0201, 0x0203, 0x0204,
                    0x0205, 0x027E, 0x00FD, 0x00D6, 0x00BD, 0x00BE}
    while position + 4 <= len(memory):
        opcode, length = struct.unpack_from("<HH", memory, position)
        position += 4
        end = position + length
        if end > len(memory):
            raise UserDataError("أحد سجلات ملف XLS مبتور؛ أعد حفظ نسخة سليمة.")
        if opcode in {0x0009, 0x0209, 0x0409, 0x0809}:
            nesting += 1
        elif opcode == 0x000A:
            nesting -= 1
            if nesting <= 0:
                return
        elif nesting == 1:
            if opcode == 0x002F:
                raise UserDataError("ملف XLS محمي بكلمة مرور؛ ارفع نسخة غير محمية.")
            if opcode in {0x0006, 0x0206, 0x0406}:
                raise UserDataError("تحتوي ورقة XLS على صيغ لا تتيح هذه الصيغة القديمة استرجاع نصوصها بأمان؛ احفظها بصيغة XLSX. لم تُستخدم النتائج المخبأة كأرقام أو ميزانيات.")
            if opcode in {0x0200, 0x0000}:
                if book.biff_version >= 80 and length >= 12:
                    _, last_row, _, last_column = struct.unpack_from("<IIHH", memory, position)
                elif length >= 8:
                    _, last_row, _, last_column = struct.unpack_from("<HHHH", memory, position)
                else:
                    raise UserDataError("أبعاد ورقة XLS غير سليمة.")
                _check_limits(last_row, last_column)
            if opcode in cell_opcodes:
                if length < 4:
                    raise UserDataError("سجل خلية XLS غير مكتمل.")
                row, column = struct.unpack_from("<HH", memory, position)
                if opcode in {0x00BD, 0x00BE}:
                    if length < 6:
                        raise UserDataError("سجل خلايا XLS غير مكتمل.")
                    column = max(column, struct.unpack_from("<H", memory, end - 2)[0])
                maximum_row = max(maximum_row, row + 1)
                maximum_column = max(maximum_column, column + 1)
                _check_limits(maximum_row, maximum_column)
        position = end
    raise UserDataError("لا تحتوي ورقة XLS على سجل نهاية سليم؛ قد يكون الملف تالفًا.")


def _read_xls_rows(data: bytes, sheet: str) -> tuple[list[list[str]], str, list[str]]:
    book = _open_xls(data)
    try:
        selected = _select_sheet(list(book.sheet_names()), sheet)
        _preflight_xls_sheet(book, list(book.sheet_names()).index(selected))
        worksheet = book.sheet_by_name(selected)
        _check_limits(worksheet.nrows, worksheet.ncols)
        rows = []
        for row_index in range(worksheet.nrows):
            values = []
            for column_index in range(worksheet.ncols):
                cell = worksheet.cell(row_index, column_index)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    value = ""
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat()
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    value = xlrd.error_text_from_code.get(cell.value, "خطأ في خلية المصدر")
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = "TRUE" if cell.value else "FALSE"
                else:
                    value = _raw_text(cell.value)
                values.append(value)
            rows.append(values)
        return rows, selected, ["قُرئ ملف XLS دون تشغيل صيغ أو روابط خارجية؛ يلزم استخدام XLSX للأوراق المحتوية على صيغ حتى لا تُستعمل نتائج مخبأة."]
    except UserDataError:
        raise
    except Exception as exc:
        raise UserDataError("تعذّرت قراءة بيانات ملف XLS بالكامل؛ لم تُستورد نسخة مبتورة.") from exc
    finally:
        book.release_resources()


def read_source(data: bytes, filename: str, sheet: str = "", encoding: str = "auto", delimiter: str = "auto") -> ParsedTable:
    suffix = _validate_input(data, filename)
    warnings = []
    selected_sheet = ""
    used_encoding = used_delimiter = ""
    if suffix == ".csv":
        text, used_encoding, notes = _decode_csv(data, encoding)
        warnings.extend(notes)
        used_delimiter, notes = _choose_delimiter(text, delimiter)
        warnings.extend(notes)
        raw_rows, _ = _csv_records(text, used_delimiter)
        rows, inconsistent = _pad_rows(raw_rows)
        if inconsistent:
            warnings.append("تختلف أعداد الأعمدة بين بعض الصفوف؛ أُكملت الخانات الناقصة بقيم فارغة دون حذف بيانات.")
    elif suffix == ".xlsx":
        rows, selected_sheet, notes = _read_xlsx_rows(data, sheet)
        warnings.extend(notes)
    else:
        rows, selected_sheet, notes = _read_xls_rows(data, sheet)
        warnings.extend(notes)
    header_row, confidence = detect_header(rows)
    if header_row == -1:
        warnings.append("لم تُكتشف عناوين معروفة بثقة؛ لن يُحذف أول صف بيانات، وستُستخدم أسماء أعمدة مولّدة حتى تختار صف العناوين.")
    else:
        if header_row:
            warnings.append(f"اكتُشف صف العناوين في الصف {header_row + 1}؛ احتُفظ بالصفوف السابقة في المصدر الأصلي.")
        if confidence < 0.8:
            warnings.append("الثقة في صف العناوين محدودة؛ راجع المعاينة أو اختر صف العناوين يدويًا.")
        signature = [_normal_words(value) for value in rows[header_row]]
        repeated = [index + 1 for index, row in enumerate(rows[header_row + 1:], header_row + 1)
                    if [_normal_words(value) for value in row] == signature and any(value.strip() for value in row)]
        if repeated:
            warnings.append(f"تكرّر صف العناوين داخل البيانات ({len(repeated)} مرة)؛ احتُفظ بهذه الصفوف ولم تُحذف تلقائيًا.")
    digest = hashlib.sha256()
    digest.update(str(filename).encode("utf-8", errors="surrogatepass"))
    digest.update(b"\x00")
    digest.update(selected_sheet.encode("utf-8", errors="surrogatepass"))
    digest.update(b"\x00")
    digest.update(bytes(data))
    return ParsedTable(digest.hexdigest(), str(filename), selected_sheet, rows,
                       used_encoding, used_delimiter, list(dict.fromkeys(warnings)), header_row)


def _unique_columns(values: list[str]) -> tuple[list[str], list[str]]:
    names, generated, occupied = [], [], set()
    # Reserve genuine headings, so a blank column cannot steal a real heading.
    reserved = {_raw_text(value).strip() for value in values if _raw_text(value).strip()}
    for position, value in enumerate(values, 1):
        base = _raw_text(value).strip()
        was_generated = not base
        if was_generated:
            base = f"عمود {position}"
        name, suffix = base, 2
        while name in occupied or (was_generated and name in reserved):
            name = f"{base} ({suffix})"
            suffix += 1
        # A duplicate must not consume the literal heading 'name (2)' later on.
        if name != base:
            while name in reserved or name in occupied:
                name = f"{base} ({suffix})"
                suffix += 1
        occupied.add(name)
        names.append(name)
        if was_generated:
            generated.append(name)
    return names, generated


def table_from_source(source: ParsedTable, header_row: int) -> pd.DataFrame:
    if not isinstance(source, ParsedTable):
        raise UserDataError("مصدر الجدول غير صالح؛ أعد قراءة الملف.")
    if isinstance(header_row, bool) or not isinstance(header_row, numbers.Integral):
        raise UserDataError("رقم صف العناوين غير صالح.")
    header_row = int(header_row)
    if header_row < -1 or header_row >= len(source.rows):
        raise UserDataError("صف العناوين المحدد خارج حدود المصدر.")
    width = max((len(row) for row in source.rows), default=0)
    _check_limits(len(source.rows), width)
    raw_headers = (list(source.rows[header_row]) + [""] * (width - len(source.rows[header_row]))) if header_row >= 0 else [""] * width
    columns, generated = _unique_columns(raw_headers)
    first_data = header_row + 1 if header_row >= 0 else 0
    data, source_rows = [], []
    warnings = list(source.warnings)
    repeated_count = 0
    header_signature = [_normal_words(value) for value in raw_headers]
    for row_number in range(first_data, len(source.rows)):
        row = [_raw_text(value) for value in source.rows[row_number]]
        if not any(value.strip() for value in row):
            continue
        row += [""] * (width - len(row))
        data.append(row)
        source_rows.append(row_number + 1)
        if header_row >= 0 and [_normal_words(value) for value in row] == header_signature:
            repeated_count += 1
    if repeated_count and not any("تكرّر صف العناوين" in warning for warning in warnings):
        warnings.append(f"تكرّر صف العناوين داخل البيانات ({repeated_count} مرة)؛ احتُفظ بهذه الصفوف ولم تُحذف تلقائيًا.")
    frame = pd.DataFrame(data, columns=columns, dtype=object)
    frame.attrs = {
        "source_id": source.source_id, "filename": source.filename, "sheet": source.sheet,
        "encoding": source.encoding, "delimiter": source.delimiter, "header_row": header_row,
        "source_rows": source_rows, "generated_columns": generated, "original_headers": raw_headers,
        "metadata_rows": copy.deepcopy(source.rows[:header_row]) if header_row >= 0 else [],
        "warnings": list(dict.fromkeys(warnings)),
    }
    return frame


def guess_mapping(frame: pd.DataFrame) -> dict[str, str | None]:
    fields = tuple(FIELD_LABELS)
    generated = set(frame.attrs.get("generated_columns", []))
    candidates = []
    for column in frame.columns:
        if not isinstance(column, str) or column in generated or _GENERATED_HEADER.fullmatch(column.strip()):
            continue
        scores = _header_scores(column)
        if scores:
            candidates.append((column, scores))
    # A 2**7 dynamic program finds the best one-to-one assignment, not a greedy
    # match that can consume the only viable column for another field.
    states = {0: (0.0, {})}
    for column, scores in candidates:
        updated = dict(states)
        for mask, (score, assignment) in states.items():
            for index, field in enumerate(fields):
                if mask & (1 << index) or field not in scores:
                    continue
                next_mask = mask | (1 << index)
                total = score + scores[field]
                if next_mask not in updated or total > updated[next_mask][0]:
                    updated[next_mask] = (total, {**assignment, field: column})
        states = updated
    assignment = max(states.values(), key=lambda state: state[0])[1]
    return {field: assignment.get(field) for field in fields}


def _phone_failure(raw: str, message: str) -> dict:
    return {"raw": raw, "e164": "", "display": "", "valid": False, "region": "",
            "error": message, "tel_url": "", "wa_url": ""}


def _parsed_valid_phone(text: str, region: str | None):
    try:
        number = phonenumbers.parse(text, region, keep_raw_input=False)
        if number.extension or not phonenumbers.is_valid_number(number):
            return None
        return number
    except (phonenumbers.NumberParseException, TypeError, ValueError):
        return None


def _phone_success(raw: str, number, notes: list[str]) -> dict:
    e164 = phonenumbers.format_number(number, phonenumbers.PhoneNumberFormat.E164)
    return {
        "raw": raw, "e164": e164,
        "display": phonenumbers.format_number(number, phonenumbers.PhoneNumberFormat.INTERNATIONAL),
        "valid": True, "region": phonenumbers.region_code_for_number(number) or "",
        "error": "؛ ".join(dict.fromkeys(notes)), "tel_url": f"tel:{e164}",
        "wa_url": f"https://wa.me/{e164[1:]}",
    }


def _phone_url_number(text: str) -> tuple[str, bool]:
    candidate = text
    if re.match(r"^(?:www\.)?wa\.me/", candidate, re.I):
        candidate = "https://" + candidate
    if not re.match(r"^[a-z][a-z0-9+.-]*://", candidate, re.I):
        return text, False
    try:
        parsed = urlsplit(candidate)
        if parsed.scheme.lower() not in {"https", "http", "whatsapp"} or parsed.username or parsed.password or parsed.port:
            raise ValueError
        host = (parsed.hostname or "").lower()
        if host in {"wa.me", "www.wa.me"}:
            if parsed.scheme.lower() not in {"http", "https"}:
                raise ValueError
            path = unquote(parsed.path).strip("/")
            if "/" in path or not path or parsed.fragment:
                raise ValueError
            return path, True
        if ((host in {"api.whatsapp.com", "web.whatsapp.com"} and parsed.path.rstrip("/") == "/send")
                or (parsed.scheme.lower() == "whatsapp" and host == "send")):
            query = parse_qs(parsed.query, keep_blank_values=True)
            phones = query.get("phone", [])
            if len(phones) != 1 or not phones[0].strip() or parsed.fragment:
                raise ValueError
            return phones[0].strip(), True
    except (ValueError, UnicodeError):
        pass
    raise UserDataError("رابط التواصل غير صالح أو لا يحتوي على رقم واتساب واحد واضح.")


def normalize_phone(value, default_region: str = "EG") -> dict:
    """Validate the numbering plan, not carrier activity or WhatsApp ownership.

    A nonempty error on a valid result is an explicit normalization/inference
    note. Only ``valid`` grants permission to generate a contact link.
    """
    raw = _raw_text(value)
    text = _INVISIBLE_FORMATTING.sub("", _digits(raw)).strip()
    if not text:
        return _phone_failure(raw, "لم يُدخل رقم.")
    if _is_formula(text):
        return _phone_failure(raw, "الخانة تحتوي على صيغة وليست رقم تواصل؛ لم تُنفّذ الصيغة.")
    text = re.sub(r"^(?:(?:p|tel|telephone|phone|mobile|هاتف|الهاتف|موبايل|جوال)\s*:\s*)+", "", text, flags=re.I).strip()
    if _is_formula(text):
        return _phone_failure(raw, "الخانة تحتوي على صيغة وليست رقم تواصل؛ لم تُنفّذ الصيغة.")
    notes = []
    if text.startswith("'") and re.match(r"^'[+\d]", text):
        text = text[1:]
        notes.append("أُزيلت علامة حفظ الرقم كنص، مع الاحتفاظ بالقيمة الأصلية.")
    try:
        text, was_url = _phone_url_number(text)
    except UserDataError as exc:
        return _phone_failure(raw, str(exc))
    text = _INVISIBLE_FORMATTING.sub("", _digits(text)).strip().replace("٫", ".")
    if _is_formula(text):
        return _phone_failure(raw, "الخانة تحتوي على صيغة وليست رقم تواصل؛ لم تُنفّذ الصيغة.")
    if not text:
        return _phone_failure(raw, "لا يوجد رقم بعد بادئة التواصل.")
    if text.startswith(("-", "−", "–", "—")):
        return _phone_failure(raw, "رقم التواصل لا يمكن أن يكون قيمة سالبة.")
    # Convert spreadsheet artifacts with Decimal, never with a lossy float cast.
    numeric_artifact = bool(re.fullmatch(r"\+?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text))
    if numeric_artifact and ("." in text or "e" in text.lower()):
        try:
            number = Decimal(text)
            if not number.is_finite() or number <= 0 or number != number.to_integral_value():
                return _phone_failure(raw, "القيمة العشرية أو الكسرية ليست رقم تواصل صحيحًا؛ لم تُقرّب أو تُحذف أرقام.")
            if number.adjusted() > 16:
                return _phone_failure(raw, "عدد أرقام التواصل يتجاوز الحد الدولي المسموح.")
            if "e" not in text.lower():
                # Preserve meaningful leading zeros in an integral '.0' artifact.
                text = text.split(".", 1)[0]
            else:
                leading_plus = text.startswith("+")
                text = ("+" if leading_plus else "") + format(number, "f").split(".", 1)[0]
            notes.append("حُوّلت الكتابة العشرية أو العلمية إلى عدد صحيح دون تقريب؛ راجع الأصفار التي ربما فُقدت في المصدر.")
        except (InvalidOperation, ValueError, OverflowError):
            return _phone_failure(raw, "تعذّر تفسير الصيغة الرقمية بأمان؛ استخدم الرقم الأصلي كنص.")
    if re.search(r"[,،;؛|/\\\r\n]", text) or text.count("+") > 1:
        return _phone_failure(raw, "تحتوي الخانة على عدة أرقام أو فواصل غير آمنة؛ ضع رقمًا واحدًا في كل خانة.")
    if len(re.findall(r"\d{7,}", text)) > 1:
        return _phone_failure(raw, "يبدو أن الخانة تحتوي على أكثر من رقم؛ لم تُدمج الأرقام معًا.")
    if "." in text:
        return _phone_failure(raw, "النقاط أو الكسور داخل الرقم ملتبسة؛ استخدم رقمًا صحيحًا واحدًا دون تغيير أرقامه.")
    if not re.fullmatch(r"[+0-9()\s\-]+", text):
        return _phone_failure(raw, "يحتوي الرقم على حروف أو رموز أو امتداد غير مدعوم؛ لم تُحذف الأجزاء غير المفهومة.")
    depth = 0
    for character in text:
        if character == "(":
            depth += 1
        elif character == ")":
            depth -= 1
        if depth < 0:
            return _phone_failure(raw, "تنسيق الأقواس في الرقم غير سليم.")
    if depth:
        return _phone_failure(raw, "تنسيق الأقواس في الرقم غير سليم.")
    compact = re.sub(r"[\s()\-]", "", text)
    if "+" in compact[1:] or not re.fullmatch(r"\+?\d+", compact):
        return _phone_failure(raw, "صيغة رمز الدولة غير سليمة؛ ضع علامة الجمع في بداية رقم واحد فقط.")
    explicit = compact.startswith(("+", "00")) or was_url
    if compact.startswith("00"):
        compact = "+" + compact[2:]
    elif was_url and not compact.startswith("+"):
        compact = "+" + compact
    if len(compact.lstrip("+")) > 15:
        return _phone_failure(raw, "عدد أرقام التواصل يتجاوز الحد الدولي المسموح؛ ربما جُمعت عدة أرقام في خانة واحدة.")
    if explicit:
        parsed = _parsed_valid_phone(compact, None)
        if parsed is None:
            return _phone_failure(raw, "الرقم غير صالح حسب خطة الترقيم الدولية؛ لم يُستبدل رمز الدولة الصريح بتخمين محلي.")
        return _phone_success(raw, parsed, notes)
    inferred = {}
    for calling_code in _SUPPORTED_CALLING_CODES:
        if compact.startswith(calling_code):
            parsed = _parsed_valid_phone("+" + compact, None)
            if parsed is not None:
                key = phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164)
                inferred[key] = parsed
    if len(inferred) == 1:
        parsed = next(iter(inferred.values()))
        notes.append("اعتُبر الرقم دوليًا دون علامة + لأن رمز الدولة معروف؛ أكّد رمز الدولة قبل التواصل.")
        return _phone_success(raw, parsed, notes)
    if len(inferred) > 1:
        return _phone_failure(raw, "الرقم الدولي دون بادئة ملتبس؛ أضف رمز الدولة مسبوقًا بعلامة +.")
    region = str(default_region or "").strip().upper()
    if region not in phonenumbers.SUPPORTED_REGIONS:
        return _phone_failure(raw, "الدولة الافتراضية غير صالحة؛ حدّد الدولة أو اكتب رمزها الدولي صراحة.")
    parsed = _parsed_valid_phone(compact, region)
    if parsed is not None:
        label = COUNTRY_OPTIONS.get(region, region)
        notes.append(f"فُسّر الرقم المحلي وفق الدولة الافتراضية: {label}؛ غيّر الدولة إذا لم تكن صحيحة.")
        return _phone_success(raw, parsed, notes)
    gcc_matches = {}
    for gcc_region in _GCC_REGIONS:
        if gcc_region == region:
            continue
        parsed = _parsed_valid_phone(compact, gcc_region)
        if parsed is not None:
            key = phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164)
            gcc_matches[key] = parsed
    if len(gcc_matches) > 1:
        return _phone_failure(raw, "الرقم المحلي يطابق أكثر من دولة خليجية؛ لن نخمّن الدولة. أضف رمز الدولة الدولي.")
    if len(gcc_matches) == 1:
        parsed = next(iter(gcc_matches.values()))
        detected = phonenumbers.region_code_for_number(parsed) or ""
        label = COUNTRY_OPTIONS.get(detected, detected)
        notes.append(f"الرقم غير صالح للدولة الافتراضية؛ رُجّحت {label} بوصفها المطابقة الخليجية الوحيدة. أكّد رمز الدولة.")
        return _phone_success(raw, parsed, notes)
    return _phone_failure(raw, "الرقم غير صالح للدولة المحددة ولا يمكن استنتاج دولة واحدة بأمان؛ راجع الأرقام ورمز الدولة.")


_CURRENCY_ALIASES_RAW = {
    "EGP": ("EGP", "جنيه مصري", "جنيهات مصرية", "جنيه مصرى", "ج.م", "ج م", "جم", "E£", "EG£", "L.E.", "LE"),
    "SAR": ("SAR", "ريال سعودي", "ريال سعودى", "ريالات سعودية", "ر.س", "ر س", "رس"),
    "AED": ("AED", "درهم إماراتي", "درهم اماراتى", "دراهم إماراتية", "د.إ", "د إ", "د ا", "د.ا"),
    "QAR": ("QAR", "ريال قطري", "ريال قطرى", "ريالات قطرية", "ر.ق", "ر ق"),
    "KWD": ("KWD", "دينار كويتي", "دينار كويتى", "دنانير كويتية", "د.ك", "د ك"),
    "BHD": ("BHD", "دينار بحريني", "دينار بحرينى", "دنانير بحرينية", "د.ب", "د ب"),
    "OMR": ("OMR", "ريال عماني", "ريال عمانى", "ريالات عمانية", "ر.ع", "ر ع"),
    "USD": ("USD", "US$", "$", "دولار أمريكي", "دولار امريكى", "دولار اميركي", "دولار", "دولارات", "dollars", "dollar"),
    "EUR": ("EUR", "€", "يورو", "euro", "euros"),
}
_CURRENCY_ALIAS_LOOKUP = {_fold_arabic(alias).casefold(): currency
                          for currency, aliases in _CURRENCY_ALIASES_RAW.items() for alias in aliases}
_CURRENCY_TOKEN_PATTERN = re.compile(
    r"(?<![a-z\u0621-\u064a])(?:" + "|".join(
        re.escape(alias).replace(r"\ ", r"\s+")
        for alias in sorted(_CURRENCY_ALIAS_LOOKUP, key=len, reverse=True)
    ) + r")(?![a-z\u0621-\u064a])", re.I,
)
_GENERIC_CURRENCY_PATTERNS = (
    (re.compile(r"(?<![a-z\u0621-\u064a])(?:جنيهات|جنيه)(?![a-z\u0621-\u064a])"), {"EGP"}),
    (re.compile(r"(?<![a-z\u0621-\u064a])(?:ريالات|ريال)(?![a-z\u0621-\u064a])"), {"SAR", "QAR", "OMR"}),
    (re.compile(r"(?<![a-z\u0621-\u064a])(?:دراهم|درهم)(?![a-z\u0621-\u064a])"), {"AED"}),
    (re.compile(r"(?<![a-z\u0621-\u064a])(?:دنانير|دينار)(?![a-z\u0621-\u064a])"), {"KWD", "BHD"}),
)
_SCALE_LOOKUP = {"k": Decimal(1000), "thousand": Decimal(1000), "thousands": Decimal(1000),
                 "الف": Decimal(1000), "الاف": Decimal(1000), "الوف": Decimal(1000),
                 "m": Decimal(1000000), "million": Decimal(1000000), "millions": Decimal(1000000),
                 "مليون": Decimal(1000000), "ملايين": Decimal(1000000)}
_SCALE_END = re.compile(r"\s*(" + "|".join(sorted(_SCALE_LOOKUP, key=len, reverse=True)) + r")\s*$", re.I)


def _budget_currency(text: str, default_currency: str) -> tuple[str, str, bool]:
    matches = []
    def remove_currency(match):
        alias = " ".join(match.group(0).split())
        matches.append(_CURRENCY_ALIAS_LOOKUP[alias])
        return " "
    # Removing a currency between two values must not turn '100 USD 200'
    # into the apparently well-grouped number '100 200'.
    for match in _CURRENCY_TOKEN_PATTERN.finditer(text):
        before, after = text[:match.start()].rstrip(), text[match.end():].lstrip()
        if before and after and before[-1].isdigit() and after[0].isdigit():
            return text, "", True
    remaining = _CURRENCY_TOKEN_PATTERN.sub(remove_currency, text)
    currencies = set(matches)
    if len(currencies) > 1:
        return remaining, "", True
    selected = next(iter(currencies), default_currency)
    for pattern, possibilities in _GENERIC_CURRENCY_PATTERNS:
        if pattern.search(remaining):
            if selected not in possibilities:
                return remaining, "", True
            for match in pattern.finditer(remaining):
                before, after = remaining[:match.start()].rstrip(), remaining[match.end():].lstrip()
                if before and after and before[-1].isdigit() and after[0].isdigit():
                    return remaining, "", True
            remaining = pattern.sub(" ", remaining)
    return remaining, selected, selected not in CURRENCY_OPTIONS


def _budget_decimal(text: str) -> Decimal | None:
    value = text.strip()
    if not value:
        return None
    if re.fullmatch(r"\+?\d+(?:\.\d+)?[eE][+-]?\d+", value):
        try:
            parsed = Decimal(value)
            return parsed if parsed.is_finite() and parsed >= 0 and parsed.adjusted() <= 308 else None
        except InvalidOperation:
            return None
    value = value.lstrip("+") if value.startswith("+") and not value.startswith("++") else value
    # Whitespace grouping is valid only in complete groups of three digits.
    if re.search(r"\s", value):
        if not re.fullmatch(r"\d{1,3}(?:\s\d{3})+(?:[.,٫]\d{1,2})?", value):
            return None
        value = re.sub(r"\s", "", value)
    if "٬" in value:
        if not re.fullmatch(r"\d{1,3}(?:٬\d{3})+(?:[.٫]\d+)?", value):
            return None
        value = value.replace("٬", "")
    if "٫" in value:
        if "." in value or "," in value or value.count("٫") != 1:
            return None
        value = value.replace("٫", ".")
        explicit_decimal = True
    else:
        explicit_decimal = False
    if "," in value and "." in value:
        if re.fullmatch(r"\d{1,3}(?:,\d{3})+\.\d{1,2}", value):
            value = value.replace(",", "")
        elif re.fullmatch(r"\d{1,3}(?:\.\d{3})+,\d{1,2}", value):
            value = value.replace(".", "").replace(",", ".")
        else:
            return None
    elif "," in value:
        if re.fullmatch(r"\d{1,3}(?:,\d{3})+", value):
            value = value.replace(",", "")
        elif re.fullmatch(r"\d+,\d{1,2}", value):
            value = value.replace(",", ".")
        else:
            return None
    elif "." in value:
        if not re.fullmatch(r"\d+\.\d+", value):
            return None
        integer, fraction = value.split(".")
        # '1.000' can mean one or one thousand; never silently select the larger value.
        if not explicit_decimal and len(fraction) == 3 and len(integer) <= 3:
            return None
    if not re.fullmatch(r"\d+(?:\.\d+)?", value):
        return None
    try:
        parsed = Decimal(value)
        return parsed if parsed.is_finite() and parsed >= 0 and parsed.adjusted() <= 308 else None
    except InvalidOperation:
        return None


def _amount_and_scale(text: str) -> tuple[Decimal | None, Decimal | None]:
    part = text.strip()
    matched = _SCALE_END.search(part)
    scale = None
    if matched:
        scale = _SCALE_LOOKUP[matched.group(1).lower()]
        part = part[:matched.start()].strip()
        if not part:
            part = "1"
    return _budget_decimal(part), scale


def parse_budget(value, default_currency: str = "EGP") -> tuple[float | None, str]:
    """Return a conservative amount in its own currency, with no conversion."""
    raw = _raw_text(value)
    currency = str(default_currency or "").upper().strip()
    if currency not in CURRENCY_OPTIONS:
        return None, ""
    if not raw.strip() or _is_formula(raw):
        return None, currency
    text = _INVISIBLE_FORMATTING.sub("", _fold_arabic(raw)).casefold().strip()
    text, currency, ambiguous = _budget_currency(text, currency)
    if ambiguous:
        return None, currency
    text = text.strip().replace("−", "-").replace("–", "-").replace("—", "-")
    if not text or any(character in text for character in "()<>%٪=;"):
        return None, currency
    text = re.sub(r"^(?:حوالي|تقريبا|قرابه|نحو|about|approximately|approx\.?)\s*", "", text).strip()
    between = bool(re.match(r"^(?:بين|between)\s*", text))
    text = re.sub(r"^(?:من|from|بين|between)\s*", "", text).strip()
    # Number-unit combinations such as 'one million and a half' are otherwise
    # ambiguous if arbitrary words are discarded. Accept only this explicit form.
    text = re.sub(r"(مليون|million)\s*(?:و\s*نصف|و\s*نص|and\s+a\s+half)$", r"\1 ونصف", text)
    text = re.sub(r"^مليون\s*ونصف$", "1.5 مليون", text)
    text = re.sub(r"^million\s*ونصف$", "1.5 million", text)
    if text.startswith("-") or re.search(r"(?<![eE])-\s*-", text):
        return None, currency
    connector = r"\s*(?:(?<![eE])-|\bto\b|(?<![\u0621-\u064a])الي(?![\u0621-\u064a]))\s*"
    if between:
        connector = r"\s*(?:(?<![eE])-|\bto\b|\band\b|(?<![\u0621-\u064a])(?:الي|و)(?![\u0621-\u064a]))\s*"
    pieces = re.split(connector, text)
    if len(pieces) > 2 or any(not piece.strip() for piece in pieces):
        return None, currency
    if len(pieces) == 1:
        piece = pieces[0].strip()
        if piece.endswith("+"):
            piece = piece[:-1].strip()
        number, scale = _amount_and_scale(piece)
        if number is None:
            return None, currency
        amount = number * (scale or Decimal(1))
    else:
        left, left_scale = _amount_and_scale(pieces[0])
        right, right_scale = _amount_and_scale(pieces[1])
        if left is None or right is None:
            return None, currency
        if right_scale is not None and left_scale is None and left <= right:
            # A shared trailing unit: '100-200k' or '100 إلى 200 ألف'.
            left_scale = right_scale
        if left_scale is not None and right_scale is None and right < left * left_scale:
            # A missing trailing unit cannot safely be inferred from '100k-200'.
            return None, currency
        amount = min(left * (left_scale or Decimal(1)), right * (right_scale or Decimal(1)))
    try:
        result = float(amount)
    except (OverflowError, ValueError):
        return None, currency
    if not math.isfinite(result) or result < 0 or (amount != 0 and result == 0):
        return None, currency
    return result, currency


_READINESS_NEGATION = re.compile(
    r"\b(?:not\s+(?:now|ready|yet|interested|immediately|today|available)|not\s+.*\bready|never|no\s*(?:thanks)?|"
    r"don\s+t\s+(?:want|need)|do\s+not\s+(?:want|need)|isn\s+t\s+ready|aren\s+t\s+ready|"
    r"غير\s+(?:جاهز|جاهزه|مستعد|مستعده|مهتم|متاح)|مش\s+(?:جاهز|جاهزه|دلوقتي|دلوقت|مستعد|مهتم|الان|عايز|عاوز|قادر)|"
    r"مو\s+(?:جاهز|الحين|مستعد)|مب\s+جاهز|ليس\s+(?:جاهزا|جاهز|الان|مستعدا)|"
    r"لست\s+(?:جاهزا|جاهز|مستعدا)|لسه\s+(?:لا|مش|بدري)|لا\s+(?:الان|حاليا|اريد|استطيع)|"
    r"لم\s+(?:استعد|اجهز))\b"
)
_READINESS_LATER = re.compile(
    r"\b(?:later|soon|tomorrow|next|after|waiting|wait|eventually|months?|weeks?|years?|"
    r"لاحقا|بعد|بعدين|قريبا|بكره|بكرا|غدا|انتظار|الانتظار|مؤجل|مؤجله|التاجيل|"
    r"شهر|شهور|اشهر|شهرين|اسبوع|اسابيع|اسبوعين|سنه|سنوات|السنه|عام|اعوام|القادم|القادمه)\b"
)
_READINESS_NOW = re.compile(
    r"\b(?:yes|y|now|ready|immediate|immediately|today|asap|"
    r"نعم|ايوه|ايوا|اجل|فوري|فورا|حالا|حاليا|الان|دلوقتي|دلوقت|جاهز|جاهزه|مستعد|مستعده|اليوم|الحين)\b"
)


def classify_readiness(value) -> str:
    raw = _raw_text(value)
    if not raw.strip() or _is_formula(raw):
        return "غير محدد"
    text = _normal_words(raw)
    if _READINESS_NEGATION.search(text) or text in {"لا", "كلا", "غير مستعد للشراء", "not", "false", "0"}:
        return "لاحقًا"
    if _READINESS_LATER.search(text):
        return "لاحقًا"
    if _READINESS_NOW.search(text) or text in {"true", "1"}:
        return "فوري"
    return "غير محدد"


def _frame_source_rows(frame: pd.DataFrame) -> list[int]:
    if len(frame) == 0:
        return []
    stored = frame.attrs.get("source_rows")
    if stored is None:
        header = frame.attrs.get("header_row", -1)
        offset = int(header) + 2 if isinstance(header, numbers.Integral) and header >= 0 else 1
        return list(range(offset, offset + len(frame)))
    if not isinstance(stored, (list, tuple)):
        raise UserDataError("بيانات ترقيم الصفوف الأصلية غير صالحة؛ أعد بناء الجدول من المصدر.")
    if all(isinstance(index, numbers.Integral) and 0 <= int(index) < len(stored) for index in frame.index):
        chosen = [stored[int(index)] for index in frame.index]
    elif len(stored) == len(frame):
        chosen = list(stored)
    else:
        raise UserDataError("فُقد الربط بين الجدول وأرقام صفوف المصدر؛ أعد بناء الجدول قبل التنظيف.")
    if any(isinstance(value, bool) or not isinstance(value, numbers.Integral) or value < 1 for value in chosen):
        raise UserDataError("يجب أن تكون أرقام صفوف المصدر أعدادًا صحيحة تبدأ من واحد.")
    chosen = [int(value) for value in chosen]
    if len(set(chosen)) != len(chosen):
        raise UserDataError("توجد أرقام صفوف أصلية مكررة؛ أعد بناء الجدول للحفاظ على هوية كل صف.")
    return chosen


def _clean_frame_types(frame: pd.DataFrame) -> pd.DataFrame:
    for column in ("phone_valid", "whatsapp_valid", "vip", "duplicate"):
        frame[column] = frame[column].astype(bool)
    frame["source_row"] = frame["source_row"].astype("int64")
    # Object dtype preserves None rather than silently changing unknown budgets to NaN.
    frame["budget_value"] = pd.Series(list(frame["budget_value"]), index=frame.index, dtype=object)
    return frame


def clean_leads(frame: pd.DataFrame, mapping: dict[str, str | None], source_label: str,
                default_region: str = "EG", vip_threshold: float = 100000,
                default_currency: str = "EGP", wa_fallback: bool = False) -> pd.DataFrame:
    if not isinstance(frame, pd.DataFrame):
        raise UserDataError("البيانات المدخلة ليست جدولًا صالحًا.")
    _check_limits(len(frame), len(frame.columns))
    if not frame.columns.is_unique or any(not isinstance(column, str) for column in frame.columns):
        raise UserDataError("يجب أن تكون أسماء الأعمدة نصية وفريدة؛ أنشئ الجدول من المصدر قبل التنظيف.")
    if not isinstance(mapping, dict):
        raise UserDataError("ربط الأعمدة غير صالح؛ اختر عمودًا واحدًا لكل حقل.")
    selected = {}
    for field in FIELD_LABELS:
        column = mapping.get(field)
        if column is not None:
            if not isinstance(column, str) or column not in frame.columns:
                raise UserDataError(f"العمود المرتبط بحقل {FIELD_LABELS[field]} غير موجود في الجدول.")
            if column in selected.values():
                raise UserDataError("لا يمكن ربط العمود نفسه بأكثر من حقل؛ راجع خيارات الربط.")
            selected[field] = column
    region = str(default_region or "").strip().upper()
    if region not in COUNTRY_OPTIONS:
        raise UserDataError("اختر دولة افتراضية مدعومة لأرقام التواصل.")
    currency = str(default_currency or "").strip().upper()
    if currency not in CURRENCY_OPTIONS:
        raise UserDataError("اختر عملة افتراضية مدعومة للميزانية وحد العميل المميز.")
    try:
        if isinstance(vip_threshold, bool):
            raise ValueError
        threshold = float(vip_threshold)
        if not math.isfinite(threshold) or threshold < 0:
            raise ValueError
    except (TypeError, ValueError, OverflowError) as exc:
        raise UserDataError("حد العميل المميز يجب أن يكون مبلغًا غير سالب ومحدودًا بالعملة الافتراضية.") from exc
    if not isinstance(wa_fallback, bool):
        raise UserDataError("خيار استخدام الهاتف للواتساب يجب أن يكون مفعّلًا أو غير مفعّل.")
    source_rows = _frame_source_rows(frame)
    columns = list(frame.columns)
    source = _raw_text(source_label)
    source_identity = _raw_text(frame.attrs.get("source_id"))
    mapped_columns = set(selected.values())
    raw_headers = frame.attrs.get("original_headers", columns)
    header_signature = tuple(_normal_words(value) for value in raw_headers) if len(raw_headers) == len(columns) else ()
    records = []
    row_iterator = frame.itertuples(index=False, name=None) if columns else (() for _ in range(len(frame)))
    for position, values in enumerate(row_iterator):
        original = {column: _raw_text(value) for column, value in zip(columns, values)}
        raw = {field: original.get(selected.get(field), "") for field in FIELD_LABELS}
        phone = normalize_phone(raw["phone"], region)
        whatsapp = normalize_phone(raw["whatsapp"], region)
        whatsapp_origin = "من المصدر" if raw["whatsapp"].strip() else "غير متوفر"
        notes = []
        if wa_fallback and not raw["whatsapp"].strip() and phone["valid"]:
            whatsapp = dict(phone)
            # Do not invent an original WhatsApp value when the source cell was empty.
            whatsapp["raw"] = raw["whatsapp"]
            whatsapp["error"] = "استُخدم الهاتف الصالح لأن خانة واتساب فارغة؛ وجود حساب واتساب غير مؤكد."
            whatsapp_origin = "من الهاتف (افتراضي)"
        name = unicodedata.normalize("NFC", raw["name"]).strip()
        if not name:
            name = "بدون اسم"
            notes.append("الاسم غير موجود؛ احتُفظ بالصف باسم بديل.")
        if phone["error"]:
            notes.append("الهاتف: " + phone["error"])
        if whatsapp["error"]:
            notes.append("واتساب: " + whatsapp["error"])
        budget_value, budget_currency = parse_budget(raw["budget"], currency)
        if raw["budget"].strip() and budget_value is None:
            notes.append("الميزانية غير مفهومة أو سالبة أو ملتبسة؛ لم تُحوّل إلى مبلغ تقديري.")
        elif budget_value is not None and budget_currency != currency:
            notes.append("عملة الميزانية تختلف عن عملة حد العميل المميز؛ لم يُجرَ تحويل عملات أو تصنيف مميز على هذا الأساس.")
        vip = bool(budget_value is not None and budget_currency == currency and budget_value >= threshold)
        readiness = classify_readiness(raw["timeline"])
        if raw["timeline"].strip() and readiness == "غير محدد":
            notes.append("جاهزية الشراء غير محددة من النص الأصلي.")
        if any(_is_formula(value) for value in original.values() if value):
            notes.append("توجد صيغة في الصف؛ احتُفظ بنصها دون تنفيذ.")
        # Repeated headings remain ordinary records, with a visible review flag.
        if any(header_signature) and all(_normal_words(original[column]) == header_signature[index]
                                         for index, column in enumerate(columns)):
            notes.append("قد يكون هذا صف عناوين مكررًا؛ احتُفظ به للمراجعة ولم يُحذف.")
        row_number = source_rows[position]
        identity = json.dumps([source_identity, source, row_number], ensure_ascii=False, separators=(",", ":"))
        lead_id = hashlib.sha256(identity.encode("utf-8", errors="surrogatepass")).hexdigest()
        details = {column: value for column, value in original.items() if column not in mapped_columns}
        records.append({
            "lead_id": lead_id, "source": source, "source_row": row_number, "name": name,
            "phone_raw": raw["phone"], "phone_e164": phone["e164"], "phone_display": phone["display"],
            "phone_valid": phone["valid"], "phone_region": phone["region"], "phone_error": phone["error"],
            "tel_url": phone["tel_url"], "whatsapp_raw": raw["whatsapp"], "whatsapp_e164": whatsapp["e164"],
            "whatsapp_display": whatsapp["display"], "whatsapp_valid": whatsapp["valid"],
            "whatsapp_error": whatsapp["error"], "wa_url": whatsapp["wa_url"], "whatsapp_origin": whatsapp_origin,
            "job": unicodedata.normalize("NFC", raw["job"]).strip(), "budget_raw": raw["budget"],
            "budget_value": budget_value, "budget_currency": budget_currency, "vip": vip,
            "timeline": raw["timeline"].strip(), "readiness": readiness,
            "city": unicodedata.normalize("NFC", raw["city"]).strip(),
            "details": json.dumps(details, ensure_ascii=False, separators=(",", ":")),
            "original": json.dumps(original, ensure_ascii=False, separators=(",", ":")),
            "issues": "؛ ".join(dict.fromkeys(notes)), "duplicate": False,
        })
    result = pd.DataFrame(records, columns=CLEAN_COLUMNS, dtype=object)
    result = _clean_frame_types(result)
    result.attrs = copy.deepcopy(frame.attrs)
    result.attrs.update({
        "source_rows": list(source_rows), "mapping": dict(selected), "source_label": source,
        "default_region": region, "default_currency": currency, "vip_threshold": threshold,
        "wa_fallback": wa_fallback,
        "validation_note": "صلاحية الأرقام تعني مطابقة خطة الترقيم فقط، ولا تثبت نشاط الخط أو ملكية حساب واتساب.",
    })
    return result


def _true_flags(frame: pd.DataFrame, column: str) -> list[bool]:
    if column not in frame.columns:
        return [False] * len(frame)
    result = []
    for value in frame[column].tolist():
        if value is None or value is pd.NA:
            result.append(False)
        elif isinstance(value, str):
            result.append(value.strip().casefold() in {"true", "1", "نعم"})
        else:
            try:
                result.append(bool(value == True))
            except (ValueError, TypeError):
                result.append(False)
    return result


def mark_duplicates(frame: pd.DataFrame) -> pd.DataFrame:
    """Mark every participating row, including phone-to-WhatsApp collisions."""
    result = frame.copy(deep=True)
    result.attrs = copy.deepcopy(frame.attrs)
    contacts = {}
    phone_flags = _true_flags(frame, "phone_valid")
    wa_flags = _true_flags(frame, "whatsapp_valid")
    phone_values = frame["phone_e164"].tolist() if "phone_e164" in frame else [""] * len(frame)
    wa_values = frame["whatsapp_e164"].tolist() if "whatsapp_e164" in frame else [""] * len(frame)
    for position, (phone, wa) in enumerate(zip(phone_values, wa_values)):
        row_contacts = set()
        for valid, value in ((phone_flags[position], phone), (wa_flags[position], wa)):
            key = _raw_text(value).strip()
            if valid and re.fullmatch(r"\+[1-9]\d{6,14}", key):
                row_contacts.add(key)
        for key in row_contacts:
            contacts.setdefault(key, []).append(position)
    duplicate_flags = [False] * len(frame)
    for positions in contacts.values():
        if len(positions) > 1:
            for position in positions:
                duplicate_flags[position] = True
    result["duplicate"] = pd.Series(duplicate_flags, index=frame.index, dtype=bool)
    return result


def calculate_kpis(frame: pd.DataFrame) -> dict:
    total = len(frame)
    phone_flags = _true_flags(frame, "phone_valid")
    wa_flags = _true_flags(frame, "whatsapp_valid")
    valid_phone, valid_wa = sum(phone_flags), sum(wa_flags)
    ready_count = int(frame["readiness"].eq("فوري").fillna(False).sum()) if "readiness" in frame else 0
    if not (valid_phone or valid_wa):
        top_contact = "لا توجد وسيلة تواصل صالحة"
    elif valid_phone == valid_wa:
        top_contact = "الهاتف وواتساب بالتساوي"
    elif valid_phone > valid_wa:
        top_contact = "الهاتف"
    else:
        top_contact = "واتساب"
    return {
        "total": total, "valid_phone": valid_phone, "valid_wa": valid_wa,
        "vip_count": sum(_true_flags(frame, "vip")), "ready_count": ready_count,
        "ready_rate": round(ready_count * 100 / total, 2) if total else 0.0,
        "reachable_count": sum(phone or wa for phone, wa in zip(phone_flags, wa_flags)),
        "top_contact": top_contact, "duplicates": sum(_true_flags(frame, "duplicate")),
    }


# Export segment: concatenated after the cleaning core, before the UI.
# Keep all user input in memory. Only the bundled, non-user font is read from disk.
import base64 as _export_base64
import binascii as _export_binascii
import csv as _export_csv
import datetime as _export_datetime
import html as _export_html
from html.parser import HTMLParser as _ExportHTMLParser
import io as _export_io
import json as _export_json
import math as _export_math
import os as _export_os
from pathlib import Path as _ExportPath
import re as _export_re
import struct as _export_struct
import subprocess as _export_subprocess
import sys as _export_sys
import threading as _export_threading
import time as _export_time
import types as _export_types
import unicodedata as _export_unicodedata
import warnings as _export_warnings
import zipfile as _export_zipfile

import pandas as pd


EXPORT_COLUMNS = (
    "lead_id", "source", "source_row", "name", "phone_raw", "phone_e164",
    "phone_display", "phone_valid", "phone_region", "phone_error", "tel_url",
    "whatsapp_raw", "whatsapp_e164", "whatsapp_display", "whatsapp_valid",
    "whatsapp_error", "wa_url", "whatsapp_origin", "job", "budget_raw",
    "budget_value", "budget_currency", "vip", "timeline", "readiness", "city",
    "details", "original", "issues", "duplicate",
)
EXPORT_EXCEL_HEADERS = (
    "معرّف العميل", "الملف المصدر", "صف المصدر الأصلي", "الاسم", "الهاتف الأصلي",
    "الهاتف الدولي", "الهاتف للعرض", "الهاتف صالح", "بلد الهاتف", "مراجعة الهاتف",
    "رابط الاتصال", "واتساب الأصلي", "واتساب الدولي", "واتساب للعرض", "واتساب صالح",
    "مراجعة واتساب", "رابط واتساب", "مصدر رقم واتساب", "المهنة", "الميزانية الأصلية",
    "قيمة الميزانية", "عملة الميزانية", "عميل مميز", "موعد الشراء", "الجاهزية",
    "المدينة", "الحقول الإضافية", "السجل الأصلي", "ملاحظات المراجعة", "تكرار محتمل",
)
PDF_MAX_HTML_BYTES = 4 * 1024 * 1024
PDF_RENDER_TIMEOUT_SECONDS = 90
PDF_WORKER_MEMORY_BYTES = 700 * 1024 * 1024
PDF_MAX_OUTPUT_BYTES = 32 * 1024 * 1024
_EXPORT_LOGO_MAX_BYTES = 2 * 1024 * 1024
_EXPORT_LOGO_MAX_PIXELS = 12_000_000
_EXPORT_FONT_MAX_BYTES = 2 * 1024 * 1024
_EXPORT_CELL_CHUNK = 30_000
_EXPORT_PREVIEW_MARK = "… [مختصر]"
_EXPORT_HEADER_ROW = 7
_EXPORT_FIRST_DATA_ROW = 8
_EXPORT_LEADS_SHEET = "العملاء"
_EXPORT_SUMMARY_SHEET = "الملخص"
_EXPORT_ORIGINAL_SHEET = "الخصائص الأصلية"
_EXPORT_E164_RE = _export_re.compile(r"\+[1-9][0-9]{6,14}\Z", _export_re.ASCII)
_EXPORT_TEL_RE = _export_re.compile(r"tel:\+[1-9][0-9]{6,14}\Z", _export_re.ASCII)
_EXPORT_WA_RE = _export_re.compile(r"https://wa\.me/[1-9][0-9]{6,14}\Z", _export_re.ASCII)
_EXPORT_DATA_RE = _export_re.compile(
    r"data:(image/png|font/ttf);base64,([A-Za-z0-9+/]*={0,2})\Z", _export_re.ASCII
)
_EXPORT_XML_BAD_RE = _export_re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_EXPORT_NAVY = "142C42"
_EXPORT_SLATE = "536779"
_EXPORT_TEAL = "087F83"
_EXPORT_LIGHT = "EFF6F7"

# Streamlit can execute a fresh app module for each rerun/session. This process-
# scoped module holds a semaphore ONLY, never a frame, HTML, logo or PDF cache.
_export_gate_module = _export_sys.modules.setdefault(
    "_folio_pdf_render_gate", _export_types.ModuleType("_folio_pdf_render_gate")
)
_EXPORT_PDF_SEMAPHORE = _export_gate_module.__dict__.setdefault(
    "semaphore", _export_threading.BoundedSemaphore(1)
)


def _export_text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    try:
        if bool(pd.isna(value)):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, (dict, list, tuple)):
        return _export_json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def _export_bool(value):
    # Match the core's _true_flags rule exactly, including non-boolean inputs.
    if value is None or value is pd.NA:
        return False
    if isinstance(value, str):
        return value.strip().casefold() in {"true", "1", "نعم"}
    try:
        return bool(value == True)
    except (TypeError, ValueError):
        return False


def _export_number(value):
    if isinstance(value, bool) or not _export_text(value).strip():
        return None
    try:
        result = float(value)
        return result if _export_math.isfinite(result) else None
    except (TypeError, ValueError, OverflowError):
        return None


def _export_check_frame(frame):
    if not isinstance(frame, pd.DataFrame):
        raise UserDataError("تعذّر التصدير: جدول العملاء غير صالح.")
    if frame.columns.has_duplicates or set(frame.columns) != set(EXPORT_COLUMNS):
        raise UserDataError("تعذّر التصدير: أعمدة جدول العملاء لا تطابق النسخة المعتمدة.")
    return frame.loc[:, list(EXPORT_COLUMNS)]


def _export_plain_option(value, label, limit, default=""):
    text = _export_text(value).strip() or default
    if len(text) > limit or _EXPORT_XML_BAD_RE.search(text):
        raise UserDataError(f"{label}: استخدم نصاً عادياً لا يتجاوز {limit} حرفاً دون محارف تحكّم.")
    return text


def _export_preview(value, limit=90):
    text = " ".join(_export_text(value).split())
    text = _EXPORT_XML_BAD_RE.sub("", text)
    if len(text) > limit:
        return text[:max(1, limit - len(_EXPORT_PREVIEW_MARK))] + _EXPORT_PREVIEW_MARK
    return text


def _export_escape(value):
    return _export_html.escape(_export_text(value), quote=True)


def _export_map(value):
    text = _export_text(value)
    if not text:
        return {}
    try:
        parsed = _export_json.loads(text)
    except (ValueError, TypeError, RecursionError):
        raise UserDataError("تعذّر التصدير: السجل الأصلي أو الحقول الإضافية ليست خريطة نصية صالحة.") from None
    if not isinstance(parsed, dict):
        raise UserDataError("تعذّر التصدير: يجب أن تكون الخصائص الأصلية خريطة من أسماء الحقول وقيمها.")
    return {_export_text(key): _export_text(item) for key, item in parsed.items()}


def _export_contact_url(row, channel):
    # Never trust raw tel_url/wa_url cells, even if they look like links.
    # Links are rebuilt only from a core-validated canonical number.
    if channel == "phone":
        number = _export_text(row.get("phone_e164"))
        if _export_bool(row.get("phone_valid")) and _EXPORT_E164_RE.fullmatch(number):
            return "tel:" + number
    elif channel == "whatsapp":
        number = _export_text(row.get("whatsapp_e164"))
        if _export_bool(row.get("whatsapp_valid")) and _EXPORT_E164_RE.fullmatch(number):
            return "https://wa.me/" + number[1:]
    return ""


def sanitize_logo(data: bytes) -> bytes:
    """Decode a bounded, single-frame raster and return a fresh metadata-free PNG."""
    from PIL import Image, ImageOps

    if not isinstance(data, bytes) or not data or len(data) > _EXPORT_LOGO_MAX_BYTES:
        raise UserDataError("ارفع شعاراً بصيغة صورة لا يتجاوز حجمه ٢ ميغابايت.")
    try:
        with _export_warnings.catch_warnings():
            _export_warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(_export_io.BytesIO(data)) as image:
                if image.format not in {"PNG", "JPEG", "WEBP"}:
                    raise UserDataError("الشعار يجب أن يكون صورة PNG أو JPEG أو WebP فقط؛ الصور المتجهة غير مسموحة.")
                width, height = image.size
                if width < 1 or height < 1 or width * height > _EXPORT_LOGO_MAX_PIXELS:
                    raise UserDataError("أبعاد الشعار تتجاوز الحد المسموح: ١٢ مليون بكسل.")
                if getattr(image, "n_frames", 1) != 1 or getattr(image, "is_animated", False):
                    raise UserDataError("استخدم شعاراً ثابتاً من إطار واحد؛ الصور المتحركة غير مسموحة.")
                image.verify()
            with Image.open(_export_io.BytesIO(data)) as image:
                image.load()
                oriented = ImageOps.exif_transpose(image)
                oriented.thumbnail((512, 256), Image.Resampling.LANCZOS)
                rgba = oriented.convert("RGBA")
                clean = Image.new("RGBA", rgba.size)
                clean.paste(rgba, (0, 0))
                output = _export_io.BytesIO()
                # Saving a new image without pnginfo/exif/icc_profile strips metadata.
                clean.save(output, format="PNG", optimize=False)
                return output.getvalue()
    except UserDataError:
        raise
    except Exception:
        # Image plugins can raise several exception types on malformed input.
        # Never expose decoder diagnostics or raw file content to the UI/logs.
        raise UserDataError("تعذّر قراءة الشعار بأمان. اختر صورة ثابتة سليمة بصيغة PNG أو JPEG أو WebP.") from None


def _export_font_bytes():
    """Read the trusted local font, with a lossless text-transport alternative."""
    path = _ExportPath(__file__).resolve().parent / "assets" / "Cairo.ttf"
    expected_hash = "667c987182391c91f4e57a2f455b1794fb5e3ee6ca4ef3383e86bb690fa9c964"
    try:
        if path.is_file():
            if not 12 <= path.stat().st_size <= _EXPORT_FONT_MAX_BYTES:
                raise OSError
            data = path.read_bytes()
        else:
            encoded_path = path.with_suffix(".ttf.b64")
            if not 12 <= encoded_path.stat().st_size <= 2 * _EXPORT_FONT_MAX_BYTES:
                raise OSError
            encoded = b"".join(encoded_path.read_bytes().split())
            data = _export_base64.b64decode(encoded, validate=True)
            if _export_base64.b64encode(data) != encoded:
                raise ValueError
        if not 12 <= len(data) <= _EXPORT_FONT_MAX_BYTES:
            raise ValueError
        if data[:4] not in {b"\x00\x01\x00\x00", b"true"}:
            raise ValueError
        if hashlib.sha256(data).hexdigest() != expected_hash:
            raise ValueError
        return data
    except (OSError, ValueError, _export_binascii.Error):
        raise UserDataError("الخط العربي المضمّن غير متاح أو لا يطابق النسخة المعتمدة. أعد تجهيز ملفات التطبيق قبل التصدير.") from None


def _export_data_url(mime, data):
    return "data:" + mime + ";base64," + _export_base64.b64encode(data).decode("ascii")


def _export_check_sanitized_png(data):
    # Accept only the shape produced by sanitize_logo: RGBA PNG, no metadata,
    # animation, alternate profiles or trailing payloads. Inspect before decoding.
    if not data.startswith(b"\x89PNG\r\n\x1a\n") or len(data) > _EXPORT_LOGO_MAX_BYTES:
        raise ValueError("مورد صورة غير مسموح.")
    offset, kinds, ended = 8, [], False
    while offset + 12 <= len(data):
        size = _export_struct.unpack(">I", data[offset:offset + 4])[0]
        kind = data[offset + 4:offset + 8]
        end = offset + 12 + size
        if end > len(data) or kind not in {b"IHDR", b"IDAT", b"IEND"}:
            raise ValueError("مورد صورة غير مسموح.")
        payload = data[offset + 8:offset + 8 + size]
        checksum = _export_struct.unpack(">I", data[offset + 8 + size:end])[0]
        if (_export_binascii.crc32(kind + payload) & 0xFFFFFFFF) != checksum:
            raise ValueError("مورد صورة غير مسموح.")
        kinds.append(kind)
        if kind == b"IHDR":
            if len(kinds) != 1 or size != 13:
                raise ValueError("مورد صورة غير مسموح.")
            width, height, depth, mode, compression, filtering, interlace = _export_struct.unpack(">IIBBBBB", payload)
            if not (1 <= width <= 512 and 1 <= height <= 256 and depth == 8
                    and mode == 6 and compression == 0 and filtering == 0 and interlace == 0):
                raise ValueError("مورد صورة غير مسموح.")
        if kind == b"IEND":
            if size or end != len(data):
                raise ValueError("مورد صورة غير مسموح.")
            ended = True
            break
        offset = end
    if not ended or not kinds or kinds[0] != b"IHDR" or b"IDAT" not in kinds:
        raise ValueError("مورد صورة غير مسموح.")
    from PIL import Image
    try:
        with Image.open(_export_io.BytesIO(data)) as image:
            image.load()
    except (OSError, SyntaxError, ValueError):
        raise ValueError("مورد صورة غير مسموح.") from None


def safe_pdf_url_fetcher(url, *args, **kwargs):
    """WeasyPrint 68 legacy callable. No default fetcher or network fallback."""
    if not isinstance(url, str) or len(url) > 4 * _EXPORT_FONT_MAX_BYTES // 3 + 128:
        raise ValueError("طلب مورد خارجي محظور.")
    match = _EXPORT_DATA_RE.fullmatch(url)
    if not match:
        raise ValueError("طلب مورد خارجي محظور.")
    mime, encoded = match.groups()
    try:
        data = _export_base64.b64decode(encoded, validate=True)
    except (ValueError, _export_binascii.Error):
        raise ValueError("مورد مضمّن غير صالح.") from None
    # Canonical base64 only: no percent escapes, extra padding or whitespace.
    if _export_base64.b64encode(data).decode("ascii") != encoded:
        raise ValueError("مورد مضمّن غير صالح.")
    if mime == "font/ttf":
        if data != _export_font_bytes():
            raise ValueError("لا يسمح إلا بالخط المضمّن المعتمد.")
    else:
        _export_check_sanitized_png(data)
    return {"string": data, "mime_type": mime, "redirected_url": url}


# 68.0 recognizes this opt-in and aborts instead of warning-and-continuing when
# its legacy callable raises. It is not a network allowlist; our callable is one.
safe_pdf_url_fetcher._fail_on_errors = True


_EXPORT_REPORT_CSS = """
@page {
  size: A4 landscape; margin: 13mm 13mm 16mm;
  @bottom-right { content: "تقرير العملاء · نسخة للتصدير"; font-family: Cairo; font-size: 8pt; color: #536779; }
  @bottom-left { content: "صفحة " counter(page) " / " counter(pages); font-family: Cairo; font-size: 8pt; color: #536779; }
}
* { box-sizing: border-box; }
html { direction: rtl; }
body { font-family: Cairo, sans-serif; color: #142c42; font-size: 9pt; line-height: 1.6; margin: 0; }
h1, h2, p { margin: 0; }
.masthead { display: table; width: 100%; padding-bottom: 4mm; border-bottom: 1.2mm solid #087f83; }
.heading { display: table-cell; vertical-align: middle; }
.eyebrow { font-size: 8pt; font-weight: 700; color: #087f83; letter-spacing: .2pt; }
h1 { font-size: 23pt; line-height: 1.6; overflow-wrap: anywhere; }
.subtitle, .muted { color: #536779; }
.logo-cell { display: table-cell; width: 44mm; text-align: left; vertical-align: middle; }
.logo { max-width: 40mm; max-height: 19mm; }
.context { margin-top: 3mm; padding: 2mm 3mm; border-right: 1mm solid #087f83; background: #eff6f7; overflow-wrap: anywhere; }
.metrics { display: table; width: 100%; table-layout: fixed; border-spacing: 2mm 0; margin: 5mm -2mm 3mm; }
.metric { display: table-cell; background: #142c42; color: white; border-radius: 2mm; padding: 2mm 3mm; vertical-align: top; }
.metric-value { font-size: 21pt; font-weight: 700; line-height: 1.3; }
.metric-label { font-size: 8pt; color: #dbe9ee; }
.quick { color: #536779; font-size: 8pt; margin: 2mm 0; overflow-wrap: anywhere; }
.notice { background: #f0f7f7; border: .25mm solid #ccdddd; border-radius: 2mm; padding: 2.5mm 3mm; margin: 3mm 0 4mm; font-size: 8pt; }
.notice strong { color: #087f83; }
.leads { width: 100%; border-collapse: collapse; table-layout: fixed; direction: rtl; font-size: 8pt; }
.leads col:nth-child(1) { width: 7%; }
.leads col:nth-child(2) { width: 17%; }
.leads col:nth-child(3) { width: 14%; }
.leads col:nth-child(4) { width: 14%; }
.leads col:nth-child(5) { width: 11%; }
.leads col:nth-child(6) { width: 11%; }
.leads col:nth-child(7) { width: 10%; }
.leads col:nth-child(8) { width: 16%; }
.leads thead { display: table-header-group; }
.leads tfoot { display: table-footer-group; }
.leads th { background: #142c42; color: white; text-align: right; font-weight: 700; padding: 2.2mm 1.7mm; }
.leads td { vertical-align: top; padding: 2.2mm 1.7mm; border-bottom: .2mm solid #dae4e9; overflow-wrap: anywhere; }
.leads tbody tr:nth-child(even) { background: #f3f7f9; }
.leads tr { break-inside: avoid; page-break-inside: avoid; }
.leads tfoot td { color: #536779; background: white; font-size: 7pt; padding: 2mm; }
.client { font-size: 9pt; font-weight: 700; }
.trace { color: #536779; font-size: 7pt; }
.contact { direction: ltr; unicode-bidi: embed; display: inline-block; white-space: nowrap; font-size: 9pt; font-weight: 700; }
a { color: #087f83; text-decoration: none; }
.pills { margin-top: 1mm; }
.pill { display: inline-block; padding: .1mm 1.4mm; border-radius: 1.2mm; margin: .3mm 0 .3mm .7mm; font-size: 7pt; }
.pill-vip { background: #eee4c9; color: #725618; }
.pill-ready { background: #dcefeb; color: #126455; }
.pill-invalid { background: #f8e5e3; color: #922e28; }
.pill-duplicate { background: #ede7f5; color: #67408c; }
.detail { font-size: 7pt; color: #536779; }
.issue { font-size: 7pt; color: #922e28; margin-top: 1mm; }
.empty { text-align: center; padding: 12mm; color: #536779; }
"""


def _export_report_css():
    return ("@font-face { font-family: Cairo; src: url('" +
            _export_data_url("font/ttf", _export_font_bytes()) +
            "') format('truetype'); font-weight: 100 900; font-style: normal; }\n" +
            _EXPORT_REPORT_CSS)


def _export_is_ready(row):
    # The core's classified label, never a keyword search of the raw timeline.
    return _export_text(row.get("readiness")) == "فوري"


def _export_budget_label(row):
    value = _export_number(row.get("budget_value"))
    currency = _export_preview(row.get("budget_currency"), 14)
    if value is not None:
        amount = f"{value:,.0f}" if value.is_integer() else f"{value:,.2f}"
        return _export_preview(amount + (" " + currency if currency else ""), 38)
    return _export_preview(row.get("budget_raw"), 38) or "غير محددة"


def _export_top_budgets(frame):
    maxima = {}
    for row in frame.to_dict("records"):
        amount = _export_number(row.get("budget_value"))
        currency = _export_text(row.get("budget_currency")).strip()
        if amount is not None and currency:
            maxima[currency] = max(amount, maxima.get(currency, amount))
    if not maxima:
        return "لا توجد ميزانيات رقمية بعملة محددة في هذه الدفعة."
    values = [f"{_export_preview(currency, 14)}: {amount:,.0f}"
              for currency, amount in sorted(maxima.items())[:4]]
    suffix = "؛ توجد عملات أخرى في إكسل." if len(maxima) > 4 else "."
    return "أعلى ميزانية لكل عملة في هذه الدفعة — " + " · ".join(values) + suffix


def _export_contact_html(row, channel):
    url = _export_contact_url(row, channel)
    prefix = "phone" if channel == "phone" else "whatsapp"
    canonical = _export_text(row.get(prefix + "_e164"))
    raw = row.get(prefix + "_raw")
    if url:
        return (f'<a class="contact" dir="ltr" href="{_export_escape(url)}">'
                f'{_export_escape(canonical)}</a>')
    label = _export_preview(raw, 27) or "غير متوفر"
    return (f'<span dir="ltr">{_export_escape(label)}</span>'
            '<div class="trace">لا يوجد رابط صالح</div>')


def build_report_html(frame, brand, logo=None, generated_at: str | None = None,
                      context_note: str = "") -> str:
    frame = _export_check_frame(frame)
    if len(frame) > PDF_BATCH_SIZE:
        raise UserDataError(f"تقرير PDF يقبل {PDF_BATCH_SIZE} عميلاً في الدفعة كحد أقصى. اختر دفعة أصغر؛ لم تُحذف أي صفوف.")
    brand = _export_plain_option(brand, "اسم الجهة", 120, "تقرير العملاء")
    context_note = _export_plain_option(context_note, "وصف الدفعة", 600)
    stamp = _export_plain_option(
        generated_at, "وقت التصدير", 80,
        _export_datetime.datetime.now(_export_datetime.timezone.utc).strftime("%Y-%m-%d %H:%M") + " — التوقيت العالمي",
    )
    css = _export_report_css()
    logo_html = ""
    if logo is not None:
        cleaned = sanitize_logo(logo)
        logo_html = ('<div class="logo-cell"><img class="logo" alt="شعار الجهة" src="' +
                     _export_data_url("image/png", cleaned) + '"></div>')
    kpis = calculate_kpis(frame)
    rate = _export_number(kpis.get("ready_rate")) or 0
    metric_values = (
        (str(int(kpis.get("total", 0))), "عملاء هذه الدفعة"),
        (str(int(kpis.get("reachable_count", 0))), "قابلون للتواصل"),
        (str(int(kpis.get("vip_count", 0))), "عملاء مميزون"),
        (f"{rate:.1f}٪", "نسبة الجاهزين الآن"),
        (str(int(kpis.get("duplicates", 0))), "تكرارات محتملة"),
    )
    metrics = "".join(
        '<div class="metric"><div class="metric-value">' + _export_escape(value) +
        '</div><div class="metric-label">' + _export_escape(label) + '</div></div>'
        for value, label in metric_values
    )
    quick = (f"هواتف صالحة: {int(kpis.get('valid_phone', 0))} · "
             f"أرقام واتساب صالحة تنسيقياً: {int(kpis.get('valid_wa', 0))} · "
             f"جاهزون الآن: {int(kpis.get('ready_count', 0))} · "
             "قناة التواصل الأوفر: " + _export_preview(kpis.get("top_contact"), 40))
    rows = []
    for row in frame.to_dict("records"):
        pills = []
        if _export_bool(row.get("vip")):
            pills.append('<span class="pill pill-vip">مميز</span>')
        if _export_is_ready(row):
            pills.append('<span class="pill pill-ready">فوري</span>')
        if not _export_contact_url(row, "phone"):
            pills.append('<span class="pill pill-invalid">هاتف غير صالح</span>')
        if _export_text(row.get("whatsapp_raw")) and not _export_contact_url(row, "whatsapp"):
            pills.append('<span class="pill pill-invalid">واتساب غير صالح</span>')
        if _export_bool(row.get("duplicate")):
            pills.append('<span class="pill pill-duplicate">تكرار محتمل</span>')
        extras = _export_map(row.get("details"))
        detail_parts = [
            f"{_export_preview(key, 20)}: {_export_preview(value, 38)}"
            for key, value in list(extras.items())[:2]
        ]
        detail_text = " · ".join(detail_parts) or "لا توجد حقول إضافية"
        if len(extras) > 2:
            detail_text += " · " + _EXPORT_PREVIEW_MARK
        issues = _export_preview(row.get("issues"), 64)
        issue_html = ('<div class="issue">' + _export_escape(issues) + '</div>') if issues else ""
        source = _export_preview(row.get("source"), 35)
        trace = f"{source} · صف {_export_preview(row.get('source_row'), 12)}"
        origin = _export_preview(row.get("whatsapp_origin"), 30)
        origin_html = '<div class="trace">' + _export_escape(origin) + '</div>' if origin else ""
        rows.append(
            '<tr><td><strong>' + _export_escape(_export_preview(row.get("lead_id"), 19)) +
            '</strong><div class="trace">' + _export_escape(trace) + '</div></td>' +
            '<td><div class="client">' + _export_escape(_export_preview(row.get("name"), 54) or "دون اسم") +
            '</div><div class="pills">' + "".join(pills) + '</div></td>' +
            '<td>' + _export_contact_html(row, "phone") + '</td>' +
            '<td>' + _export_contact_html(row, "whatsapp") + origin_html + '</td>' +
            '<td>' + _export_escape(_export_preview(row.get("job"), 32) or "غير محددة") +
            '<div class="trace">' + _export_escape(_export_preview(row.get("city"), 25)) + '</div></td>' +
            '<td>' + _export_escape(_export_budget_label(row)) + '</td>' +
            '<td>' + _export_escape(_export_preview(row.get("readiness"), 34) or "غير محددة") +
            '<div class="trace">' + _export_escape(_export_preview(row.get("timeline"), 32)) + '</div></td>' +
            '<td><div class="detail">' + _export_escape(detail_text) + '</div>' + issue_html + '</td></tr>'
        )
    if not rows:
        rows.append('<tr><td class="empty" colspan="8">لا توجد سجلات ضمن هذه الدفعة.</td></tr>')
    note_html = '<div class="context">' + _export_escape(context_note) + '</div>' if context_note else ""
    report = (
        '<!DOCTYPE html><html lang="ar" dir="rtl"><head><meta charset="utf-8">'
        '<title>' + _export_escape(brand) + ' — تقرير العملاء</title><style>' + css +
        '</style></head><body><header class="masthead"><div class="heading">'
        '<div class="eyebrow">تقرير العملاء المحتملين</div><h1>' + _export_escape(brand) +
        '</h1><p class="subtitle">' + _export_escape(stamp) + '</p></div>' + logo_html +
        '</header>' + note_html + '<section class="metrics">' + metrics + '</section>' +
        '<p class="quick">' + _export_escape(quick) + '</p>' +
        '<p class="quick">' + _export_escape(_export_top_budgets(frame)) + '</p>' +
        '<div class="notice"><strong>البيانات الكاملة في إكسل.</strong> '
        'يعرض هذا التقرير معاينة مختصرة للحقول الطويلة؛ علامة «… [مختصر]» تعني اختصار النص. '
        'تحتفظ ورقة «الخصائص الأصلية» في إكسل بكل الحقول والقيم الأصلية، ويمكن تنزيل الملفات المرفوعة ضمن أرشيف المصادر. '
        'المؤشرات تخص هذه الدفعة فقط. صلاحية تنسيق الرقم لا تؤكد وجود حساب واتساب أو هوية صاحبه.</div>'
        '<table class="leads"><colgroup><col><col><col><col><col><col><col><col></colgroup>'
        '<thead><tr><th>السجل والمصدر</th><th>العميل والحالة</th><th>الهاتف</th><th>واتساب</th>'
        '<th>المهنة والمدينة</th><th>الميزانية</th><th>الجاهزية</th><th>تفاصيل ومراجعة</th></tr></thead>'
        '<tfoot><tr><td colspan="8">معرّف السجل والملف وصف المصدر يتيحان التتبّع. '
        'روابط الاتصال تُفتح يدوياً، والنصوص المطوّلة متاحة بالكامل في إكسل.</td></tr></tfoot>'
        '<tbody>' + "".join(rows) + '</tbody></table></body></html>'
    )
    if len(report.encode("utf-8")) > PDF_MAX_HTML_BYTES:
        raise UserDataError("حجم التقرير يتجاوز حد الأمان البالغ ٤ ميغابايت. قلّل حجم الدفعة أو الشعار.")
    return report


class _ExportReportGuard(_ExportHTMLParser):
    """Validate our fixed template again at the worker trust boundary."""
    _tags = {
        "html", "head", "meta", "title", "style", "body", "header", "div", "h1",
        "p", "section", "strong", "span", "a", "img", "table", "colgroup", "col",
        "thead", "tfoot", "tbody", "tr", "th", "td",
    }
    _void = {"meta", "img", "col"}
    _extra_attrs = {"meta": {"charset"}, "a": {"href"}, "img": {"src", "alt"}, "td": {"colspan"}}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack = []
        self.styles = []
        self.counts = {}
        self.body_rows = 0
        self.nodes = 0
        self.doctype = False

    @staticmethod
    def reject():
        raise ValueError("بنية تقرير غير مسموحة.")

    def handle_decl(self, decl):
        if decl.casefold() != "doctype html" or self.doctype:
            self.reject()
        self.doctype = True

    def handle_pi(self, data):
        self.reject()

    def handle_comment(self, data):
        self.reject()

    def handle_starttag(self, tag, attrs):
        self.nodes += 1
        if tag not in self._tags or self.nodes > 40_000 or len(self.stack) > 15:
            self.reject()
        self.counts[tag] = self.counts.get(tag, 0) + 1
        if len(dict(attrs)) != len(attrs):
            self.reject()
        for key, value in attrs:
            if key not in {"class", "dir", "lang"} | self._extra_attrs.get(tag, set()):
                self.reject()
            if value is None:
                self.reject()
            if key == "class" and not _export_re.fullmatch(r"[a-z -]{1,80}", value):
                self.reject()
            if key == "dir" and value not in {"rtl", "ltr"}:
                self.reject()
            if key == "lang" and value != "ar":
                self.reject()
            if key == "charset" and value != "utf-8":
                self.reject()
            if key == "colspan" and value != "8":
                self.reject()
            if key == "href" and not (_EXPORT_TEL_RE.fullmatch(value) or _EXPORT_WA_RE.fullmatch(value)):
                self.reject()
            if key == "alt" and value != "شعار الجهة":
                self.reject()
            if key == "src":
                if not value.startswith("data:image/png;base64,"):
                    self.reject()
                safe_pdf_url_fetcher(value)
        if tag == "tr" and self.stack and self.stack[-1] == "tbody":
            self.body_rows += 1
            if self.body_rows > PDF_BATCH_SIZE:
                self.reject()
        if tag not in self._void:
            self.stack.append(tag)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in self._void:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if not self.stack or self.stack[-1] != tag:
            self.reject()
        self.stack.pop()

    def handle_data(self, data):
        if self.stack and self.stack[-1] == "style":
            self.styles.append(data)
        elif len(data) > 2048 or _EXPORT_XML_BAD_RE.search(data):
            self.reject()

    def finish(self):
        self.close()
        required = ("html", "head", "title", "style", "body", "table", "thead", "tfoot", "tbody")
        if self.stack or not self.doctype or any(self.counts.get(tag) != 1 for tag in required):
            self.reject()
        if self.counts.get("img", 0) > 1 or "".join(self.styles) != _export_report_css():
            self.reject()


def _export_validate_worker_html(html_bytes):
    if not isinstance(html_bytes, bytes) or not html_bytes or len(html_bytes) > PDF_MAX_HTML_BYTES:
        raise ValueError("حجم تقرير غير مسموح.")
    text = html_bytes.decode("utf-8", errors="strict")
    guard = _ExportReportGuard()
    guard.feed(text)
    guard.finish()
    return text


def _export_set_worker_limits():
    if not _export_sys.platform.startswith("linux"):
        raise RuntimeError("يتطلب التصدير الآمن بيئة لينكس الداعمة لحدود الموارد.")
    import resource

    def limit(which, soft, hard=None):
        _, old_hard = resource.getrlimit(which)
        hard = soft if hard is None else hard
        if old_hard != resource.RLIM_INFINITY:
            hard = min(hard, old_hard)
        resource.setrlimit(which, (min(soft, hard), hard))

    # RLIMIT_AS bounds address space (not merely resident memory). Keep BLAS
    # single-threaded at process launch to avoid gigabyte virtual reservations.
    limit(resource.RLIMIT_AS, PDF_WORKER_MEMORY_BYTES)
    limit(resource.RLIMIT_CPU, 80, 85)
    # Pango may materialize the bundled public font. Never write HTML, logos or
    # PDFs; all user-derived resources stay in bytes/BytesIO with an in-memory cache.
    limit(resource.RLIMIT_FSIZE, 8 * 1024 * 1024)
    limit(resource.RLIMIT_CORE, 0)
    limit(resource.RLIMIT_NOFILE, 128)


def handle_pdf_worker():
    """Call before any Streamlit UI code; exits only for the exact worker flag."""
    if _export_sys.argv[1:] != ["--render-pdf-worker"]:
        return False
    try:
        _export_set_worker_limits()
        import contextlib
        import logging
        logging.disable(logging.CRITICAL)
        # Nothing from imports/rendering may contaminate stdout or disclose input.
        with contextlib.redirect_stdout(_export_sys.stderr):
            raw_html = _export_sys.stdin.buffer.read(PDF_MAX_HTML_BYTES + 1)
            html_text = _export_validate_worker_html(raw_html)
            from weasyprint import HTML
            from weasyprint.text.fonts import FontConfiguration
            font_config = FontConfiguration()
            output = _export_io.BytesIO()
            HTML(
                string=html_text, encoding="utf-8", base_url=None,
                url_fetcher=safe_pdf_url_fetcher,
            ).write_pdf(
                target=output, font_config=font_config, presentational_hints=False,
                full_fonts=False, optimize_images=False, pdf_forms=False,
                cache={},
            )
            pdf_bytes = output.getvalue()
            if not pdf_bytes.startswith(b"%PDF-") or len(pdf_bytes) > PDF_MAX_OUTPUT_BYTES:
                raise ValueError("تعذّر إنشاء تقرير ضمن الحدود الآمنة.")
        _export_sys.stdout.buffer.write(pdf_bytes)
        _export_sys.stdout.buffer.flush()
    except Exception:
        # Exit status only: never echo input, URLs, HTML or exception details.
        raise SystemExit(2) from None
    raise SystemExit(0)


def _export_worker_pdf(html_bytes):
    env = _export_os.environ.copy()
    for variable in ("OPENBLAS_NUM_THREADS", "OMP_NUM_THREADS", "OMP_THREAD_LIMIT", "MKL_NUM_THREADS",
                     "NUMEXPR_NUM_THREADS", "NUMEXPR_MAX_THREADS", "BLIS_NUM_THREADS", "ARROW_IO_THREADS"):
        env[variable] = "1"
    env["MALLOC_ARENA_MAX"] = "2"
    env["ARROW_DEFAULT_MEMORY_POOL"] = "system"
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env["PYTHONUNBUFFERED"] = "1"
    env["PYTHONWARNINGS"] = "ignore"
    path = _ExportPath(__file__).resolve()
    process = None
    readers = []
    output = bytearray()
    output_overflow = _export_threading.Event()
    pipe_failure = _export_threading.Event()
    deadline = _export_time.monotonic() + PDF_RENDER_TIMEOUT_SECONDS
    try:
        process = _export_subprocess.Popen(
            [_export_sys.executable, "-B", str(path), "--render-pdf-worker"],
            stdin=_export_subprocess.PIPE, stdout=_export_subprocess.PIPE,
            stderr=_export_subprocess.DEVNULL, env=env, close_fds=True,
            start_new_session=True,
        )

        def send_input():
            try:
                process.stdin.write(html_bytes)
                process.stdin.close()
            except (BrokenPipeError, OSError, ValueError):
                pipe_failure.set()

        def read_output():
            try:
                while True:
                    chunk = process.stdout.read(65_536)
                    if not chunk:
                        return
                    if len(output) + len(chunk) > PDF_MAX_OUTPUT_BYTES:
                        output_overflow.set()
                        return
                    output.extend(chunk)
            except (OSError, ValueError):
                pipe_failure.set()

        readers = [
            _export_threading.Thread(target=send_input, daemon=True),
            _export_threading.Thread(target=read_output, daemon=True),
        ]
        for thread in readers:
            thread.start()
        while process.poll() is None:
            if _export_time.monotonic() >= deadline or output_overflow.is_set():
                raise TimeoutError
            try:
                process.wait(timeout=min(0.15, max(0.001, deadline - _export_time.monotonic())))
            except _export_subprocess.TimeoutExpired:
                pass
        for thread in readers:
            thread.join(timeout=max(0.001, deadline - _export_time.monotonic()))
        if (process.returncode != 0 or output_overflow.is_set() or pipe_failure.is_set()
                or any(thread.is_alive() for thread in readers)
                or not output.startswith(b"%PDF-") or not output.rstrip().endswith(b"%%EOF")):
            raise ValueError
        return bytes(output)
    except TimeoutError:
        raise UserDataError("توقّف إنشاء PDF عند حد الوقت أو الحجم الآمن. جرّب دفعة أصغر أو نزّل إكسل.") from None
    except (OSError, ValueError, _export_subprocess.SubprocessError):
        raise UserDataError("تعذّر إنشاء PDF ضمن حدود الأمان. جرّب دفعة أصغر، أو تحقّق من تثبيت مكوّنات الطباعة والخط العربي.") from None
    finally:
        if process is not None:
            if process.poll() is None:
                process.kill()
            try:
                process.wait(timeout=5)
            except _export_subprocess.TimeoutExpired:
                pass
            for thread in readers:
                thread.join(timeout=1)
            for pipe in (process.stdin, process.stdout):
                if pipe is not None:
                    try:
                        pipe.close()
                    except (OSError, ValueError):
                        pass


def build_pdf(frame, brand, logo=None, context_note: str = "") -> bytes:
    frame = _export_check_frame(frame)
    if len(frame) > PDF_BATCH_SIZE:
        raise UserDataError(f"اختر دفعة تضم {PDF_BATCH_SIZE} عميلاً أو أقل لتصدير PDF؛ لم تُختصر قائمة العملاء تلقائياً.")
    context_note = _export_plain_option(context_note, "وصف الدفعة", 600)
    # Fail fast instead of accumulating concurrent render jobs or user frames.
    if not _EXPORT_PDF_SEMAPHORE.acquire(blocking=False):
        raise UserDataError("توجد عملية إنشاء PDF قيد التنفيذ. انتظر اكتمالها ثم أعد المحاولة.")
    try:
        report = build_report_html(frame, brand, logo=logo, context_note=context_note)
        return _export_worker_pdf(report.encode("utf-8"))
    finally:
        _EXPORT_PDF_SEMAPHORE.release()


def _export_xml_text(value):
    return _EXPORT_XML_BAD_RE.sub(lambda match: f"\\u{ord(match.group()):04x}", _export_text(value))


def _export_excel_literal(cell, value):
    # Explicit string type defeats all formula prefixes without changing text.
    # A leading apostrophe would alter raw data, so it is deliberately not used.
    text = _export_xml_text(value)
    if len(text.encode("utf-16-le")) // 2 >= 32767:
        raise UserDataError("تجاوزت خلية حد إكسل؛ يجب تقسيم القيمة إلى أجزاء قبل التصدير.")
    cell.value = text
    cell.data_type = "s"
    cell.number_format = "@"
    return cell


def _export_excel_preview(value, limit=450):
    original = _export_text(value)
    text = _export_xml_text(original)
    changed = text != original or len(text) > limit
    if len(text) > limit:
        text = text[:limit - len(_EXPORT_PREVIEW_MARK)] + _EXPORT_PREVIEW_MARK
    elif text != original:
        text += " [محارف خاصة ممثّلة نصياً]"
    return text, changed


def _export_original_encoding(text):
    text = _export_text(text)
    if _EXPORT_XML_BAD_RE.search(text):
        # JSON safely and reversibly represents XML-forbidden control codepoints.
        return _export_json.dumps(text, ensure_ascii=True), "نص JSON؛ اجمع الأجزاء ثم فك ترميز JSON"
    return text, "نص حرفي؛ اجمع الأجزاء بحسب ترتيبها"


def _export_cell_chunks(text):
    # At most 15,000 Unicode codepoints / 30,000 UTF-16 units per cell, including
    # astral characters. No openpyxl 32,767-character silent truncation is possible.
    size = _EXPORT_CELL_CHUNK // 2
    return [text[index:index + size] for index in range(0, len(text), size)] or [""]


def _export_style_sheet(sheet):
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "صفحة &P من &N"
    sheet.oddFooter.center.size = 9
    sheet.oddFooter.center.font = "Cairo"
    sheet.oddFooter.right.text = "تقرير العملاء"
    sheet.sheet_properties.outlinePr.summaryRight = False


def _export_banner(sheet, brand, last_column, subtitle, logo_png=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.drawing.image import Image as ExcelImage
    from openpyxl.utils import get_column_letter

    end = get_column_letter(last_column)
    sheet.merge_cells(f"A1:{end}2")
    title = _export_excel_literal(sheet["A1"], brand)
    title.font = Font(name="Cairo", size=23, bold=True, color="FFFFFF")
    title.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2, indent=1)
    title.fill = PatternFill("solid", fgColor=_EXPORT_NAVY)
    for line in (1, 2):
        sheet.row_dimensions[line].height = 26
        for column in range(1, last_column + 1):
            sheet.cell(line, column).fill = PatternFill("solid", fgColor=_EXPORT_NAVY)
    sheet.merge_cells(f"A3:{end}3")
    cell = _export_excel_literal(sheet["A3"], subtitle)
    cell.font = Font(name="Cairo", size=10, color=_EXPORT_SLATE)
    cell.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2, wrap_text=True)
    sheet.row_dimensions[3].height = 32
    if logo_png:
        image = ExcelImage(_export_io.BytesIO(logo_png))
        scale = min(145 / image.width, 55 / image.height, 1)
        image.width, image.height = image.width * scale, image.height * scale
        sheet.add_image(image, f"{get_column_letter(max(2, last_column - 1))}1")


def _export_table_header(sheet, headers):
    from openpyxl.styles import Alignment, Font, PatternFill
    for index, label in enumerate(headers, start=1):
        cell = _export_excel_literal(sheet.cell(_EXPORT_HEADER_ROW, index), label)
        cell.fill = PatternFill("solid", fgColor=_EXPORT_NAVY)
        cell.font = Font(name="Cairo", color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2, wrap_text=True)
    sheet.row_dimensions[_EXPORT_HEADER_ROW].height = 32
    sheet.freeze_panes = "E8"
    sheet.print_title_rows = f"1:{_EXPORT_HEADER_ROW}"


def _export_finish_table(sheet, display_name, width, last_row):
    from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
    from openpyxl.worksheet.filters import AutoFilter
    from openpyxl.utils import get_column_letter
    # A header-only table is valid for an empty export; do not invent a lead.
    ref = f"A{_EXPORT_HEADER_ROW}:{get_column_letter(width)}{max(_EXPORT_HEADER_ROW, last_row)}"
    # Predeclare headers: openpyxl otherwise materializes the whole table range
    # just to discover its first row while saving.
    table = Table(displayName=display_name, ref=ref, autoFilter=AutoFilter(ref=ref), tableColumns=[
        TableColumn(id=index, name=str(sheet.cell(_EXPORT_HEADER_ROW, index).value))
        for index in range(1, width + 1)
    ])
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False)
    sheet.add_table(table)
    sheet.auto_filter.ref = ref
    sheet.print_area = f"A1:{get_column_letter(width)}{max(_EXPORT_HEADER_ROW, last_row)}"


def _export_save_excel_in_memory(book):
    # openpyxl's normal save(BytesIO) still writes sheet XML to temporary files.
    # This small 3.1.5-pinned override streams each sheet directly into the ZIP.
    # No global monkeypatches and no customer-data temporary files are involved.
    from openpyxl.writer.excel import ExcelWriter
    from openpyxl.worksheet._writer import WorksheetWriter
    from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing

    class MemoryExcelWriter(ExcelWriter):
        def write_worksheet(self, sheet):
            sheet._drawing = SpreadsheetDrawing()
            sheet._drawing.charts = sheet._charts
            sheet._drawing.images = sheet._images
            with self._archive.open(sheet.path[1:], "w", force_zip64=True) as stream:
                writer = WorksheetWriter(sheet, out=stream)
                writer.write()
                sheet._rels = writer._rels
            self.manifest.append(sheet)

    output = _export_io.BytesIO()
    with _export_zipfile.ZipFile(output, "w", compression=_export_zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        MemoryExcelWriter(book, archive).write_data()
    return output.getvalue()


def _export_cache_summary(data, frame):
    """Cache our fixed summary formulas without evaluating any uploaded formula."""
    from xml.etree import ElementTree as ET

    metrics = calculate_kpis(frame)
    values = {
        "B5": metrics["total"], "B6": metrics["valid_phone"], "B7": metrics["valid_wa"],
        "B8": metrics["vip_count"], "B9": metrics["ready_count"],
        "B10": metrics["ready_count"] / metrics["total"] if metrics["total"] else 0,
        "B11": metrics["reachable_count"], "B12": metrics["duplicates"],
        "B13": metrics["top_contact"],
    }
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    destination = _export_io.BytesIO()
    with _export_zipfile.ZipFile(_export_io.BytesIO(data)) as original:
        with _export_zipfile.ZipFile(destination, "w", compression=_export_zipfile.ZIP_DEFLATED) as target:
            for info in original.infolist():
                content = original.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    sheet = ET.fromstring(content)
                    seen = set()
                    for cell in sheet.iter(namespace + "c"):
                        reference = cell.get("r")
                        if reference not in values or cell.find(namespace + "f") is None:
                            continue
                        cached = cell.find(namespace + "v")
                        if cached is None:
                            cached = ET.SubElement(cell, namespace + "v")
                        value = values[reference]
                        cell.set("t", "str" if isinstance(value, str) else "n")
                        cached.text = str(value)
                        seen.add(reference)
                    if seen != set(values):
                        raise UserDataError("تعذّر تجهيز قيم ملخص Excel؛ بنية المعادلات غير متطابقة.")
                    content = ET.tostring(sheet, encoding="utf-8", xml_declaration=True)
                target.writestr(info, content)
    return destination.getvalue()


def build_excel(frame: pd.DataFrame, brand: str, logo: bytes | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    frame = _export_check_frame(frame)
    if len(frame) > 25_000:
        raise UserDataError("تصدير إكسل يقبل ٢٥ ألف عميل كحد أقصى في هذه العملية.")
    brand = _export_plain_option(brand, "اسم الجهة", 120, "تقرير العملاء")
    cleaned_logo = sanitize_logo(logo) if logo is not None else None
    book = Workbook()
    summary = book.active
    summary.title = _EXPORT_SUMMARY_SHEET
    leads = book.create_sheet(_EXPORT_LEADS_SHEET)
    originals = book.create_sheet(_EXPORT_ORIGINAL_SHEET)
    book.calculation = CalcProperties(
        calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto"
    )
    book.properties.creator = PRODUCT_NAME
    book.properties.title = "تقرير العملاء"
    book.properties.description = "بيانات العملاء مع المصدر والخصائص الأصلية؛ يعاد حساب المعادلات عند فتح إكسل."
    body_font = Font(name="Cairo", size=10, color=_EXPORT_NAVY)
    row_alignment = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)
    ltr_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, readingOrder=1)
    zebra_fill = PatternFill("solid", fgColor=_EXPORT_LIGHT)
    thin_border = Border(bottom=Side(style="hair", color="DCE6EB"))
    for sheet in (summary, leads, originals):
        _export_style_sheet(sheet)

    _export_banner(
        summary, brand, 8,
        "ملخص المجموعة المصدّرة فقط — تُعاد حسابات المعادلات عند فتح الملف في إكسل.", cleaned_logo,
    )
    _export_banner(
        leads, brand, len(EXPORT_COLUMNS),
        "معاينة للحقول الطويلة: علامة [مختصر] تعني اختصار النص. جميع القيم الكاملة في ورقة «الخصائص الأصلية» وأرشيف المصادر.",
    )
    original_headers = (
        "معرّف العميل", "الملف المصدر", "صف المصدر الأصلي", "مجموعة البيانات", "رقم الخاصية",
        "جزء الاسم", "أجزاء الاسم", "جزء القيمة", "أجزاء القيمة", "اسم الحقل الأصلي",
        "القيمة الأصلية", "ترميز الاسم", "ترميز القيمة",
    )
    _export_banner(
        originals, brand, len(original_headers),
        "كل الحقول المعيّنة وغير المعيّنة محفوظة. أعد تركيب الأسماء والقيم الطويلة بترتيب الأجزاء؛ يوضح عمود الترميز متى يلزم فك JSON.",
    )
    _export_table_header(leads, EXPORT_EXCEL_HEADERS)
    _export_table_header(originals, original_headers)
    originals.freeze_panes = "F8"
    for index, width in enumerate(
        (20, 28, 14, 28, 23, 23, 23, 13, 15, 28, 29, 23, 23, 23, 13,
         28, 32, 28, 25, 28, 20, 17, 14, 26, 21, 22, 43, 43, 43, 16), start=1
    ):
        leads.column_dimensions[get_column_letter(index)].width = width
    for index, width in enumerate((20, 30, 14, 20, 13, 12, 12, 12, 12, 36, 76, 35, 35), start=1):
        originals.column_dimensions[get_column_letter(index)].width = width
    for index, width in enumerate((28, 22, 18, 18, 18, 18, 18, 18), start=1):
        summary.column_dimensions[get_column_letter(index)].width = width

    original_row = _EXPORT_FIRST_DATA_ROW
    boolean_columns = {"phone_valid", "whatsapp_valid", "vip", "duplicate"}
    phone_columns = {"phone_raw", "phone_e164", "phone_display", "whatsapp_raw", "whatsapp_e164", "whatsapp_display"}
    for excel_row, record in enumerate(frame.to_dict("records"), start=_EXPORT_FIRST_DATA_ROW):
        overflow = []
        for index, key in enumerate(EXPORT_COLUMNS, start=1):
            cell = leads.cell(excel_row, index)
            raw_value = record.get(key)
            if key in boolean_columns:
                _export_excel_literal(cell, "نعم" if _export_bool(raw_value) else "لا")
            elif key == "budget_value":
                number = _export_number(raw_value)
                if number is None:
                    _export_excel_literal(cell, "")
                else:
                    cell.value = number
                    cell.number_format = '#,##0.00;[Red]-#,##0.00'
            elif key in {"tel_url", "wa_url"}:
                link = _export_contact_url(record, "phone" if key == "tel_url" else "whatsapp")
                _export_excel_literal(cell, link)
                if link:
                    cell.hyperlink = link
                if _export_text(raw_value) != link and _export_text(raw_value):
                    overflow.append((key, _export_text(raw_value)))
            else:
                limit = 900 if key in {"details", "original", "issues"} else 450
                text, shortened = _export_excel_preview(raw_value, limit)
                _export_excel_literal(cell, text)
                if shortened:
                    overflow.append((key, _export_text(raw_value)))
            cell.font = body_font
            cell.alignment = ltr_alignment if key in phone_columns | {"tel_url", "wa_url"} else row_alignment
            cell.border = thin_border
            if excel_row % 2 == 0:
                cell.fill = zebra_fill
            if key in {"phone_e164", "phone_display", "whatsapp_e164", "whatsapp_display"}:
                channel = "phone" if key.startswith("phone") else "whatsapp"
                link = _export_contact_url(record, channel)
                if link:
                    cell.hyperlink = link
            if cell.hyperlink:
                cell.font = Font(name="Cairo", size=10, color=_EXPORT_TEAL, underline="single")
            if key == "vip" and _export_bool(raw_value):
                cell.fill = PatternFill("solid", fgColor="EEE4C9")
            elif key == "duplicate" and _export_bool(raw_value):
                cell.fill = PatternFill("solid", fgColor="EDE7F5")
            elif key in {"phone_valid", "whatsapp_valid"} and not _export_bool(raw_value):
                cell.fill = PatternFill("solid", fgColor="F8E5E3")
        leads.row_dimensions[excel_row].height = 55
        original_map = _export_map(record.get("original"))
        extra_map = _export_map(record.get("details"))
        attributes = [("أصلي", key, value) for key, value in original_map.items()]
        attributes.extend(("إضافي", key, value) for key, value in extra_map.items()
                          if key not in original_map or original_map[key] != value)
        attributes.extend(("بيانات التصدير", EXPORT_EXCEL_HEADERS[EXPORT_COLUMNS.index(key)], value) for key, value in overflow)
        for attribute_index, (group, key, value) in enumerate(attributes, start=1):
            encoded_key, key_encoding = _export_original_encoding(key)
            encoded_value, value_encoding = _export_original_encoding(value)
            key_parts, value_parts = _export_cell_chunks(encoded_key), _export_cell_chunks(encoded_value)
            part_count = max(len(key_parts), len(value_parts))
            if original_row + part_count - 1 > 1_048_576:
                raise UserDataError("الخصائص الأصلية تتجاوز حد صفوف إكسل؛ صدّر مجموعة أصغر مع أرشيف المصادر.")
            for part_index in range(part_count):
                output_values = (
                    record.get("lead_id"), record.get("source"), record.get("source_row"), group,
                    attribute_index, part_index + 1 if part_index < len(key_parts) else "", len(key_parts),
                    part_index + 1 if part_index < len(value_parts) else "", len(value_parts),
                    key_parts[part_index] if part_index < len(key_parts) else "",
                    value_parts[part_index] if part_index < len(value_parts) else "", key_encoding, value_encoding,
                )
                for column_index, item in enumerate(output_values, start=1):
                    target = originals.cell(original_row, column_index)
                    if column_index in {5, 6, 7, 8, 9} and isinstance(item, int):
                        target.value = item
                    else:
                        # Source identifiers should already be short. Refuse, not
                        # silently truncate, if an upstream contract is violated.
                        _export_excel_literal(target, item)
                    target.font = body_font
                    target.alignment = row_alignment
                    if original_row % 2 == 0:
                        target.fill = zebra_fill
                originals.row_dimensions[original_row].height = 45
                original_row += 1

    last_row = _EXPORT_FIRST_DATA_ROW + len(frame) - 1
    _export_finish_table(leads, "LeadsExport", len(EXPORT_COLUMNS), last_row)
    _export_finish_table(originals, "OriginalAttributes", len(original_headers), original_row - 1)
    reference_last = max(_EXPORT_FIRST_DATA_ROW, last_row)

    def column_range(key):
        letter = get_column_letter(EXPORT_COLUMNS.index(key) + 1)
        return f"'{_EXPORT_LEADS_SHEET}'!${letter}${_EXPORT_FIRST_DATA_ROW}:${letter}${reference_last}"

    phone_range, wa_range = column_range("phone_valid"), column_range("whatsapp_valid")
    # Static criteria mirror calculate_kpis. No user text is interpolated into
    # any Excel formula, even inside quotes.
    ready_formula = f'=IFERROR(COUNTIF({column_range("readiness")},"فوري"),0)'
    total_formula = f'=IFERROR(ROWS({column_range("lead_id")}),0)' if len(frame) else '=IFERROR(0,0)'
    formulas = (
        ("إجمالي العملاء", total_formula),
        ("هواتف صالحة", f'=IFERROR(COUNTIF({phone_range},"نعم"),0)'),
        ("أرقام واتساب صالحة تنسيقياً", f'=IFERROR(COUNTIF({wa_range},"نعم"),0)'),
        ("عملاء مميزون", f'=IFERROR(COUNTIF({column_range("vip")},"نعم"),0)'),
        ("جاهزون الآن", ready_formula),
        ("نسبة الجاهزين الآن", "=IFERROR(B9/B5,0)"),
        ("قابلون للتواصل", f'=IFERROR(COUNTIF({phone_range},"نعم")+COUNTIF({wa_range},"نعم")-COUNTIFS({phone_range},"نعم",{wa_range},"نعم"),0)'),
        ("تكرارات محتملة", f'=IFERROR(COUNTIF({column_range("duplicate")},"نعم"),0)'),
        ("قناة التواصل الأوفر", '=IFERROR(IF(MAX(B6:B7)=0,"لا توجد وسيلة تواصل صالحة",IF(B6=B7,"الهاتف وواتساب بالتساوي",IF(B7>B6,"واتساب","الهاتف"))),"لا توجد وسيلة تواصل صالحة")'),
    )
    for index, (label, formula) in enumerate(formulas, start=5):
        _export_excel_literal(summary.cell(index, 1), label)
        target = summary.cell(index, 2)
        target.value = formula  # Only application-owned formulas are written.
        target.font = Font(name="Cairo", size=15, bold=True, color=_EXPORT_TEAL)
        target.number_format = "0.0%" if index == 10 else "0"
        for column in (1, 2):
            summary.cell(index, column).alignment = Alignment(horizontal="right", vertical="center", readingOrder=2, wrap_text=True)
            summary.cell(index, column).fill = zebra_fill if index % 2 else PatternFill("solid", fgColor="FFFFFF")
        summary.cell(index, 1).font = body_font
        summary.row_dimensions[index].height = 33
    summary["B13"].number_format = "General"
    notes = (
        "الملخص يحتوي معادلات مع قيم محسوبة مسبقاً من هذه المجموعة. يعيد إكسل حسابها عند فتح الملف؛ لم تُنفّذ أي صيغة من المصدر.",
        "الميزانيات أرقام في ورقة العملاء، والعملات منفصلة. لا تُجمع عملات مختلفة ولا يحدث أي تحويل للعملة.",
        "الهواتف نصوص للحفاظ على + والأصفار. روابط الاتصال مبنية فقط من الأرقام المصنفة صالحة، ولا تؤكد وجود حساب واتساب.",
        "الخصائص الأصلية تحفظ كل الحقول، المعيّنة والإضافية. اجمع أجزاء القيمة بالترتيب؛ الخلايا التي يذكر ترميزها JSON تحتاج فك ترميز بعد الجمع.",
        "معاينة ورقة العملاء قد تختصر النصوص الطويلة صراحة. القيم الكاملة محفوظة في الخصائص الأصلية، وأرشيف المصادر يحفظ الملفات المرفوعة كما هي.",
    )
    for index, note in enumerate(notes, start=16):
        summary.merge_cells(start_row=index, start_column=1, end_row=index, end_column=8)
        cell = _export_excel_literal(summary.cell(index, 1), note)
        cell.font = Font(name="Cairo", size=10, color=_EXPORT_SLATE)
        cell.alignment = row_alignment
        summary.row_dimensions[index].height = 34
    summary.freeze_panes = "A5"
    summary.print_area = "A1:H20"
    try:
        return _export_cache_summary(_export_save_excel_in_memory(book), frame)
    finally:
        book.close()


def demo_csv() -> bytes:
    """Synthetic records only. Phone examples do not represent real customers."""
    headers = ["اسم العميل", "رقم الموبايل", "رقم الواتساب", "المهنة", "الميزانية", "موعد الشراء",
               "المدينة", "نوع العقار", "قناة الإعلان", "ملاحظات إضافية", "نوع البيانات"]
    examples = [
        ["عميل تجريبي ٠١", "+201000000001", "", "طبيب", "٣٬٥٠٠٬٠٠٠ جنيه مصري", "جاهز فوراً", "القاهرة", "شقة", "حملة تجريبية", "يرغب في معاينة تجريبية"],
        ["عميل تجريبي ٠٢", "+966500000001", "+966500000001", "مهندس", "1.2 مليون ريال سعودي", "خلال شهر", "الرياض", "فيلا", "حملة تجريبية", "تفصيل إضافي غير معيّن"],
        ["عميل تجريبي ٠٣", "+971500000001", "+971500000001", "صاحب شركة", "AED 900000", "غير جاهز الآن", "دبي", "مكتب", "حملة تجريبية", "النفي لا يعني جاهزية فورية"],
        ["عميل تجريبي ٠٤", "٠١٠٠٠٠٠٠٠٠٤", "", "مدرس", "750,000 EGP", "بعد ٣ أشهر", "الجيزة", "شقة", "حملة تجريبية", "رقم محلي تجريبي بصيغة مصر"],
        ["عميل تجريبي ٠٥", "+966550000002", "+966550000002", "طبيبة", "$150,000", "جاهز الآن", "جدة", "فيلا", "حملة تجريبية", "عملة مختلفة دون تحويل"],
        ["عميل تجريبي ٠٦", "+971520000002", "غير صالح", "محاسب", "١٬١٠٠٬٠٠٠ درهم إماراتي", "لست جاهزاً", "أبوظبي", "شقة", "حملة تجريبية", "واتساب صريح غير صالح؛ يحتاج مراجعة"],
        ["عميل تجريبي ٠٧", "0123?", "", "مصمم", "غير محددة", "لا أريد الشراء الآن", "الإسكندرية", "استوديو", "حملة تجريبية", "هاتف معيب مقصود لاختبار المراجعة"],
        ["عميل تجريبي ٠٨", "", "+201100000008", "مهندسة", "2 مليون جنيه", "فوراً", "القاهرة", "شقة", "حملة تجريبية", "الهاتف مفقود؛ يوجد واتساب تجريبي"],
        ["عميل تجريبي ٠٩", "+201000000001", "", "طبيب", "٣٬٥٠٠٬٠٠٠ جنيه مصري", "جاهز فوراً", "القاهرة", "شقة", "حملة تجريبية", "تكرار مقصود للرقم في السجل الأول"],
        ["عميل تجريبي ١٠", "+966540000010", "", "طالب", "SAR 450000", "مش جاهز حالياً", "الدمام", "شقة", "حملة تجريبية", "نفي عامي للجاهزية"],
        ["عميل تجريبي ١١", "+971560000011", "", "مدير", "EUR 120000", "خلال أسبوع", "الشارقة", "مكتب", "حملة تجريبية", "ميزانية باليورو محفوظة بعملتها"],
        ["عميل تجريبي ١٢", "+201200000012", "", "طبيبة", "EGP 1800000", "لم أحدد موعداً", "المنصورة", "عيادة", "حملة تجريبية", "سجل تجريبي دون موعد شراء مؤكد"],
    ]
    output = _export_io.StringIO(newline="")
    writer = _export_csv.writer(output, lineterminator="\r\n")
    writer.writerow(headers)
    for row in examples:
        writer.writerow(row + ["بيانات اصطناعية للتجربة فقط؛ ليست عملاء حقيقيين ولا تُستخدم للاتصال"])
    return output.getvalue().encode("utf-8-sig")


def _export_safe_basename(name):
    text = _export_unicodedata.normalize("NFC", _export_text(name)).replace("\\", "/")
    text = text.rsplit("/", 1)[-1]
    text = "".join(char for char in text if _export_unicodedata.category(char) not in {"Cc", "Cf", "Cs"})
    text = _export_re.sub(r'[<>:"/\\|?*]', "_", text).strip(" .")
    if not text:
        text = "مصدر"
    stem, suffix = _export_os.path.splitext(text)
    if not stem:
        stem = "مصدر"
    if stem.upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        stem = "مصدر_" + stem
    suffix = suffix[:20]
    stem = stem[:120]
    # Stay within common 255-byte filesystem filename limits, even in Arabic.
    while len((stem + suffix).encode("utf-8")) > 200:
        stem = stem[:-1]
    return (stem or "مصدر") + suffix


def build_sources_archive(uploads: list[tuple[str, bytes]]) -> bytes:
    if not isinstance(uploads, list) or len(uploads) > 5:
        raise UserDataError("أرشيف المصادر يقبل الملفات المرفوعة صراحة، بحد أقصى خمسة ملفات.")
    output = _export_io.BytesIO()
    used = set()
    with _export_zipfile.ZipFile(output, "w", compression=_export_zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for upload in uploads:
            if not isinstance(upload, (tuple, list)) or len(upload) != 2:
                raise UserDataError("تعذّر تجهيز أرشيف المصادر: ملف مرفوع غير صالح.")
            raw_name, data = upload
            if not isinstance(raw_name, str) or not isinstance(data, bytes) or len(data) > 20 * 1024 * 1024:
                raise UserDataError("تعذّر تجهيز أرشيف المصادر: يجب ألا يتجاوز الملف ٢٠ ميغابايت.")
            name = _export_safe_basename(raw_name)
            stem, suffix = _export_os.path.splitext(name)
            unique, counter = name, 2
            while unique.casefold() in used:
                unique = f"{stem} ({counter}){suffix}"
                counter += 1
            used.add(unique.casefold())
            # No filesystem paths, symlinks, extra metadata, manifests or other
            # session data are ever added. Bytes are exactly the explicit upload.
            info = _export_zipfile.ZipInfo(unique, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = _export_zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o100600 << 16
            archive.writestr(info, data)
    return output.getvalue()


# This guard must stay before the UI segment: worker subprocesses must never
# import or execute Streamlit UI/authentication/session code.
if __name__ == "__main__" and _export_sys.argv[1:] == ["--render-pdf-worker"]:
    handle_pdf_worker()


# Streamlit interface. Processing and export functions above remain importable.
import html as _ui_html
import os as _ui_os
from pathlib import Path as _UIPath
from dataclasses import asdict as _ui_asdict
import time as _ui_time

import streamlit as st


APP_TITLE = PRODUCT_NAME
ENCODING_OPTIONS = {"auto": "تلقائي", "utf-8-sig": "UTF-8", "cp1256": "Windows-1256 · عربي",
                    "cp1252": "Windows-1252 · لاتيني", "utf-16": "UTF-16", "utf-32": "UTF-32",
                    "iso-8859-6": "ISO-8859-6 · عربي"}
DELIMITER_OPTIONS = {"auto": "تلقائي", ",": "فاصلة ,", ";": "فاصلة منقوطة ;", "\t": "علامة تبويب", "|": "فاصل |"}
CONTACT_FILTERS = {"all": "كل السجلات", "reachable": "قابل للتواصل", "phone": "هاتف صالح",
                   "whatsapp": "واتساب صالح تنسيقياً", "unreachable": "بلا وسيلة صالحة"}
FLAG_FILTERS = {"all": "كل الحالات", "vip": "عملاء مميزون", "duplicate": "تكرارات محتملة",
               "review": "بها ملاحظات مراجعة"}


def _ui_escape(value):
    return _ui_html.escape(_raw_text(value), quote=True)


def _ui_short(value, limit=100):
    text = _INVISIBLE_FORMATTING.sub("", _raw_text(value))
    return text if len(text) <= limit else text[:limit - 1] + "…"


def _ui_settings(name, default=""):
    if name in _ui_os.environ:
        return _ui_os.environ[name]
    try:
        return st.secrets.get(name, default)
    except FileNotFoundError:
        paths = st.get_option("secrets.files") or []
        if any(_UIPath(path).expanduser().is_file() for path in paths):
            raise UserDataError("تعذّرت قراءة إعدادات المصادقة. راجع Secrets؛ لن يُفتح التطبيق تلقائياً.") from None
        return default


def auth_access_allowed(claims, allowed_emails, now=None):
    """Authorize an OIDC identity, not just a successful authentication event."""
    if not isinstance(claims, dict) or claims.get("email_verified") is not True:
        return False
    allowlist = {str(email).strip().casefold() for email in allowed_emails if str(email).strip()}
    email = str(claims.get("email", "")).strip().casefold()
    if not email or email not in allowlist:
        return False
    expiry = claims.get("exp")
    if expiry is not None:
        try:
            if isinstance(expiry, bool) or not math.isfinite(float(expiry)) or float(expiry) <= (_ui_time.time() if now is None else now):
                return False
        except (TypeError, ValueError, OverflowError):
            return False
    return True


def _clear_workspace():
    next_epoch = int(st.session_state.get("lc_epoch", 0)) + 1
    for key in list(st.session_state):
        if str(key).startswith("lc_"):
            del st.session_state[key]
    st.session_state["lc_epoch"] = next_epoch


def _logout():
    _clear_workspace()
    st.logout()


def _require_auth():
    try:
        mode = str(_ui_settings("AUTH_MODE", "open")).strip().casefold()
        if mode == "open":
            return False
        if mode != "oidc":
            raise UserDataError("قيمة AUTH_MODE غير صالحة. اختر oidc أو open صراحةً.")
        raw_allowlist = _ui_settings("ALLOWED_EMAILS", "")
        allowlist = raw_allowlist.split(",") if isinstance(raw_allowlist, str) else list(raw_allowlist)
        if not any(str(email).strip() for email in allowlist):
            raise UserDataError("المصادقة مفعّلة لكن قائمة ALLOWED_EMAILS فارغة. الوصول مغلق حتى ضبطها.")
        try:
            auth = st.secrets.get("auth", {})
        except FileNotFoundError:
            auth = {}
        required = ("redirect_uri", "cookie_secret", "client_id", "client_secret", "server_metadata_url")
        if not all(auth.get(key) for key in required):
            raise UserDataError("أكمل إعداد مزوّد OIDC داخل قسم auth في Secrets. الوصول مغلق حتى اكتمال الإعداد.")
        if len(str(auth.get("cookie_secret", ""))) < 32:
            raise UserDataError("مفتاح جلسات تسجيل الدخول قصير. استخدم سراً عشوائياً بطول ٣٢ محرفاً أو أكثر في Secrets.")
        if not st.user.is_logged_in:
            st.title("دخول آمن إلى " + PRODUCT_NAME)
            st.caption("سجّل الدخول بالحساب المعتمد لدى الجهة. لا تُقرأ أي ملفات قبل التحقق من الهوية.")
            st.button("تسجيل الدخول", on_click=st.login, type="primary", icon=":material/lock:")
            st.stop()
        if not auth_access_allowed(st.user.to_dict(), allowlist):
            st.error("هذا الحساب غير مسموح، أو البريد غير موثّق، أو انتهت صلاحية الهوية. اطلب من مسؤول الجهة مراجعة إعدادات الوصول.")
            st.button("تسجيل الخروج", on_click=_logout)
            st.stop()
        st.sidebar.button("تسجيل الخروج", on_click=_logout, icon=":material/logout:")
        return True
    except (UserDataError, TypeError, ValueError):
        st.error("الوصول مغلق: إعدادات المصادقة غير مكتملة أو غير صالحة. راجع AUTH_MODE وALLOWED_EMAILS وقسم auth في Secrets وفق دليل النشر.")
        st.stop()


def _ui_style(dark=False):
    palette = ({"bg": "#0b1424", "panel": "#142033", "ink": "#e5edf7", "muted": "#a4b6ca",
                "line": "#2b3b50", "soft": "#1a2b40", "accent": "#42c5b5"} if dark else
               {"bg": "#f5f7fb", "panel": "#ffffff", "ink": "#172b46", "muted": "#596f85",
                "line": "#dde5ed", "soft": "#edf6f7", "accent": "#087f83"})
    font = _export_data_url("font/ttf", _export_font_bytes())
    css = """
    @font-face {font-family:Cairo;src:url('__FONT__') format('truetype');font-weight:200 1000;font-display:swap}
    :root {--lc-bg:__bg__;--lc-panel:__panel__;--lc-ink:__ink__;--lc-muted:__muted__;--lc-line:__line__;--lc-soft:__soft__;--lc-accent:__accent__}
    .stApp, [data-testid="stAppViewContainer"], [data-testid="stHeader"] {background:var(--lc-bg);color:var(--lc-ink)}
    .stApp, input, textarea, button, h1,h2,h3,h4, label, p, [data-testid="stMetricValue"], [data-testid="stMarkdownContainer"] {font-family:Cairo,Arial,sans-serif}
    [data-testid="stMainBlockContainer"] {max-width:1560px;padding:4.8rem 2.4rem 4rem;direction:rtl}
    [data-testid="stSidebar"] {background:var(--lc-panel);border-left:1px solid var(--lc-line);direction:rtl}
    [data-testid="stSidebar"] [data-testid="stSidebarUserContent"] {padding:1.2rem 1.3rem 2rem}
    [data-testid="stWidgetLabel"], [data-testid="stCaptionContainer"], .stMarkdown {text-align:right;direction:rtl}
    [data-testid="stCaptionContainer"] {color:var(--lc-muted)}
    [data-testid="stWidgetLabel"] p, h1,h2,h3,h4 {color:var(--lc-ink)}
    [data-testid="stExpander"] {background:var(--lc-panel);border-radius:14px;border-color:var(--lc-line)}
    [data-testid="stExpanderDetails"] {padding:1.2rem}
    [data-baseweb="input"], [data-baseweb="base-input"], [data-baseweb="select"] > div, textarea {
      background:var(--lc-panel)!important;color:var(--lc-ink)!important;border-color:var(--lc-line)!important}
    input {color:var(--lc-ink)!important;text-align:right}
    [data-testid="stTextInputRootElement"], [data-testid="stNumberInputContainer"], .react-aria-ComboBox [role="group"],
    [data-testid="stMultiSelect"] [role="group"] {background:var(--lc-panel)!important;border-color:var(--lc-line)!important}
    [data-testid="stTextInputField"], input[type="number"], input[role="combobox"], textarea {
      background:var(--lc-panel)!important;color:var(--lc-ink)!important}
    [data-testid="stNumberInput"] button, .react-aria-ComboBox button {background:transparent!important;color:var(--lc-ink)!important}
    input::placeholder,textarea::placeholder {color:var(--lc-muted)!important;opacity:1}
    [data-heading-text] {color:var(--lc-ink)!important}.lc-hero [data-heading-text] {color:white!important}
    [data-testid="stCaptionContainer"] p {color:var(--lc-muted)!important}
    [data-testid="stTabs"] button[role="tab"] {color:var(--lc-muted)!important}
    [data-testid="stTabs"] button[aria-selected="true"] {color:var(--lc-accent)!important}
    [data-testid="stExpander"] summary {color:var(--lc-ink)!important}
    [data-testid="stBaseButton-secondary"], [data-testid="stBaseButton-secondaryFormSubmit"] {
      background:var(--lc-panel)!important;color:var(--lc-ink)!important;border-color:var(--lc-line)!important}
    [data-testid="stFileUploaderDropzone"] {background:var(--lc-soft);border:1px dashed var(--lc-accent);border-radius:14px}
    [data-testid="stFileUploaderDropzoneInstructions"] > div > span {font-size:0}
    [data-testid="stFileUploaderDropzoneInstructions"] > div > span:after {content:"اسحب الملفات هنا";font:600 14px Cairo;color:var(--lc-ink)}
    [data-testid="stFileUploaderDropzoneInstructions"] > div > small {font-size:0}
    [data-testid="stFileUploaderDropzoneInstructions"] > div > small:after {content:"أو اخترها من جهازك · حتى ٢٠ ميجابايت للملف";font:12px Cairo;color:var(--lc-muted)}
    [data-testid="stFileUploaderDropzone"] button:not([aria-label^="Remove"]) {font-size:0;color:var(--lc-ink)!important;background:var(--lc-panel)!important}
    [data-testid="stFileUploaderDropzone"] button:not([aria-label^="Remove"]) > * {display:none}
    [data-testid="stFileUploaderDropzone"] button:not([aria-label^="Remove"]):after {content:"اختيار الملفات";font:600 13px Cairo}
    .stButton button, .stDownloadButton button, .stLinkButton a {border-radius:10px;font-weight:700}
    .stButton button[kind="primary"] {background:#087f83;border-color:#087f83;color:#fff}
    [data-testid="stDataFrame"] {direction:ltr}
    .lc-top {display:flex;align-items:center;justify-content:space-between;gap:18px;margin:0 0 24px}
    .lc-wordmark {display:flex;align-items:center;gap:12px;font-size:23px;font-weight:800;color:var(--lc-ink);min-width:0}
    .lc-brand-lockup {display:flex;flex-direction:column;align-items:flex-start;gap:2px}
    .lc-product {white-space:nowrap;font-size:23px;line-height:1.5}
    .lc-product-label {font-size:11px;font-weight:400;color:var(--lc-muted)}
    .lc-mark {width:42px;height:42px;flex-shrink:0;border-radius:12px;background:#087f83;display:grid;place-items:center;color:white;font-size:25px}
    .lc-secure {font-size:12px;color:var(--lc-accent);background:var(--lc-soft);padding:7px 13px;border:1px solid var(--lc-line);border-radius:99px}
    .lc-hero {background:#142c42;color:white;border-radius:20px;padding:30px 34px;display:grid;grid-template-columns:1fr auto;gap:24px;align-items:center;margin-bottom:24px;overflow:hidden}
    .lc-hero h1 {font:800 clamp(25px,3.5vw,38px)/1.6 Cairo;margin:0;color:white;letter-spacing:-.6px}
    .lc-hero p {color:#c4d5e3;max-width:740px;font-size:14px;line-height:1.9;margin:7px 0 0}
    .lc-overline {font-size:11px;color:#7fe1d3;font-weight:700;letter-spacing:1px}
    .lc-brand {width:136px;min-height:85px;border-right:1px solid #36546b;padding-right:22px;text-align:center;overflow-wrap:anywhere}
    .lc-brand img {max-width:110px;max-height:60px;object-fit:contain;filter:drop-shadow(0 0 1px white)}
    .lc-brand div {font-size:12px;line-height:1.8;margin-top:8px;color:#d9e6ee}
    .lc-workflow {display:grid;grid-template-columns:repeat(3,1fr);border:1px solid var(--lc-line);border-radius:12px;background:var(--lc-panel);margin:18px 0 24px;overflow:hidden}
    .lc-step {padding:15px 18px;display:flex;gap:12px;align-items:center;border-left:1px solid var(--lc-line);font-weight:700;font-size:13px}
    .lc-step:last-child {border:0}.lc-num {color:var(--lc-accent);font-size:12px;background:var(--lc-soft);padding:5px 8px;border-radius:6px}
    .lc-step small {display:block;color:var(--lc-muted);font-weight:400;font-size:11px}
    .lc-kpis {display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:14px;margin:15px 0 20px}
    .lc-kpi {padding:20px 20px 17px;background:var(--lc-panel);border:1px solid var(--lc-line);border-radius:14px;position:relative;overflow:hidden}
    .lc-kpi:before {content:"";position:absolute;right:0;top:22px;width:3px;height:35px;background:var(--lc-accent);border-radius:3px}
    .lc-kpi small {color:var(--lc-muted);font-size:12px}.lc-kpi strong {font-size:33px;line-height:1.65;display:block;color:var(--lc-ink);font-variant-numeric:tabular-nums}
    .lc-kpi em {font-size:10px;font-style:normal;color:var(--lc-muted)}
    .lc-micro {display:flex;gap:12px;flex-wrap:wrap;color:var(--lc-muted);font-size:12px;margin:12px 0 20px}
    .lc-micro b {color:var(--lc-accent)}
    .lc-table-wrap {overflow-x:auto;border:1px solid var(--lc-line);border-radius:14px;background:var(--lc-panel);margin:14px 0}
    .lc-table {width:100%;min-width:920px;border-collapse:collapse;text-align:right;direction:rtl;font-size:12px;line-height:1.7}
    .lc-table th {padding:13px 14px;background:var(--lc-soft);color:var(--lc-muted);font-weight:700;white-space:nowrap}
    .lc-table td {padding:14px;vertical-align:top;border-top:1px solid var(--lc-line);color:var(--lc-ink);max-width:230px;overflow-wrap:anywhere}
    .lc-table tr:nth-child(even) td {background:color-mix(in srgb,var(--lc-soft) 28%,var(--lc-panel))}
    .lc-table tr:hover td {background:var(--lc-soft)}
    .lc-table strong {font-size:13px}.lc-table small {display:block;color:var(--lc-muted);font-size:10px;margin-top:4px}
    .lc-ltr {direction:ltr;unicode-bidi:isolate;display:inline-block;white-space:nowrap;font-variant-numeric:tabular-nums}
    .lc-link {display:inline-flex;align-items:center;gap:5px;background:var(--lc-soft);border:1px solid var(--lc-line);color:var(--lc-accent)!important;text-decoration:none;padding:5px 7px;border-radius:7px}
    .lc-link svg {width:14px;height:14px;flex-shrink:0}
    .lc-pill {display:inline-block;padding:2px 7px;border-radius:5px;margin:4px 0 0 4px;font-size:10px;background:var(--lc-soft);color:var(--lc-muted);white-space:nowrap}
    .lc-pill.good {background:#d9f1e9;color:#146352}.lc-pill.vip {background:#faf0ce;color:#765d0f}.lc-pill.bad {background:#fbe7e8;color:#a23443}.lc-pill.dup {background:#ede7fa;color:#67509b}
    .lc-table summary {cursor:pointer;color:var(--lc-accent);font-weight:700;white-space:nowrap}
    .lc-table details dl {max-width:230px;font-size:11px}.lc-table dt {font-weight:700;margin-top:7px}.lc-table dd {margin:0;color:var(--lc-muted)}
    .lc-bars {border:1px solid var(--lc-line);border-radius:14px;padding:20px;background:var(--lc-panel);height:100%}
    .lc-bars h3 {margin:0 0 15px;font-size:16px}.lc-bar-row {margin:14px 0}.lc-bar-label {display:flex;justify-content:space-between;color:var(--lc-muted);font-size:12px;margin-bottom:6px}
    .lc-track {height:7px;background:var(--lc-soft);border-radius:9px;overflow:hidden}.lc-fill {height:7px;background:var(--lc-accent);border-radius:9px}
    .lc-empty {padding:26px;background:var(--lc-panel);border:1px solid var(--lc-line);border-radius:14px;margin-top:15px;text-align:center;color:var(--lc-muted)}
    .lc-foot {font-size:11px;color:var(--lc-muted);padding-top:24px;margin-top:24px;border-top:1px solid var(--lc-line);line-height:1.9}
    @media(max-width:900px){[data-testid="stMainBlockContainer"]{padding:4.8rem 1.2rem 3rem}.lc-kpis{grid-template-columns:repeat(2,minmax(0,1fr))}.lc-hero{padding:24px}.lc-brand{display:none}.lc-workflow{grid-template-columns:1fr}.lc-step{border:0;border-bottom:1px solid var(--lc-line)}}
    @media(max-width:500px){.lc-top{align-items:flex-start;flex-wrap:wrap;gap:12px}.lc-product{font-size:21px}.lc-secure{max-width:125px;font-size:10px}.lc-kpi{padding:14px}.lc-kpi strong{font-size:26px}.lc-hero{padding:21px}.lc-wordmark{font-size:21px}}
    """
    css = css.replace("__FONT__", font)
    for key, value in palette.items():
        css = css.replace("__" + key + "__", value)
    st.html("<style>" + css + "</style>")


def _ui_header(brand, logo, protected):
    logo_html = ('<img alt="شعار الشركة" src="' + _export_data_url("image/png", logo) + '">') if logo else ""
    st.markdown('<div class="lc-top"><div class="lc-wordmark"><span class="lc-mark" aria-hidden="true">ن</span><span class="lc-brand-lockup"><bdi class="lc-product" dir="ltr">' + _ui_escape(PRODUCT_NAME) + '</bdi><span class="lc-product-label">منصة العملاء</span></span></div><div class="lc-secure">' +
                ("جلسة محمية · بياناتك منفصلة" if protected else "معالجة داخل جلسة مستقلة") + '</div></div>' +
                '<section class="lc-hero"><div><div class="lc-overline">من ملفات مبعثرة إلى قرارات أوضح</div>' +
                '<h1>بيانات أنظف. فرص أوضح.</h1><p>اجمع ملفات الحملات، راجع الأرقام، واعرف مَن يستحق المتابعة. تقارير عربية بهويتك، من أول عميل لآخر تفصيلة.</p></div>' +
                '<div class="lc-brand">' + logo_html + '<div>' + _ui_escape(brand) + '</div></div></section>', unsafe_allow_html=True)


def filter_leads(frame, query="", cities=None, readiness=None, contact="all", flag="all",
                 currency=None, minimum_budget=0, sources=None):
    """Literal matching only; never execute uploaded text as a regex/expression."""
    result = frame.copy()
    if query.strip():
        needle = _normal_words(query)
        columns = ["name", "phone_raw", "phone_e164", "whatsapp_raw", "whatsapp_e164", "job", "city", "details", "source"]
        text = result[columns].fillna("").astype(str).agg(" ".join, axis=1).map(_normal_words)
        result = result[text.str.contains(needle, regex=False, na=False)]
    if cities:
        result = result[result["city"].isin(cities)]
    if readiness:
        result = result[result["readiness"].isin(readiness)]
    if sources:
        result = result[result["source"].isin(sources)]
    if contact == "reachable":
        result = result[result["phone_valid"] | result["whatsapp_valid"]]
    elif contact == "unreachable":
        result = result[~(result["phone_valid"] | result["whatsapp_valid"])]
    elif contact in {"phone", "whatsapp"}:
        result = result[result[contact + "_valid"]]
    if flag in {"vip", "duplicate"}:
        result = result[result[flag]]
    elif flag == "review":
        result = result[result["issues"].str.strip().ne("")]
    if currency:
        result = result[result["budget_currency"].eq(currency)]
        if minimum_budget > 0:
            result = result[pd.to_numeric(result["budget_value"], errors="coerce").ge(minimum_budget)]
    return result.reset_index(drop=True)


def _lead_table_html(frame, allow_contacts=True):
    call_icon = '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" aria-hidden="true"><path d="M4 3h4l2 5-3 2c2 4 3 5 7 7l2-3 5 2v4c0 2-4 2-6 1C8 18 4 14 2 7 1 4 2 3 4 3Z"/></svg>'
    wa_icon = '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" aria-hidden="true"><path d="M21 11a9 9 0 0 1-13 8l-5 2 1-5A9 9 0 1 1 21 11Z"/><path d="m8 7 2 3-1 1 3 3 1-1 3 2"/></svg>'
    rows = []
    for row in frame.to_dict("records"):
        call_url = _export_contact_url(row, "phone")
        wa_url = _export_contact_url(row, "whatsapp")
        phone = _ui_escape(row["phone_display"] or _ui_short(row["phone_raw"], 22) or "غير متوفر")
        if call_url and allow_contacts:
            phone = '<a class="lc-link lc-ltr" title="اتصال يدوي" href="' + _ui_escape(call_url) + '">' + call_icon + phone + '</a>'
        else:
            phone = '<span class="lc-ltr">' + phone + '</span>'
        whatsapp = ('<a class="lc-link" target="_blank" rel="noopener noreferrer" title="فتح واتساب يدوياً" href="' + _ui_escape(wa_url) + '">' + wa_icon + 'واتساب</a>') if wa_url and allow_contacts else ("صالح تنسيقياً" if wa_url else "غير متوفر")
        whatsapp += '<small>' + _ui_escape(row["whatsapp_origin"]) + '</small>'
        pills = ('<span class="lc-pill vip">مميز</span>' if row["vip"] else "")
        pills += ('<span class="lc-pill dup">تكرار محتمل</span>' if row["duplicate"] else "")
        pills += ('<span class="lc-pill bad">مراجعة الهاتف</span>' if not row["phone_valid"] else "")
        extras = json.loads(row["details"])
        entries = ''.join('<dt>' + _ui_escape(_ui_short(key, 60)) + '</dt><dd>' + _ui_escape(_ui_short(value, 160)) + '</dd>' for key, value in list(extras.items())[:8])
        details = '<details><summary>عرض التفاصيل</summary><dl>' + entries + '</dl><small>' + _ui_escape(_ui_short(row["issues"], 220)) + '</small><small>القيم الكاملة في مستكشف السجل أدناه.</small></details>'
        source = _ui_escape(_ui_short(row["source"], 25)) + ' · صف ' + str(int(row["source_row"]))
        ready_class = "good" if row["readiness"] == "فوري" else ""
        rows.append('<tr><td><strong>' + _ui_escape(_ui_short(row["name"], 70)) + '</strong><small>' + _ui_escape(_ui_short(row["job"], 35)) + '</small>' + pills + '</td><td>' + phone + '</td><td>' + whatsapp + '</td><td><span class="lc-ltr">' + _ui_escape(_export_budget_label(row)) + '</span></td><td>' + _ui_escape(_ui_short(row["city"], 40)) + '</td><td><span class="lc-pill ' + ready_class + '">' + _ui_escape(row["readiness"]) + '</span></td><td>' + details + '</td><td><small>' + source + '</small></td></tr>')
    return '<div class="lc-table-wrap"><table class="lc-table"><thead><tr><th>العميل</th><th>الهاتف</th><th>واتساب</th><th>الميزانية</th><th>المدينة</th><th>الجاهزية</th><th>تفاصيل</th><th>المصدر</th></tr></thead><tbody>' + ''.join(rows) + '</tbody></table></div>'


def _bar_panel(title, counts):
    maximum = max([value for _, value in counts] or [1]) or 1
    bars = []
    for label, value in counts:
        width = max(1, min(100, 100 * value / maximum))
        bars.append('<div class="lc-bar-row"><div class="lc-bar-label"><span>' + _ui_escape(_ui_short(label, 40)) + '</span><b>' + f'{value:,}' + '</b></div><div class="lc-track"><div class="lc-fill" style="width:' + f'{width:.2f}' + '%"></div></div></div>')
    st.markdown('<section class="lc-bars"><h3>' + _ui_escape(title) + '</h3>' + ''.join(bars) + '</section>', unsafe_allow_html=True)


def _show_metrics(frame, total):
    k = calculate_kpis(frame)
    cards = [(f'{k["total"]:,}', "العملاء في العرض", f"من إجمالي {total:,} سجلاً"),
             (f'{k["reachable_count"]:,}', "قابلون للتواصل", "هاتف أو واتساب صالح تنسيقياً"),
             (f'{k["vip_count"]:,}', "عملاء مميزون", "بحسب العملة والحد المختارين"),
             (f'{k["ready_rate"]:.1f}٪', "جاهزون الآن", f'{k["ready_count"]:,} سجلاً بعبارات جاهزية واضحة')]
    st.markdown('<div class="lc-kpis">' + ''.join('<div class="lc-kpi"><small>' + label + '</small><strong>' + value + '</strong><em>' + note + '</em></div>' for value, label, note in cards) + '</div><div class="lc-micro"><span>قناة التواصل الأوفر: <b>' + _ui_escape(k["top_contact"]) + '</b></span><span>·</span><span>هواتف صالحة: <b>' + str(k["valid_phone"]) + '</b></span><span>·</span><span>واتساب صالح تنسيقياً: <b>' + str(k["valid_wa"]) + '</b></span><span>·</span><span>تكرارات محتملة: <b>' + str(k["duplicates"]) + '</b></span></div>', unsafe_allow_html=True)


def _source_controls(uploads, epoch):
    frames, signatures, audit = [], [], []
    blocked = False
    active_cache_keys = set()
    cache = st.session_state.setdefault("lc_parse_cache", {})
    sheet_cache = st.session_state.setdefault("lc_sheet_cache", {})
    total_cells = 0
    for file_index, (filename, content) in enumerate(uploads):
        file_hash = hashlib.sha256(filename.encode("utf-8", errors="surrogatepass") + content).hexdigest()[:20]
        file_key = f"lc_{epoch}_{file_index}_{file_hash}"
        with st.expander(f"مصدر {file_index + 1} · {_ui_short(filename, 100)}", expanded=len(uploads) == 1):
            try:
                if file_hash not in sheet_cache:
                    sheet_cache[file_hash] = workbook_sheets(content, filename)
                sheets = sheet_cache[file_hash]
                if sheets == [""]:
                    selected = [""]
                    c1, c2 = st.columns(2)
                    encoding = c1.selectbox("ترميز CSV", list(ENCODING_OPTIONS), format_func=ENCODING_OPTIONS.get, key=file_key + "_enc")
                    delimiter = c2.selectbox("فاصل الأعمدة", list(DELIMITER_OPTIONS), format_func=DELIMITER_OPTIONS.get, key=file_key + "_sep")
                else:
                    selected = st.multiselect("الأوراق المراد معالجتها", sheets, default=sheets[:1], key=file_key + "_sheets")
                    if not selected:
                        st.warning("اختر ورقة واحدة على الأقل من هذا الملف.")
                        blocked = True
                    st.caption("تُعالج الأوراق المختارة فقط؛ كل أوراق المصنف تظل محفوظة داخل أرشيف الملف الأصلي.")
                    encoding = delimiter = "auto"
                for sheet in selected:
                    parse_key = hashlib.sha256(json.dumps([file_hash, sheet, encoding, delimiter], ensure_ascii=False).encode()).hexdigest()[:24]
                    active_cache_keys.add(parse_key)
                    if parse_key not in cache:
                        cache[parse_key] = _ui_asdict(read_source(content, filename, sheet, encoding, delimiter))
                    # Script-defined class identity changes on a Streamlit rerun.
                    # Store only builtin structures and rehydrate for this run.
                    source = ParsedTable(**cache[parse_key])
                    if sheet:
                        st.markdown("**الورقة:** " + _ui_escape(sheet))
                    if not source.rows:
                        st.error("المصدر المختار فارغ؛ اختر ورقة أخرى أو احذف الملف الفارغ من الرفع.")
                        blocked = True
                        continue
                    key = file_key + "_" + parse_key
                    detected = source.header_row + 1
                    header_one = st.number_input("صف العناوين — اكتب 0 إذا لم توجد عناوين", min_value=0, max_value=len(source.rows), value=detected, step=1, key=key + "_header")
                    frame = table_from_source(source, int(header_one) - 1)
                    total_cells += len(frame) * len(frame.columns)
                    for note in frame.attrs.get("warnings", []):
                        st.caption(note)
                    c1, c2 = st.columns([1, 3])
                    c1.caption(f"{len(frame):,} صفاً · {len(frame.columns):,} عموداً")
                    c2.caption(f"الترميز: {source.encoding or 'Excel'} · صف العناوين: {header_one or 'لا يوجد'}")
                    st.dataframe(frame.head(6), hide_index=True, width="stretch", height="auto")
                    with st.expander("فحص الصفوف قبل اختيار العناوين"):
                        first_rows = pd.DataFrame(source.rows[:min(15, len(source.rows))])
                        first_rows.columns = [f"عمود {i + 1}" for i in range(len(first_rows.columns))]
                        first_rows.index = range(1, len(first_rows) + 1)
                        st.dataframe(first_rows, width="stretch")
                        st.caption("أرقام هذه المعاينة تبدأ من أول صف في المصدر، بما فيه صفوف المعلومات السابقة للعناوين.")
                    guessed = guess_mapping(frame)
                    mapping = {}
                    mapping_columns = st.columns(3)
                    options = [None] + list(frame.columns)
                    for index, (field, label) in enumerate(FIELD_LABELS.items()):
                        with mapping_columns[index % 3]:
                            mapping[field] = st.selectbox(label, options, index=options.index(guessed[field]),
                                                          format_func=lambda value: "غير محدد" if value is None else str(value),
                                                          key=key + f"_{header_one}_map_{field}")
                    assigned = [value for value in mapping.values() if value is not None]
                    if len(set(assigned)) != len(assigned):
                        st.error("تم اختيار العمود نفسه لأكثر من حقل. اجعل لكل حقل عموداً مستقلاً؛ أو فعّل خيار استخدام الهاتف لواتساب.")
                        blocked = True
                    if mapping["phone"] is None and mapping["whatsapp"] is None:
                        st.warning("لم تُربط أعمدة تواصل. يمكنك الاحتفاظ بالبيانات، لكن لن تُنشأ روابط اتصال إلا بعد ربط الأعمدة.")
                    if frame.empty:
                        st.error("لا توجد صفوف بيانات بعد العناوين المختارة. راجع رقم الصف.")
                        blocked = True
                    label = f"{file_index + 1} · {_ui_short(filename, 120)}" + (f" / {_ui_short(sheet, 80)}" if sheet else "")
                    frames.append((frame, mapping, label))
                    signatures.append([source.source_id, header_one, mapping, label])
                    audit.append({"المصدر": label, "الصفوف": len(frame), "صف العناوين": int(header_one),
                                  "حقول إضافية": len(frame.columns) - len(set(assigned)),
                                  "الترميز": source.encoding or "Excel", "خريطة الحقول": mapping})
            except UserDataError as exc:
                st.error(str(exc))
                blocked = True
            except Exception:
                st.error("تعذّرت قراءة هذا المصدر بأمان. أعد حفظه كملف CSV أو Excel سليم ثم جرّب مجدداً؛ لم تُعالج بقية الملفات جزئياً.")
                blocked = True
    for key in list(cache):
        if key not in active_cache_keys:
            del cache[key]
    active_files = {hashlib.sha256(name.encode("utf-8", errors="surrogatepass") + data).hexdigest()[:20] for name, data in uploads}
    for key in list(sheet_cache):
        if key not in active_files:
            del sheet_cache[key]
    total_rows = sum(len(frame) for frame, _, _ in frames)
    if total_rows > MAX_ROWS or total_cells > MAX_CELLS:
        st.error(f"المجموع يتجاوز حد الجلسة: {MAX_ROWS:,} صفاً أو {MAX_CELLS:,} خلية. اختر ملفات أو أوراقاً أقل؛ لا يحدث اختصار تلقائي.")
        blocked = True
    return frames, signatures, audit, blocked


def _preview_results(frame, demo):
    if frame.empty:
        st.info("لا توجد نتائج مطابقة لهذه المرشحات. غيّر البحث أو المدينة أو الحالة.")
        return
    c1, c2, c3 = st.columns([2, 1, 1])
    order = c1.selectbox("ترتيب العرض", ["الأولوية", "ترتيب المصدر", "الاسم"], key="lc_sort")
    size = c2.selectbox("سجلات الصفحة", [10, 25, 50, 100], index=1, key="lc_page_size")
    if order == "الأولوية":
        ranked = frame.assign(_priority=frame["vip"].astype(int) * 2 + frame["readiness"].eq("فوري").astype(int))
        ordered = ranked.sort_values("_priority", ascending=False, kind="stable").drop(columns="_priority").reset_index(drop=True)
    elif order == "الاسم":
        ordered = frame.sort_values("name", kind="stable").reset_index(drop=True)
    else:
        ordered = frame
    page_count = max(1, math.ceil(len(ordered) / size))
    page = c3.selectbox("الصفحة", list(range(1, page_count + 1)), key=f"lc_page_{page_count}_{size}")
    start = (page - 1) * size
    page_frame = ordered.iloc[start:start + size]
    st.caption(f"عرض {start + 1:,}–{start + len(page_frame):,} من {len(ordered):,}. اختيار الصفحة لا يختصر نطاق Excel أو PDF؛ لهما إعداد مستقل.")
    st.markdown(_lead_table_html(page_frame, allow_contacts=not demo), unsafe_allow_html=True)
    with st.expander("مستكشف السجل · جميع القيم الأصلية دون اختصار"):
        selected = st.selectbox("اختر سجلاً من الصفحة الحالية", list(range(len(page_frame))),
                                format_func=lambda i: f'{_ui_short(page_frame.iloc[i]["name"], 65)} · صف {page_frame.iloc[i]["source_row"]}',
                                key=f"lc_detail_{page}_{len(page_frame)}")
        row = page_frame.iloc[selected]
        st.caption(f'{row["source"]} · الصف {row["source_row"]} · المعرّف {row["lead_id"]}')
        if row["issues"]:
            st.info(row["issues"])
        c1, c2 = st.columns(2)
        c1.markdown("**حقول إضافية**")
        c1.json(json.loads(row["details"]), expanded=True)
        c2.markdown("**كل الحقول كما قُرئت**")
        c2.json(json.loads(row["original"]), expanded=True)


def _export_controls(full, filtered, uploads, brand, logo, demo):
    scope = st.radio("نطاق التصدير", ["النتائج بعد التصفية", "كل البيانات المنظّفة"], horizontal=True, key="lc_export_scope")
    target = filtered if scope == "النتائج بعد التصفية" else full
    st.caption(f"النطاق المحدد: {len(target):,} سجلاً. المؤشرات داخل كل ملف تُحسب على سجلاته فقط. التكرار محسوب على مجموعة الاستيراد الأصلية.")
    if target.empty:
        st.info("النطاق فارغ. غيّر المرشحات أو اختر كل البيانات.")
        return
    select_signature = hashlib.sha256((st.session_state.get("lc_result_sig", "") + json.dumps(target["lead_id"].tolist()) + brand).encode() + (logo or b"")).hexdigest()[:24]
    if st.session_state.get("lc_export_sig") != select_signature:
        for key in ["lc_excel_bytes", "lc_pdf_bytes", "lc_pdf_html", "lc_pdf_sig"]:
            st.session_state.pop(key, None)
        st.session_state["lc_export_sig"] = select_signature
    left, right = st.columns(2)
    with left:
        st.markdown("### ملف Excel المنسّق")
        st.caption("ملخص تنفيذي + بيانات قابلة للتصفية + الخصائص الأصلية كاملة. اتجاه عربي وروابط قابلة للنقر.")
        if st.button("تجهيز Excel", key="lc_build_excel", icon=":material/table_view:", width="stretch"):
            try:
                with st.spinner("تجهيز ملف Excel داخل الجلسة…"):
                    st.session_state["lc_excel_bytes"] = build_excel(target, brand, logo)
            except UserDataError as exc:
                st.error(str(exc))
            except Exception:
                st.error("تعذّر تجهيز Excel ضمن موارد التشغيل المتاحة. صدّر نطاقاً أصغر أو نزّل أرشيف المصادر.")
        if st.session_state.get("lc_excel_bytes"):
            st.download_button("تنزيل Excel", st.session_state["lc_excel_bytes"], f"{PRODUCT_SLUG}-leads.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="lc_download_excel", width="stretch")
        st.caption("معادلات الملخص يعيد Excel حسابها عند الفتح. لا ينفّذ التطبيق أي صيغ قادمة من ملفك.")
    with right:
        st.markdown("### تقرير PDF للطباعة")
        batch_size = st.selectbox("حجم دفعة PDF", [25, 50, 100, 250, PDF_BATCH_SIZE], index=2, key="lc_pdf_batch_size")
        batch_count = math.ceil(len(target) / batch_size)
        batch = st.selectbox("دفعة التقرير", list(range(1, batch_count + 1)),
                             format_func=lambda i: f"دفعة {i} من {batch_count}", key=f"lc_pdf_batch_{batch_count}_{batch_size}")
        offset = (batch - 1) * batch_size
        part = target.iloc[offset:offset + batch_size]
        note = f"{scope} · دفعة {batch} من {batch_count} · السجلات {offset + 1}–{offset + len(part)} من {len(target)}. ترتيب المصدر، وليس ترتيب جدول المعاينة."
        if demo:
            note += " بيانات اصطناعية للتجربة فقط؛ لا تستخدم الأرقام للتواصل."
        pdf_sig = f"{select_signature}_{batch}_{batch_size}"
        if st.session_state.get("lc_pdf_sig") != pdf_sig:
            st.session_state.pop("lc_pdf_bytes", None)
            st.session_state.pop("lc_pdf_html", None)
        st.caption(note)
        if st.button("تجهيز PDF", key="lc_build_pdf", icon=":material/picture_as_pdf:", width="stretch"):
            try:
                with st.spinner("تنسيق الصفحات العربية وتضمين الخط…"):
                    html = build_report_html(part, brand, logo, context_note=note)
                    pdf = build_pdf(part, brand, logo, context_note=note)
                    st.session_state.update({"lc_pdf_bytes": pdf, "lc_pdf_html": html, "lc_pdf_sig": pdf_sig})
            except UserDataError as exc:
                st.error(str(exc))
            except Exception:
                st.error("تعذّر تجهيز PDF. تحقّق من مكونات الطباعة أو اختر دفعة أصغر. ملف Excel يظل متاحاً.")
        if st.session_state.get("lc_pdf_bytes") and st.session_state.get("lc_pdf_sig") == pdf_sig:
            st.download_button("تنزيل PDF", st.session_state["lc_pdf_bytes"], f"{PRODUCT_SLUG}-report-{batch:03}.pdf", "application/pdf", key="lc_download_pdf", width="stretch")
            st.download_button("تنزيل نسخة HTML للطباعة", st.session_state["lc_pdf_html"].encode("utf-8"),
                               f"{PRODUCT_SLUG}-report-{batch:03}.html", "text/html;charset=utf-8", key="lc_download_html", width="stretch")
    st.divider()
    st.markdown("#### نسخة المصادر دون أي تغيير")
    st.caption("هذا الأرشيف يشمل كل الملفات المرفوعة، بكل الأوراق وصفوف المعلومات السابقة للعناوين. لا يتبع المرشحات الحالية. قد تحتوي الملفات الأصلية على صيغ؛ افتحها فقط إذا وثقت بالمصدر.")
    if st.button("تجهيز أرشيف المصادر الأصلية", key="lc_build_sources", icon=":material/folder_zip:"):
        st.session_state["lc_sources_zip"] = build_sources_archive(uploads)
    if st.session_state.get("lc_sources_zip"):
        st.download_button("تنزيل المصادر ZIP", st.session_state["lc_sources_zip"], f"{PRODUCT_SLUG}-original-sources.zip", "application/zip", key="lc_download_sources")


def run_app():
    st.set_page_config(page_title=APP_TITLE, page_icon=":material/filter_alt:", layout="wide", initial_sidebar_state="auto")
    epoch = int(st.session_state.setdefault("lc_epoch", 0))
    protected = _require_auth()
    with st.sidebar:
        st.markdown("## إعدادات مساحة العمل")
        theme = st.selectbox("مظهر الواجهة", ["فاتح", "داكن"], key="lc_theme")
        brand = st.text_input("اسم الشركة أو الوكالة", value="وكالتك", max_chars=120, key="lc_brand").strip() or "تقرير العملاء"
        logo_upload = st.file_uploader("شعار الشركة — حتى ٢ ميجابايت", type=["png", "jpg", "jpeg", "webp"], max_upload_size=2, key=f"lc_logo_{epoch}")
        logo = None
        logo_error = False
        if logo_upload:
            try:
                logo = sanitize_logo(logo_upload.getvalue())
            except UserDataError as exc:
                st.error(str(exc))
                logo_error = True
        st.divider()
        region = st.selectbox("دولة الأرقام المحلية", list(COUNTRY_OPTIONS), format_func=COUNTRY_OPTIONS.get, key="lc_region")
        currency = st.selectbox("العملة الافتراضية وحد المميز", CURRENCY_OPTIONS, key="lc_currency")
        threshold = st.number_input("الحد الأدنى لميزانية العميل المميز", min_value=0.0, max_value=1_000_000_000_000.0, value=1_000_000.0, step=100_000.0, key="lc_vip_threshold")
        fallback = st.checkbox("استخدام الهاتف لواتساب إذا كانت خانته فارغة", value=False, key="lc_wa_fallback")
        st.caption("يظهر هذا الرقم كافتراض واضح؛ لا يُستبدل به رقم واتساب صريح غير صالح. لا نتحقق من نشاط الخط أو وجود حساب واتساب.")
        st.caption("تُحترم بادئة الدولة المكتوبة. الأرقام المحلية الملتبسة تُعلّم للمراجعة. العملات المختلفة لا تُحوّل ولا تُجمع معاً.")
        st.divider()
        st.button("مسح بيانات الجلسة", on_click=_clear_workspace, icon=":material/delete_sweep:", width="stretch")
        st.caption("يزيل الملفات والنتائج من حالة هذه الجلسة. لا يحذف النسخ التي سبق تنزيلها على جهازك.")
    try:
        _ui_style(theme == "داكن")
    except UserDataError as exc:
        st.error(str(exc))
        st.stop()
    _ui_header(brand, logo, protected)
    if not protected:
        st.info("وضع العرض المفتوح: استخدم بيانات اصطناعية على الروابط العامة. قبل رفع بيانات حقيقية، فعّل OIDC وقائمة السماح أو قيّد المشاهدين لدى الاستضافة.")
    mode = st.radio("مصدر البيانات", ["رفع ملفات", "تجربة جاهزة"], horizontal=True, key="lc_mode")
    demo = mode == "تجربة جاهزة"
    if demo:
        uploads = [("synthetic-leads.csv", demo_csv())]
        st.warning("هذه بيانات اصطناعية وليست عملاء حقيقيين. روابط التواصل معطّلة في المعاينة التجريبية؛ لا تستخدم أرقام التقارير للتواصل.")
    else:
        files = st.file_uploader("ملفات العملاء من Meta أو TikTok أو Google Ads أو Excel", type=["xlsx", "xls", "csv"],
                                 accept_multiple_files=True, key=f"lc_files_{epoch}")
        if len(files) > MAX_FILES:
            st.error(f"يمكن معالجة {MAX_FILES} ملفات كحد أقصى في المرة الواحدة. احذف الملفات الإضافية من قائمة الرفع.")
            st.stop()
        uploads = [(file.name, file.getvalue()) for file in files]
    st.markdown('<div class="lc-workflow"><div class="lc-step"><span class="lc-num">01</span><div>ارفع ملفاتك<small>Excel أو CSV بأي ترتيب أعمدة</small></div></div><div class="lc-step"><span class="lc-num">02</span><div>راجع الربط<small>كشف تلقائي مع تعديل يدوي</small></div></div><div class="lc-step"><span class="lc-num">03</span><div>حلّل وصدّر<small>لوحة مؤشرات وتقارير بهويتك</small></div></div></div>', unsafe_allow_html=True)
    if not uploads:
        for key in ["lc_result", "lc_result_sig", "lc_excel_bytes", "lc_pdf_bytes", "lc_pdf_html", "lc_sources_zip", "lc_parse_cache", "lc_sheet_cache"]:
            st.session_state.pop(key, None)
        st.markdown('<div class="lc-empty"><strong>ابدأ بملف واحد. كل تفصيلة تفضل محفوظة.</strong><p>اختار «تجربة جاهزة» فوق لاستكشاف الأدوات ببيانات اصطناعية، أو ارفع ملفات التصدير من حملاتك.</p></div>', unsafe_allow_html=True)
        st.download_button("تنزيل CSV تجريبي", demo_csv(), "synthetic-leads.csv", "text/csv", key="lc_demo_download")
        return
    uploads_sig = hashlib.sha256(b"".join(hashlib.sha256(name.encode() + data).digest() for name, data in uploads)).hexdigest()
    if st.session_state.get("lc_uploads_sig") != uploads_sig:
        st.session_state.pop("lc_sources_zip", None)
        st.session_state["lc_uploads_sig"] = uploads_sig
    st.markdown("### مراجعة الاستيراد وربط الأعمدة")
    st.caption("الاقتراحات مساعدة وليست قراراً نهائياً. كل عمود غير مربوط يدخل في التفاصيل؛ الصفوف الفارغة فقط تُستبعد من قائمة العملاء.")
    frames, signatures, audit, blocked = _source_controls(uploads, epoch)
    signature = hashlib.sha256(json.dumps([signatures, region, currency, threshold, fallback], ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    if st.session_state.get("lc_result_sig") not in (None, signature):
        for key in ["lc_result", "lc_result_sig", "lc_excel_bytes", "lc_pdf_bytes", "lc_pdf_html"]:
            st.session_state.pop(key, None)
        st.info("تغيّرت إعدادات المصدر أو التنظيف. شغّل المعالجة مجدداً لتحديث النتائج والتقارير.")
    if st.button("تنظيف البيانات وبناء لوحة المؤشرات", type="primary", width="stretch", icon=":material/auto_fix_high:",
                 disabled=blocked or logo_error or not frames, key="lc_process"):
        try:
            with st.spinner("توحيد الأرقام والميزانيات ومراجعة التكرارات…"):
                results = [clean_leads(frame, mapping, source, region, threshold, currency, fallback) for frame, mapping, source in frames]
                for result in results:
                    result.attrs = {}
                merged = mark_duplicates(pd.concat(results, ignore_index=True))
                st.session_state.update({"lc_result": merged, "lc_result_sig": signature, "lc_audit": audit})
                for key in ["lc_excel_bytes", "lc_pdf_bytes", "lc_pdf_html", "lc_export_sig"]:
                    st.session_state.pop(key, None)
        except UserDataError as exc:
            st.error(str(exc))
        except Exception:
            st.error("تعذّرت معالجة المجموعة كاملة. لم تُعرض نتيجة جزئية؛ راجع الملفات وربط الأعمدة ثم جرّب مجدداً.")
    if blocked or logo_error or "lc_result" not in st.session_state:
        return
    full = st.session_state["lc_result"]
    st.divider()
    st.markdown("### نظرة أوضح على فرصك")
    c1, c2, c3 = st.columns([2, 1, 1])
    query = c1.text_input("بحث بالاسم أو الرقم أو أي تفصيلة", key="lc_search", max_chars=150)
    cities = c2.multiselect("المدينة", sorted(value for value in full["city"].unique() if value), placeholder="كل المدن", key="lc_cities")
    ready = c3.multiselect("الجاهزية", ["فوري", "لاحقًا", "غير محدد"], placeholder="كل الحالات", key="lc_ready")
    with st.expander("مرشحات إضافية"):
        c1, c2, c3 = st.columns(3)
        contact = c1.selectbox("إمكانية التواصل", list(CONTACT_FILTERS), format_func=CONTACT_FILTERS.get, key="lc_contact")
        flag = c2.selectbox("حالة السجل", list(FLAG_FILTERS), format_func=FLAG_FILTERS.get, key="lc_flag")
        sources = c3.multiselect("المصدر", list(full["source"].unique()), key="lc_sources")
        c1, c2 = st.columns(2)
        budget_currency = c1.selectbox("تصفية الميزانية بعملة واحدة", [None] + CURRENCY_OPTIONS,
                                       format_func=lambda value: "دون تصفية بالميزانية" if value is None else value, key="lc_filter_currency")
        min_budget = c2.number_input("الحد الأدنى للميزانية", min_value=0.0, max_value=1_000_000_000_000.0,
                                     value=0.0, step=100_000.0, disabled=budget_currency is None, key="lc_min_budget")
    filtered = filter_leads(full, query, cities, ready, contact, flag, budget_currency, min_budget, sources)
    _show_metrics(filtered, len(full))
    table_tab, insights_tab, export_tab, audit_tab = st.tabs(["العملاء والتواصل", "قراءة المؤشرات", "التصدير والتقارير", "سجل المراجعة"])
    with table_tab:
        _preview_results(filtered, demo)
    with insights_tab:
        left, right = st.columns(2)
        with left:
            counts = [(str(label) or "غير محددة", int(value)) for label, value in filtered["city"].value_counts().head(8).items()]
            _bar_panel("أكثر المدن حضوراً", counts)
        with right:
            counts = [(label, int(filtered["readiness"].eq(label).sum())) for label in ["فوري", "لاحقًا", "غير محدد"]]
            _bar_panel("توزيع جاهزية الشراء", counts)
        st.caption("تعتمد الجاهزية على عبارات محددة مع مراعاة النفي؛ الحالات غير الواضحة تبقى غير محددة. المؤشرات تشمل التكرارات لأنها لم تُحذف.")
        st.markdown("#### الميزانيات حسب العملة — دون تحويل")
        amounts = filtered.dropna(subset=["budget_value"])
        if amounts.empty:
            st.info("لا توجد ميزانيات رقمية واضحة في هذا النطاق.")
        else:
            grouped = amounts.groupby("budget_currency")["budget_value"].agg(["count", "max", "median"]).reset_index()
            grouped.columns = ["العملة", "سجلات بميزانية", "أعلى ميزانية", "الوسيط"]
            st.dataframe(grouped, hide_index=True, width="stretch")
    with export_tab:
        _export_controls(full, filtered, uploads, brand, logo, demo)
    with audit_tab:
        st.markdown("#### قرارات الاستيراد")
        for item in st.session_state.get("lc_audit", []):
            st.json(item, expanded=False)
        st.markdown("#### سجل ملاحظات التنظيف")
        review = filtered.loc[filtered["issues"].ne(""), ["name", "source", "source_row", "issues"]].copy()
        review.columns = ["الاسم", "المصدر", "صف المصدر", "ملاحظات"]
        st.dataframe(review, hide_index=True, width="stretch")
        st.caption("لا تُحذف التكرارات تلقائياً. أرقام المصدر تشير إلى صفوف الملف الأصلية، وليس ترتيب النتائج بعد التصفية. صيغ Excel لا تُنفّذ، والميزانيات الملتبسة لا تُخمّن.")
    st.markdown('<footer class="lc-foot"><bdi dir="ltr">' + _ui_escape(PRODUCT_NAME) + '</bdi> · الملفات تبقى في ذاكرة الجلسة ولا تُستخدم للتدريب داخل هذا التطبيق ولا تُرسل لخدمة تحليل خارجية. انتهاء الجلسة أو إعادة التشغيل قد يفقد البيانات؛ نزّل النتائج التي تحتاجها. روابط الاتصال وواتساب إجراءات يدوية فقط. لا توجد قاعدة بيانات دائمة أو نظام فوترة في هذه النسخة.</footer>', unsafe_allow_html=True)


if __name__ == "__main__":
    run_app()
