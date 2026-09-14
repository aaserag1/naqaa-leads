"""Export regressions. Fixtures are synthetic and contain no customer data."""
import base64
import csv
import importlib.util
import io
import json
from pathlib import Path
import re
import sys
import unittest
from unittest import mock
import zipfile

import openpyxl
import pandas as pd
from PIL import Image, PngImagePlugin

import app


def export_frame(rows=2):
    records = []
    for index in range(rows):
        record = {key: "" for key in app.EXPORT_COLUMNS}
        record.update({
            "lead_id": f"SYN-{index + 1:04}",
            "source": "مصدر اصطناعي.csv", "source_row": str(index + 2),
            "name": "عميل اصطناعي", "phone_raw": "+201000000001",
            "phone_e164": "+201000000001", "phone_display": "+20 100 000 0001",
            "phone_valid": True, "phone_region": "EG", "phone_error": "",
            "tel_url": "tel:+201000000001", "whatsapp_raw": "+966500000001",
            "whatsapp_e164": "+966500000001", "whatsapp_display": "+966 50 000 0001",
            "whatsapp_valid": True, "whatsapp_error": "", "wa_url": "https://wa.me/966500000001",
            "whatsapp_origin": "رقم مستقل", "job": "طبيب", "budget_raw": "150000 EGP",
            "budget_value": 150000.0, "budget_currency": "EGP", "vip": index == 0,
            "timeline": "جاهز فوراً" if index == 0 else "غير جاهز الآن",
            "readiness": "فوري" if index == 0 else "لاحقاً", "city": "القاهرة",
            "details": json.dumps({"حقل إضافي": "قيمة اصطناعية"}, ensure_ascii=False),
            "original": json.dumps({"اسم العميل": "عميل اصطناعي", "هاتف": "+201000000001", "حقل إضافي": "قيمة اصطناعية"}, ensure_ascii=False),
            "issues": "", "duplicate": index == 1,
        })
        records.append(record)
    return pd.DataFrame(records, columns=app.EXPORT_COLUMNS)


def png_bytes(size=(64, 32), metadata=False):
    stream = io.BytesIO()
    image = Image.new("RGB", size, "#087f83")
    info = PngImagePlugin.PngInfo()
    if metadata:
        info.add_text("Description", "synthetic private metadata")
    image.save(stream, format="PNG", pnginfo=info)
    return stream.getvalue()


class LogoTests(unittest.TestCase):
    def test_logo_thumbnail_metadata_and_format(self):
        cleaned = app.sanitize_logo(png_bytes((1200, 600), metadata=True))
        with Image.open(io.BytesIO(cleaned)) as image:
            self.assertEqual(image.format, "PNG")
            self.assertLessEqual(image.width, 512)
            self.assertLessEqual(image.height, 256)
            self.assertEqual(image.n_frames, 1)
            self.assertFalse(image.getexif())
            self.assertNotIn("Description", image.info)
            self.assertNotIn("exif", image.info)
        app.safe_pdf_url_fetcher("data:image/png;base64," + base64.b64encode(cleaned).decode())

    def test_jpeg_and_webp_reencoded(self):
        for format_name in ("JPEG", "WEBP"):
            with self.subTest(format_name=format_name):
                stream = io.BytesIO()
                Image.new("RGB", (80, 45), "navy").save(stream, format=format_name)
                self.assertTrue(app.sanitize_logo(stream.getvalue()).startswith(b"\x89PNG"))

    def test_svg_invalid_oversized_and_many_pixels_rejected(self):
        inputs = [b'<svg xmlns="http://www.w3.org/2000/svg"/>', b"https://example.invalid/logo.png",
                  b"not an image", b"x" * (2 * 1024 * 1024 + 1), png_bytes((4001, 3000))]
        for data in inputs:
            with self.subTest(length=len(data)):
                with self.assertRaises(app.UserDataError):
                    app.sanitize_logo(data)

    def test_animation_rejected(self):
        stream = io.BytesIO()
        image = Image.new("RGB", (20, 20), "navy")
        image.save(stream, format="PNG", save_all=True,
                   append_images=[Image.new("RGB", (20, 20), "teal")], duration=100, loop=0)
        with self.assertRaises(app.UserDataError):
            app.sanitize_logo(stream.getvalue())


class ExcelExportTests(unittest.TestCase):
    def test_literals_styles_rtl_hyperlinks_and_formulas(self):
        frame = export_frame()
        frame.loc[0, "name"] = '=HYPERLINK("https://attacker.invalid","اختبار")'
        frame.loc[0, "job"] = "@SUM(A1:A10)"
        frame.loc[0, "source"] = "+ملف اصطناعي.csv"
        frame.loc[0, "phone_raw"] = "+201000000001"
        frame.loc[1, "name"] = "-SUM(1,2)"
        frame.loc[1, "tel_url"] = 'javascript:alert("bad")'
        frame.loc[1, "wa_url"] = "https://attacker.invalid/"
        frame.loc[1, "phone_valid"] = False
        frame.loc[1, "whatsapp_valid"] = False
        book = openpyxl.load_workbook(io.BytesIO(app.build_excel(frame, "=شركة اصطناعية", png_bytes())), data_only=False)
        self.assertEqual(book.sheetnames, ["الملخص", "العملاء", "الخصائص الأصلية"])
        leads, summary = book["العملاء"], book["الملخص"]
        self.assertEqual(summary["A1"].value, "=شركة اصطناعية")
        self.assertEqual(summary["A1"].data_type, "s")
        for sheet in book:
            self.assertTrue(sheet.sheet_view.rightToLeft)
            self.assertFalse(sheet.sheet_view.showGridLines)
        self.assertEqual(leads.freeze_panes, "E8")
        self.assertTrue(leads.tables)
        self.assertTrue(leads.auto_filter.ref)
        self.assertEqual(leads["D8"].value, frame.loc[0, "name"])
        self.assertEqual(leads["D8"].data_type, "s")
        self.assertEqual(leads["S8"].value, "@SUM(A1:A10)")
        self.assertEqual(leads["E8"].value, "+201000000001")
        self.assertEqual(leads["E8"].number_format, "@")
        self.assertIsInstance(leads["U8"].value, (float, int))
        self.assertEqual(leads["K8"].hyperlink.target, "tel:+201000000001")
        self.assertEqual(leads["Q8"].hyperlink.target, "https://wa.me/966500000001")
        self.assertIsNone(leads["K9"].hyperlink)
        self.assertIsNone(leads["Q9"].hyperlink)
        self.assertEqual(leads["G8"].hyperlink.target, "tel:+201000000001")
        self.assertEqual(leads["N8"].hyperlink.target, "https://wa.me/966500000001")
        self.assertEqual(summary["B10"].value, "=IFERROR(B9/B5,0)")
        self.assertIn("$H$8:$H$9", summary["B6"].value)
        self.assertIn("$O$8:$O$9", summary["B7"].value)
        self.assertIn("COUNTIFS", summary["B11"].value)
        self.assertTrue(book.calculation.fullCalcOnLoad)
        self.assertEqual(book.calculation.calcMode, "auto")
        self.assertEqual(len(summary._images), 1)
        for sheet in (leads, book["الخصائص الأصلية"]):
            self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))
        for row in leads:
            for cell in row:
                if cell.hyperlink:
                    self.assertRegex(cell.hyperlink.target, r"^(tel:\+[1-9][0-9]{6,14}|https://wa\.me/[1-9][0-9]{6,14})$")
        book.close()

    def test_original_attributes_reconstruct_and_formula_prefixes_are_literal(self):
        long_value = "=تفاصيل اصطناعية\n" + ("معلومة عربية " * 7000) + ("\U0001F4D6" * 1000)
        long_key = "مفتاح" * 7000
        controlled = "نص\x00يتضمن\x01محارف تحكم"
        originals = {"اسم العميل": "=1+1", "الميزانية": "-900", "بيانات طويلة": long_value,
                     "+حقل": "+SUM(A1)", "@حقل": "@SUM(A1)", long_key: "قيمة الاسم الطويل", "محارف": controlled}
        frame = export_frame(1)
        frame.loc[0, "original"] = json.dumps(originals, ensure_ascii=False)
        frame.loc[0, "details"] = json.dumps({"بيانات طويلة": long_value}, ensure_ascii=False)
        book = openpyxl.load_workbook(io.BytesIO(app.build_excel(frame, "جهة اصطناعية")))
        sheet = book["الخصائص الأصلية"]
        grouped = {}
        for row in sheet.iter_rows(min_row=8):
            if row[3].value != "أصلي":
                continue
            identifier = row[4].value
            entry = grouped.setdefault(identifier, {"keys": [], "values": [], "key_encoding": row[11].value, "value_encoding": row[12].value})
            if row[5].value:
                entry["keys"].append((row[5].value, row[9].value or ""))
            if row[7].value:
                entry["values"].append((row[7].value, row[10].value or ""))
            for cell in (row[9], row[10]):
                self.assertNotEqual(cell.data_type, "f")
                self.assertLess(len((cell.value or "").encode("utf-16-le")) // 2, 32767)
        reconstructed = {}
        for entry in grouped.values():
            key = "".join(value for _, value in sorted(entry["keys"]))
            value = "".join(value for _, value in sorted(entry["values"]))
            if "JSON" in entry["key_encoding"]:
                key = json.loads(key)
            if "JSON" in entry["value_encoding"]:
                value = json.loads(value)
            reconstructed[key] = value
        self.assertEqual(reconstructed, originals)
        self.assertIn("[مختصر]", book["العملاء"]["AA8"].value)
        self.assertIn("[مختصر]", book["العملاء"]["AB8"].value)
        book.close()

    def test_empty_export_keeps_headers_and_safe_summary(self):
        book = openpyxl.load_workbook(io.BytesIO(app.build_excel(export_frame(0), "جهة اصطناعية")))
        self.assertEqual(book["العملاء"].max_row, 7)
        self.assertEqual(book["العملاء"]["A7"].value, "معرّف العميل")
        self.assertIn("IFERROR", book["الملخص"]["B10"].value)
        book.close()


    def test_excel_does_not_create_temporary_customer_xml(self):
        with mock.patch("openpyxl.worksheet._writer.create_temporary_file", side_effect=AssertionError("Disk persistence is forbidden")):
            data = app.build_excel(export_frame(1), "جهة اصطناعية", png_bytes())
        book = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(book["العملاء"]["D8"].value, "عميل اصطناعي")
        book.close()


class PDFSafetyTests(unittest.TestCase):
    def test_html_escapes_every_user_context_and_contains_print_structure(self):
        payload = '<script>alert("x")</script>'
        frame = export_frame(1)
        frame.loc[0, "name"] = payload
        frame.loc[0, "source"] = '<img src="https://bad.invalid/x">'
        frame.loc[0, "city"] = '<b>مدينة</b>'
        frame.loc[0, "details"] = json.dumps({'<svg/onload=x>': '<iframe src=x>'}, ensure_ascii=False)
        frame.loc[0, "tel_url"] = "javascript:alert(1)"
        html = app.build_report_html(frame, payload, context_note=payload, generated_at=payload)
        self.assertNotIn("<script>", html)
        self.assertNotIn("<iframe", html)
        self.assertNotIn("<svg", html)
        self.assertNotIn('<img src="https://', html)
        self.assertIn("&lt;script&gt;", html)
        self.assertIn("&lt;svg/onload=x&gt;", html)
        self.assertIn("data:font/ttf;base64,", html)
        self.assertIn("size: A4 landscape", html)
        self.assertIn("<thead>", html)
        self.assertIn("<tfoot>", html)
        self.assertIn("counter(pages)", html)
        self.assertIn("break-inside: avoid", html)
        self.assertIn('href="tel:+201000000001"', html)
        self.assertIn('href="https://wa.me/966500000001"', html)
        self.assertIn("البيانات الكاملة في إكسل", html)
        self.assertIn("لا تؤكد وجود حساب واتساب", html)
        self.assertEqual(app._export_validate_worker_html(html.encode()), html)

    def test_url_fetcher_denies_every_nonapproved_resource(self):
        attempts = ["http://example.invalid/x", "https://example.invalid/x", "file:///etc/passwd",
                    "ftp://example.invalid/x", "/etc/passwd", "//example.invalid/x", "javascript:alert(1)",
                    "data:image/svg+xml;base64,PHN2Zy8+", "data:text/html;base64,PHNjcmlwdD4=",
                    "data:font/ttf;base64,YmFk", "DATA:image/png;base64,YQ==", "data:image/png;base64,YQ==\n",
                    "data:image/png;base64,YQ%3D%3D"]
        for url in attempts:
            with self.subTest(url=url):
                with self.assertRaises(ValueError):
                    app.safe_pdf_url_fetcher(url)
        font = app._export_font_bytes()
        result = app.safe_pdf_url_fetcher("data:font/ttf;base64," + base64.b64encode(font).decode())
        self.assertEqual(result["string"], font)
        self.assertEqual(result["mime_type"], "font/ttf")
        metadata_png = png_bytes(metadata=True)
        with self.assertRaises(ValueError):
            app.safe_pdf_url_fetcher("data:image/png;base64," + base64.b64encode(metadata_png).decode())

    def test_invalid_contact_links_cannot_inject_even_with_true_flags(self):
        frame = export_frame(1)
        frame.loc[0, "phone_e164"] = '+201000000001" onclick="x'
        frame.loc[0, "whatsapp_e164"] = "+966500000001?x=1"
        frame.loc[0, "phone_raw"] = '<a href="file:///etc/passwd">نص</a>'
        html = app.build_report_html(frame, "جهة اصطناعية")
        self.assertNotIn('href="tel:', html)
        self.assertNotIn('href="https://wa.me/', html)
        self.assertNotIn('onclick="x', html)

    def test_subset_kpis_are_not_all_input_kpis(self):
        frame = export_frame(3).iloc[:1]
        html = app.build_report_html(frame, "جهة اصطناعية")
        self.assertIn('<div class="metric-value">1</div><div class="metric-label">عملاء هذه الدفعة</div>', html)
        self.assertIn('<div class="metric-value">0</div><div class="metric-label">تكرارات محتملة</div>', html)

    def test_pdf_and_html_reject_over_limit_without_truncation(self):
        frame = export_frame(app.PDF_BATCH_SIZE + 1)
        with self.assertRaises(app.UserDataError):
            app.build_report_html(frame, "جهة اصطناعية")
        with self.assertRaises(app.UserDataError):
            app.build_pdf(frame, "جهة اصطناعية")
        self.assertEqual(len(frame), app.PDF_BATCH_SIZE + 1)

    def test_template_guard_forbids_css_and_active_html_changes(self):
        html = app.build_report_html(export_frame(1), "جهة اصطناعية")
        for changed in (html.replace("<body>", '<body onload="x">'),
                        html.replace("</body>", '<script>x</script></body>'),
                        html.replace("size: A4 landscape", "size: A4 portrait"),
                        html.replace('href="tel:+201000000001"', 'href="file:///etc/passwd"')):
            with self.subTest(changed=changed[-100:]):
                with self.assertRaises(ValueError):
                    app._export_validate_worker_html(changed.encode())
        with self.assertRaises(ValueError):
            app._export_validate_worker_html(b"x" * (app.PDF_MAX_HTML_BYTES + 1))

    def test_context_limit_and_semaphore_release(self):
        with self.assertRaises(app.UserDataError):
            app.build_report_html(export_frame(1), "جهة اصطناعية", context_note="س" * 601)
        with self.assertRaises(app.UserDataError):
            app.build_pdf(export_frame(1), "جهة اصطناعية", context_note="س" * 601)
        with mock.patch.object(app, "_export_worker_pdf", side_effect=app.UserDataError("اختبار اصطناعي")):
            with self.assertRaises(app.UserDataError):
                app.build_pdf(export_frame(1), "جهة اصطناعية")
        self.assertTrue(app._EXPORT_PDF_SEMAPHORE.acquire(blocking=False))
        try:
            with self.assertRaises(app.UserDataError):
                app.build_pdf(export_frame(1), "جهة اصطناعية")
        finally:
            app._EXPORT_PDF_SEMAPHORE.release()

    def test_worker_flag_is_exact_and_not_taken_during_normal_import(self):
        for args in (["app.py"], ["app.py", "--render-pdf-worker", "extra"]):
            with mock.patch.object(sys, "argv", args):
                self.assertFalse(app.handle_pdf_worker())

    def test_worker_timeout_kills_and_reaps_child(self):
        class HangingProcess:
            def __init__(self):
                self.stdin, self.stdout = io.BytesIO(), io.BytesIO()
                self.returncode = None
                self.killed = False

            def poll(self):
                return self.returncode

            def wait(self, timeout=None):
                if self.returncode is None:
                    raise app._export_subprocess.TimeoutExpired("synthetic worker", timeout)
                return self.returncode

            def kill(self):
                self.killed, self.returncode = True, -9

        child = HangingProcess()
        with mock.patch.object(app._export_subprocess, "Popen", return_value=child) as launch, \
                mock.patch.object(app, "PDF_RENDER_TIMEOUT_SECONDS", 0):
            with self.assertRaises(app.UserDataError):
                app._export_worker_pdf(b"synthetic fixture")
        self.assertTrue(child.killed)
        self.assertIsNotNone(child.returncode)
        self.assertTrue(child.stdout.closed)
        self.assertEqual(launch.call_args.args[0][-1], "--render-pdf-worker")
        self.assertEqual(launch.call_args.kwargs["stderr"], app._export_subprocess.DEVNULL)
        self.assertEqual(launch.call_args.kwargs["env"]["OPENBLAS_NUM_THREADS"], "1")

    def test_worker_limits_and_output_bounds_are_explicit(self):
        self.assertLessEqual(app.PDF_WORKER_MEMORY_BYTES, 700 * 1024 * 1024)
        self.assertLessEqual(app.PDF_RENDER_TIMEOUT_SECONDS, 90)
        self.assertLessEqual(app.PDF_MAX_HTML_BYTES, 4 * 1024 * 1024)
        child = mock.Mock()
        child.stdin, child.stdout = io.BytesIO(), io.BytesIO(b"%PDF-oversized synthetic output%%EOF")
        child.returncode = 0
        child.poll.return_value = 0
        child.wait.return_value = 0
        with mock.patch.object(app._export_subprocess, "Popen", return_value=child), \
                mock.patch.object(app, "PDF_MAX_OUTPUT_BYTES", 10):
            with self.assertRaises(app.UserDataError):
                app._export_worker_pdf(b"synthetic fixture")

    @unittest.skipUnless(importlib.util.find_spec("weasyprint") is not None and sys.platform.startswith("linux"),
                         "WeasyPrint and Linux resource isolation are required")
    def test_pdf_worker_smoke_when_renderer_is_installed(self):
        data = app.build_pdf(export_frame(1), "جهة اصطناعية", png_bytes(), context_note="دفعة اصطناعية للاختبار")
        self.assertTrue(data.startswith(b"%PDF-"))
        self.assertTrue(data.rstrip().endswith(b"%%EOF"))
        self.assertLess(len(data), app.PDF_MAX_OUTPUT_BYTES)


class SourceExportTests(unittest.TestCase):
    def test_demo_is_bom_synthetic_and_exercises_mixed_cases(self):
        data = app.demo_csv()
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(len(rows), 12)
        self.assertTrue(all("اصطناعية" in row["نوع البيانات"] for row in rows))
        phones = [row["رقم الموبايل"] for row in rows]
        self.assertTrue(any(value.startswith("+20") for value in phones))
        self.assertTrue(any(value.startswith("+966") for value in phones))
        self.assertTrue(any(value.startswith("+971") for value in phones))
        self.assertIn("", phones)
        self.assertIn("0123?", phones)
        self.assertLess(len(set(phones)), len(phones))
        self.assertTrue(any("غير جاهز" in row["موعد الشراء"] for row in rows))
        self.assertTrue(any("لست" in row["موعد الشراء"] for row in rows))

    def test_sources_preserve_exact_bytes_with_safe_unique_names(self):
        uploads = [("../../قائمة.csv", b"one"), ("C:\\private\\قائمة.csv", b"two"),
                   ("/قائمة.csv", b"three"), ("..", b"four"), ("CON.csv", b"five")]
        data = app.build_sources_archive(uploads)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self.assertEqual(len(archive.namelist()), len(uploads))
            self.assertEqual(len(set(name.casefold() for name in archive.namelist())), len(uploads))
            self.assertEqual([archive.read(name) for name in archive.namelist()], [data for _, data in uploads])
            for info in archive.infolist():
                self.assertNotIn("/", info.filename)
                self.assertNotIn("\\", info.filename)
                self.assertNotIn(":", info.filename)
                self.assertNotIn(info.filename, ("", ".", ".."))
                self.assertEqual((info.external_attr >> 16) & 0o170000, 0o100000)


if __name__ == "__main__":
    unittest.main()
